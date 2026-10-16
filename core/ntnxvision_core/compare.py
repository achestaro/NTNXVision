"""Snapshot comparison: what changed between two collections.

Objects are matched per type by their stable id (uuid / extId) — the same id Nutanix
keeps across collections — with (name, place) as the fallback for the few types that
have none. For matched objects every column marked ``diff`` in `exporters.tables` is
compared on its display value, so a change reads exactly as the report shows it
("32 GB" → "64 GB", not a raw byte count).

Raw changes are then **grouped** for the screen:

* consecutive added (or removed) objects of one type in one place whose names share a
  stem collapse into one row — ``k8s-worker-09 … 40`` with ``count`` 32;
* the same field change (same before → after) on several objects of one type in one place
  collapses likewise.

Summary counts are per object, never per row. `export_comparison` writes the ungrouped
list (one row per object and field) plus a summary sheet.
"""

from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ntnxvision_core.exporters.tables import COMPARE_TYPES, SPECS, Ctx, TableSpec
from ntnxvision_core.inventory import display, fmt_bytes
from ntnxvision_core.models import Environment
from ntnxvision_core.numfmt import NA, excel_number_format, fmt_count

_TRAILING_NUM = re.compile(r"^(.*?)(\d+)$")
_KIND_ORDER = {"changed": 0, "added": 1, "removed": 2}


def _where(spec: TableSpec, obj: Any, env: Environment) -> str:
    place = spec.where(obj)
    if place:
        return str(place)
    source = getattr(obj, "source_pc", None) or env.prism_central.host or ""
    return source.split(".")[0] if source else NA


def _key(spec: TableSpec, obj: Any, env: Environment) -> str:
    ident = spec.id(obj)
    if ident:
        return str(ident)
    return f"{spec.name(obj)}@{_where(spec, obj, env)}"


def raw_changes(base: Environment, cur: Environment) -> list[dict[str, Any]]:
    """One entry per added / removed object and per changed field (ungrouped)."""
    base_ctx, cur_ctx = Ctx(base), Ctx(cur)
    out: list[dict[str, Any]] = []
    for type_id in COMPARE_TYPES:
        spec = SPECS[type_id]
        # A type one of the two runs did not collect (area unselected, product not
        # deployed, collector failed) has no "before" or "after" — diffing it would
        # report every object as added or removed.
        if not (collected(base, spec) and collected(cur, spec)):
            continue
        before = {_key(spec, o, base): o for o in spec.items(base)}
        after = {_key(spec, o, cur): o for o in spec.items(cur)}
        # Fallback match by (name, place) for objects whose id changed between runs
        # (ESXi VMs re-registered, a PE-mode uuid vs a PC-mode one).
        unmatched_before = {k: o for k, o in before.items() if k not in after}
        by_name = {(spec.name(o), _where(spec, o, base)): k for k, o in unmatched_before.items()}
        for k, o in list(after.items()):
            if k in before:
                continue
            alt = by_name.get((spec.name(o), _where(spec, o, cur)))
            if alt is not None and alt in before:
                before[k] = before.pop(alt)
        for k, obj in after.items():
            if k not in before:
                out.append(_entry("added", spec, obj, cur, cur_ctx))
        for k, obj in before.items():
            if k not in after:
                out.append(_entry("removed", spec, obj, base, base_ctx, cur=cur))
        diff_cols = [c for c in spec.cols if c.diff]
        for k, new in after.items():
            old = before.get(k)
            if old is None:
                continue
            for col in diff_cols:
                a = display(col, col.get(old, base_ctx))
                b = display(col, col.get(new, cur_ctx))
                if a != b:
                    out.append({
                        "kind": "changed", "type": type_id, "object_type": spec.singular,
                        "object": str(spec.name(new)), "where": _where(spec, new, cur),
                        "field": col.diff, "before": a or NA, "after": b or NA,
                        "detail": None, "id": k,
                    })
    return out


def collected(env: Environment, spec: TableSpec) -> bool:
    """Did this snapshot's run collect `spec`'s area at all?"""
    if spec.items(env):
        return True
    if not spec.note_keys:
        return True
    notes = env.collection_notes
    for key in spec.note_keys:
        for k, v in notes.items():
            if (k == key or k.endswith(f": {key}")) and v.startswith("ok"):
                return True
    return False


def _entry(kind: str, spec: TableSpec, obj: Any, env: Environment, ctx: Ctx,
           cur: Environment | None = None) -> dict[str, Any]:
    return {
        "kind": kind, "type": spec.type, "object_type": spec.singular,
        "object": str(spec.name(obj)), "where": _where(spec, obj, env),
        "field": None, "before": None, "after": None,
        "detail": _describe(spec, obj, ctx, kind, cur), "id": _key(spec, obj, env),
        "_obj": obj,
    }


def _describe(spec: TableSpec, obj: Any, ctx: Ctx, kind: str, cur: Environment | None
              ) -> str | None:
    t = spec.type
    if t == "vms":
        return " · ".join(x for x in (f"{fmt_count(obj.vcpu_total or NA)} vCPU",
                                      fmt_bytes(obj.memory_bytes), obj.guest_os) if x)
    if t == "hosts":
        return " · ".join(x for x in (obj.block_model, f"{fmt_count(obj.num_cpu_cores)} cores"
                                      if obj.num_cpu_cores else None) if x)
    if t == "storage_containers":
        return " · ".join(x for x in (fmt_bytes(obj.max_capacity_bytes),
                                      "compression on" if obj.compression_enabled
                                      else "compression off") if x)
    if t == "security_policies":
        return f"{(obj.policy_type or 'Policy').title()} · {(obj.state or '').lower()}"
    if t == "protection_policies" and kind == "removed" and cur is not None:
        cats = set(obj.categories)
        cur_ctx = Ctx(cur)
        orphaned = sum(
            1 for v in cur.vms
            if any(f"{k}:{val}" in cats for k, vals in v.categories.items() for val in vals)
            and cur_ctx.vm_policy(v) is None)
        if orphaned:
            return f"{fmt_count(orphaned)} VMs it covered are now unprotected"
        return "No VMs were relying on it"
    if t == "clusters":
        return f"AOS {obj.aos_version}" if obj.aos_version else None
    parts = []
    for col in spec.ui_cols[1:3]:
        value = display(col, col.get(obj, ctx))
        if value:
            parts.append(value)
    return " · ".join(parts) or None


# ─── grouping ──────────────────────────────────────────────────────────


def _stem(name: str) -> tuple[str, int | None]:
    m = _TRAILING_NUM.match(name)
    if not m:
        return name, None
    return m.group(1), int(m.group(2))


def range_label(names: list[str]) -> str:
    """'k8s-worker-09 … 40' — first full name, then the last one's tail after the
    common prefix (backed off to the last separator so 'n5 … n8' keeps its letter)."""
    first, last = names[0], names[-1]
    if first == last:
        return f"{first} ×{fmt_count(len(names))}" if len(names) > 1 else first
    i = 0
    while i < min(len(first), len(last)) and first[i] == last[i]:
        i += 1
    cut = max(first.rfind(sep, 0, i) for sep in "-_. ") + 1
    return f"{first} … {last[cut:]}"


def _names_label(names: list[str]) -> str:
    """A run of same-stem names as a range; anything else as a short list."""
    if len(names) == 1:
        return names[0]
    stems = {_stem(n)[0] for n in names}
    if len(stems) == 1 and all(_stem(n)[1] is not None for n in names):
        return range_label(names)
    if len(names) <= 3:
        return ", ".join(names)
    return f"{names[0]}, {names[1]} + {fmt_count(len(names) - 2)} more"


def group_changes(raw: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    # Added / removed: runs of same-stem names within (kind, type, where).
    buckets: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    changed: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
    for c in raw:
        if c["kind"] == "changed":
            changed.setdefault((c["type"], c["where"], c["field"], c["before"], c["after"]),
                               []).append(c)
        else:
            buckets.setdefault((c["kind"], c["type"], c["where"]), []).append(c)
    for (kind, type_id, where), items in buckets.items():
        items.sort(key=lambda c: (_stem(c["object"])[0], _stem(c["object"])[1] or -1,
                                  c["object"]))
        i = 0
        while i < len(items):
            stem, num = _stem(items[i]["object"])
            j = i + 1
            if num is not None:
                while j < len(items):
                    s2, n2 = _stem(items[j]["object"])
                    if s2 != stem or n2 is None:
                        break
                    j += 1
            rows.append(_group_row(kind, type_id, where, items[i:j]))
            i = j
    for (type_id, where, field, before, after), items in changed.items():
        names = sorted((c["object"] for c in items), key=lambda n: (_stem(n)[0],
                                                                   _stem(n)[1] or -1, n))
        rows.append({
            "kind": "changed", "object_type": items[0]["object_type"],
            "object": _names_label(names),
            "where": where, "field": field, "before": before, "after": after,
            "detail": None, "count": len(items), "_type": type_id,
        })
    rows.sort(key=lambda r: (COMPARE_TYPES.index(r["_type"]) if r["_type"] in COMPARE_TYPES
                             else 99, _KIND_ORDER[r["kind"]], r["where"] or "", r["object"]))
    for r in rows:
        r.pop("_type", None)
    return rows


def _group_row(kind: str, type_id: str, where: str, items: list[dict[str, Any]]
               ) -> dict[str, Any]:
    spec = SPECS[type_id]
    if len(items) == 1:
        c = items[0]
        return {"kind": kind, "object_type": c["object_type"], "object": c["object"],
                "where": where, "field": None, "before": None, "after": None,
                "detail": c["detail"], "count": 1, "_type": type_id}
    names = [c["object"] for c in items]
    detail = f"{fmt_count(len(items))} {spec.unit_plural}"
    if type_id == "vms":
        objs = [c["_obj"] for c in items]
        sizes = {(o.vcpu_total, o.memory_bytes) for o in objs}
        if len(sizes) == 1:
            vcpu, mem = sizes.pop()
            detail += f" · {fmt_count(vcpu or NA)} vCPU · {fmt_bytes(mem)}"
        else:
            detail += f" · {fmt_count(sum(o.vcpu_total or 0 for o in objs))} vCPU · " \
                      f"{fmt_bytes(sum(o.memory_bytes or 0 for o in objs))} total"
        if kind == "removed":
            freed = sum(o.total_disk_bytes for o in objs)
            if freed:
                detail += f" · {fmt_bytes(freed)} freed"
    elif type_id == "hosts":
        objs = [c["_obj"] for c in items]
        models = {o.block_model for o in objs if o.block_model}
        detail = f"{fmt_count(len(items))} nodes" + (f" · {models.pop()}"
                                                     if len(models) == 1 else "")
        cores = {o.num_cpu_cores for o in objs}
        if len(cores) == 1 and None not in cores:
            detail += f" · {fmt_count(cores.pop())} cores each"
    return {"kind": kind, "object_type": spec.singular, "object": range_label(names),
            "where": where, "field": None, "before": None, "after": None, "detail": detail,
            "count": len(items), "_type": type_id}


# ─── figures ───────────────────────────────────────────────────────────


def capacity(env: Environment) -> tuple[int, int]:
    """(capacity, used) bytes across clusters.

    Containers on one cluster share its storage pool and each reports the pool as its
    ceiling, so the pool is counted once per cluster (the largest ceiling) while usage is
    summed across containers.
    """
    pools: dict[str, int] = {}
    used = 0
    for s in env.storage_containers:
        key = f"{s.source_pc or ''}/{s.cluster_uuid or s.cluster_name or s.uuid}"
        pools[key] = max(pools.get(key, 0), s.max_capacity_bytes or 0)
        used += s.used_bytes or 0
    return sum(pools.values()), used


def _capacity(env: Environment) -> int:
    return capacity(env)[0]


def _tb(value: int) -> str:
    return f"{value / 1024**4:,.0f}" if value >= 10 * 1024**4 else f"{value / 1024**4:,.1f}"


def _delta(added: int, removed: int, was: int) -> str:
    parts = []
    if added and removed:
        parts.append(f"+{fmt_count(added)} / −{fmt_count(removed)}")
    elif added:
        parts.append(f"+{fmt_count(added)}")
    elif removed:
        parts.append(f"−{fmt_count(removed)}")
    else:
        parts.append("no change")
    parts.append(f"was {fmt_count(was)}")
    return " · ".join(parts)


def compare(base: Environment, cur: Environment) -> dict[str, Any]:
    raw = raw_changes(base, cur)
    summary = {
        "added": sum(1 for c in raw if c["kind"] == "added"),
        "removed": sum(1 for c in raw if c["kind"] == "removed"),
        # Per object, not per field: two fields changed on one VM is one changed VM.
        "changed": len({(c["type"], c["id"]) for c in raw if c["kind"] == "changed"}),
    }
    summary["total"] = summary["added"] + summary["removed"] + summary["changed"]

    def counts(type_id: str) -> tuple[int, int]:
        return (sum(1 for c in raw if c["type"] == type_id and c["kind"] == "added"),
                sum(1 for c in raw if c["type"] == type_id and c["kind"] == "removed"))

    va, vr = counts("vms")
    ha, hr = counts("hosts")
    pa, pr = counts("protection_policies")
    cap_b, cap_c = _capacity(base), _capacity(cur)
    cap_delta = cap_c - cap_b
    figures = [
        {"label": "Changes", "value": summary["total"],
         "qualifier": f"{fmt_count(summary['added'])} added · "
                      f"{fmt_count(summary['removed'])} removed · "
                      f"{fmt_count(summary['changed'])} changed"},
        {"label": "VMs", "value": len(cur.vms), "qualifier": _delta(va, vr, len(base.vms))},
        {"label": "Hosts", "value": len(cur.hosts), "qualifier": _delta(ha, hr, len(base.hosts))},
        {"label": "Storage", "value": _tb(cap_c), "unit": "TB",
         "qualifier": (f"{'+' if cap_delta >= 0 else '−'}{_tb(abs(cap_delta))} TB · "
                       if cap_delta else "no change · ") + f"was {_tb(cap_b)} TB"},
        {"label": "Protection policies", "value": len(cur.protection_policies),
         "qualifier": _delta(pa, pr, len(base.protection_policies)),
         "qualifier_tone": "warn" if pr else None},
    ]
    return {"summary": summary, "figures": figures, "changes": group_changes(raw),
            "days": _days(base.collected_at, cur.collected_at)}


def _days(a: str | None, b: str | None) -> float:
    try:
        da = datetime.fromisoformat(str(a).replace("Z", "+00:00"))
        db = datetime.fromisoformat(str(b).replace("Z", "+00:00"))
    except ValueError:
        return 0.0
    return round(abs((db - da).total_seconds()) / 86400, 2)


# ─── export ────────────────────────────────────────────────────────────

_HEADERS = ["Change", "Object type", "Object", "Where", "Field", "Before", "After", "Detail",
            "Id"]


def export_comparison(base: Environment, cur: Environment, base_name: str, cur_name: str,
                      fmt: str, directory: Path) -> Path:
    raw = raw_changes(base, cur)
    result = compare(base, cur)
    directory.mkdir(parents=True, exist_ok=True)
    stem = f"changes-{base_name}-to-{cur_name}"
    rows = [[c["kind"], c["object_type"], c["object"], c["where"], c["field"], c["before"],
             c["after"], c["detail"], c["id"]] for c in raw]
    if fmt == "csv":
        out = directory / f"{stem}.csv"
        with out.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow(_HEADERS)
            writer.writerows(rows)
        return out
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    ws.append(["NTNXVision · What changed"])
    ws["A1"].font = Font(size=14, bold=True, color="1F4E78")
    ws.append([])
    meta = [
        ("Baseline", base_name), ("Baseline collected", base.collected_at),
        ("Current", cur_name), ("Current collected", cur.collected_at),
        ("Days between", result["days"]), ("", ""),
        ("Total changes", result["summary"]["total"]),
        ("Added", result["summary"]["added"]), ("Removed", result["summary"]["removed"]),
        ("Changed", result["summary"]["changed"]), ("", ""),
    ]
    for label, value in meta:
        ws.append([label, value])
        number_format = excel_number_format(value)
        if number_format:
            ws.cell(row=ws.max_row, column=2).number_format = number_format
    for fig in result["figures"][1:]:
        ws.append([fig["label"],
                   f"{fmt_count(fig['value'])}{' ' + fig['unit'] if fig.get('unit') else ''}",
                   fig.get("qualifier")])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 36
    sheet = wb.create_sheet("Changes")
    sheet.append(_HEADERS)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(len(rows) + 1, 1)}"
    for col, width in zip("ABCDEFGHI", (10, 18, 32, 18, 20, 22, 22, 44, 38), strict=True):
        sheet.column_dimensions[col].width = width
    out = directory / f"{stem}.xlsx"
    wb.save(out)
    return out
