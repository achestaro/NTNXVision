"""Excel exporter (openpyxl).

Produces an RVTools-style workbook: an Overview sheet plus one sheet per object type,
each with a frozen, styled header row and an auto-filter. Byte-valued fields are rendered
in GiB for readability. A sheet is omitted entirely when its object type has zero rows
(feature not deployed / not licensed) — Feature Usage and Collection Notes still record
that status, so the workbook stays a place to check "is X in use" without a noise sheet
per absent feature.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ntnxvision_core.exporters import tables as T
from ntnxvision_core.exporters.sourcing import source_hosts as _source_hosts
from ntnxvision_core.exporters.sourcing import sources as _sources
from ntnxvision_core.exporters.tables import Ctx
from ntnxvision_core.models import Environment

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=14, bold=True, color="1F4E78")
GIB = 1024**3


class ExcelExporter:
    #: Number of sheets written by the last export (for the "34 sheets" report detail).
    sheet_count: int = 0

    def export(self, env: Environment, path: str | Path) -> Path:
        wb = Workbook()
        ctx = Ctx(env)
        # Phase 1 — core infrastructure
        self._overview_sheet(wb, env)
        self._pc_config_sheet(wb, env)
        for spec in (T.CLUSTERS, T.HOSTS, T.VMS, T.SYSTEM_VMS, T.STORAGE, T.SUBNETS,
                     # PC catalog — templates / images / OVAs
                     T.TEMPLATES, T.IMAGES, T.OVAS,
                     # Phase 2 — data protection / DR
                     T.PROTECTION_POLICIES, T.RECOVERY_PLANS, T.AVAILABILITY_ZONES,
                     T.REMOTE_SITES,
                     # Phase 2 — Health / Alerts, then LCM firmware / software currency
                     T.ALERTS, T.LCM,
                     # Phase 2 — Flow
                     T.SECURITY_POLICIES, T.CATEGORIES, T.SERVICE_GROUPS, T.ADDRESS_GROUPS,
                     T.VPCS, T.FLOATING_IPS,
                     # Phase 2 — NUS
                     T.FILE_SERVERS, T.FILE_SHARES, T.OBJECT_STORES, T.BUCKETS,
                     T.VOLUME_GROUPS,
                     # Phase 2 — NDB
                     T.DATABASES, T.DATABASE_SERVERS, T.TIME_MACHINES,
                     # Phase 2 — IAM / automation
                     T.USERS, T.ROLES, T.PROJECTS, T.PLAYBOOKS, T.NCM_REPORTS,
                     # Licensing
                     T.LICENSES):
            self._spec_sheet(wb, env, spec, ctx)
        # Analysis + diagnostics
        self._feature_usage_sheet(wb, env)
        self._notes_sheet(wb, env)

        out = Path(path)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
        self.sheet_count = len(wb.sheetnames)
        return out

    def _spec_sheet(self, wb: Workbook, env: Environment, spec: T.TableSpec, ctx: Ctx) -> None:
        """One object-type sheet, from the column definitions the Inventory grid shares.

        The LCM sheet's `target_version` is deliberately not a column: despite the name it
        is the version LCM currently aims at (the running one unless an upgrade is
        staged), so it would only duplicate Current Version. It stays in the snapshot.
        """
        items = spec.items(env)
        headers = [c.header or "" for c in spec.excel_cols]
        rows = T.table_rows(spec, env, ctx)
        self._table(wb, spec.sheet, headers, rows, _sources(env, items) if spec.sources
                    else None)

    # ─── sheets ────────────────────────────────────────────────────

    def _overview_sheet(self, wb: Workbook, env: Environment) -> None:
        ws = cast(Worksheet, wb.active)
        ws.title = "Overview"
        ws["A1"] = "NTNXVision — Environment Inventory"
        ws["A1"].font = _TITLE_FONT
        # A combined run has no single management plane: name every contributing PC
        # rather than reporting a synthetic host with an "unknown" version.
        if env.is_combined:
            pc_rows: list[tuple[str, Any]] = [
                ("Prism Centrals", f"{len(env.source_environments)} (combined)"),
                ("Sources", ", ".join(_source_hosts(env))),
            ]
        else:
            pc_rows = [
                ("Prism Central", env.prism_central.host or ""),
                ("PC Version", env.prism_central.version or "unknown"),
            ]
        rows: list[tuple[str, Any]] = [
            *pc_rows,
            ("Collected At (UTC)", env.collected_at or ""),
            ("Tool Version", env.tool_version),
            ("", ""),
        ]
        for label, value in env.summary().items():
            rows.append((label.replace("_", " ").title(), value))
        for i, (label, value) in enumerate(rows, start=3):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 44

    def _pc_config_sheet(self, wb: Workbook, env: Environment) -> None:
        if env.is_combined:
            self._source_pcs_sheet(wb, env)
            return
        cfg = env.pc_config
        headers = ["Property", "Value"]
        rows: list[list[Any]] = [
            ["Version", cfg.version],
            ["Deployment", "Scale-out" if cfg.is_scale_out else "Single-node"
             if cfg.is_scale_out is not None else None],
            ["PC Nodes", cfg.node_count],
            ["Registered Clusters", cfg.registered_cluster_count],
            ["Total vCPU", cfg.num_vcpus],
            ["Total Memory (GiB)", _gib(cfg.memory_bytes)],
            ["NCC Version", cfg.ncc_version],
            ["PC VM(s)", ", ".join(cfg.pc_vm_names)],
            ["IP Address(es)", ", ".join(cfg.ip_addresses)],
        ]
        self._table(wb, "Prism Central", headers, rows)

    def _source_pcs_sheet(self, wb: Workbook, env: Environment) -> None:
        """One row per contributing Prism Central — the combined-run replacement for the
        single-PC property sheet, which would otherwise have to pick one PC's sizing and
        present it as the whole estate's.
        """
        headers = [
            "Prism Central", "Version", "Deployment", "PC Nodes", "Registered Clusters",
            "Total vCPU", "Total Memory (GiB)", "NCC Version", "IP Address(es)",
            "Clusters", "Hosts", "VMs",
        ]
        rows: list[list[Any]] = []
        for s in env.source_environments:
            cfg, counts = s.pc_config, s.summary
            rows.append([
                s.prism_central.host or "unknown",
                s.prism_central.version or cfg.version,
                "Scale-out" if cfg.is_scale_out else "Single-node"
                if cfg.is_scale_out is not None else None,
                cfg.node_count, cfg.registered_cluster_count, cfg.num_vcpus,
                _gib(cfg.memory_bytes), cfg.ncc_version, ", ".join(cfg.ip_addresses),
                counts.get("clusters"), counts.get("hosts"), counts.get("vms"),
            ])
        self._table(wb, "Prism Centrals", headers, rows)

    def _feature_usage_sheet(self, wb: Workbook, env: Environment) -> None:
        from ntnxvision_core.analysis import analyze_feature_usage

        rows = [fu.as_row() for fu in analyze_feature_usage(env)]
        self._table(wb, "Feature Usage", ["Product", "Feature", "In Use", "Detail"], rows)

    def _notes_sheet(self, wb: Workbook, env: Environment) -> None:
        from ntnxvision_core.notes import clean_error, pretty_label

        headers = ["Area", "Outcome", "Detail"]
        rows: list[list[Any]] = []
        for key, status in env.collection_notes.items():
            if status == "ok":
                outcome, detail = "Collected", ""
            elif status.startswith("skipped"):
                outcome = "Not present"
                detail = status.split(":", 1)[1].strip() if ":" in status else status
            else:
                outcome, detail = "Error", clean_error(status)
            rows.append([pretty_label(key), outcome, detail])
        for w in env.warnings:
            rows.append(["Compatibility", "Warning", w])
        self._table(wb, "Collection Notes", headers, rows)

    # ─── table helper ──────────────────────────────────────────────

    def _table(
        self, wb: Workbook, title: str, headers: list[str], rows: list[list[Any]],
        sources: list[str] | None = None,
    ) -> None:
        """Write one sheet. `sources` (combined runs only) prepends a "Prism Central"
        column naming the PC each row came from — object names are not unique across
        PCs, so without it two merged rows can be indistinguishable.
        """
        if not rows:
            return
        if sources is not None:
            headers = ["Prism Central", *headers]
            rows = [[src, *row] for src, row in zip(sources, rows, strict=True)]
        ws: Worksheet = wb.create_sheet(_safe_title(title))
        ws.append(headers)
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append(row)

        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 1)}"
        self._autosize(ws, headers, rows)

    @staticmethod
    def _autosize(ws: Worksheet, headers: list[str], rows: list[list[Any]]) -> None:
        for col_idx, header in enumerate(headers, start=1):
            width = len(str(header))
            for row in rows:
                if col_idx - 1 < len(row) and row[col_idx - 1] is not None:
                    width = max(width, len(str(row[col_idx - 1])))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 50)


def _safe_title(title: str) -> str:
    """Excel sheet titles disallow []:*?/\\ and cap at 31 chars."""
    for ch in "[]:*?/\\":
        title = title.replace(ch, "-")
    return title[:31]


def _update_flag(value: bool | None) -> str:
    """Yes / No / Unknown for an LCM update flag.

    None means LCM did not tell us whether a newer version exists; rendering that as
    "No" would assert the component is current, which is exactly the wrong answer to
    the one question this sheet is here to answer.
    """
    if value is None:
        return "Unknown"
    return "Yes" if value else "No"


def _gib(value: int | None) -> float | None:
    if not value:
        return None
    return round(value / GIB, 2)
