"""Word Architecture Design Document (ADD) exporter (python-docx).

This is NTNXVision's differentiator: rather than a raw data dump, it produces a narrative
architecture design deliverable: cover page, contents, executive summary, assumptions,
logical and physical design, data protection, security posture, storage services,
databases, licensing, lifecycle and an appendix on collection coverage and API routing.

The content comes from `content.build_report`, the same model the PDF renders, so both
documents carry the same chapters, numbering, tables, counts and wording. Optional
chapters are left out when the environment has no data for them (Licensing & Feature
Usage and Assumptions & Constraints say why), and chapter numbers close over the gap.

Styling is the light Runway theme in Office-safe form: Calibri text in dark ink, an iris
cover rule, an iris marker on Heading 1 and an iris rule under each table header, cool
grey table rules, and status shown as a small coloured square beside an ink label.
Headings use the real Heading 1 / Heading 2 styles so Word's navigation pane and the
contents field work; the contents field is pre-filled with the headings and Word is asked
to refresh it (page numbers) when the file is opened. A customer ``.dotx`` / ``.docx``
template keeps its own styles: only direct table and cover formatting is applied then.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

from ntnxvision_core.diagrams.render.png import render_png
from ntnxvision_core.exporters import theme as T
from ntnxvision_core.exporters.content import (
    Bullets,
    Chapter,
    Figure,
    FigureRow,
    FigureStrip,
    Paragraph,
    Report,
    SceneFigure,
    Status,
    Table,
    build_report,
    cell_text,
)
from ntnxvision_core.models import Environment
from ntnxvision_core.numfmt import NA

#: The status glyph: a small filled square in the status colour, then the ink label.
GLYPH = "■"
#: The narrowest design-summary column beside a figure (pt).
SUMMARY_MIN_W = 170.0
#: Tables with at most this many rows keep together on one page.
SMALL_TABLE = 15

_INK = RGBColor.from_string(T.INK)
_INK_2 = RGBColor.from_string(T.INK_2)
_INK_3 = RGBColor.from_string(T.INK_3)
_IRIS = RGBColor.from_string(T.IRIS)


class WordExporter:
    #: Top-level chapters and data tables written by the last export (report detail line).
    section_count: int = 0
    table_count: int = 0

    def export(self, env: Environment, path: str | Path, options: Any = None) -> Path:
        from ntnxvision_core.exporters import ExportOptions

        opts = ExportOptions.coerce(options)
        custom = opts.template not in ("", "standard")
        doc = load_template(opts.template) if custom else Document()
        report = build_report(env, opts, eyebrow="Architecture design document")
        self._tables = 0
        #: The current section is landscape / starts on a fresh page (a section break).
        self._landscape = False
        self._fresh = False
        if not custom:
            self._page_setup(doc)
            self._styles(doc)
        _ensure_styles(doc)
        self._properties(doc, report, env, opts)
        self._cover(doc, report)
        self._toc(doc, report)
        for chapter in report.chapters:
            self._chapter(doc, chapter)
        if not custom:  # a customer template keeps its own header and footer
            self._header_footer(doc, report)
        _update_fields_on_open(doc)

        out = Path(path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))
        self.section_count = len(report.chapters)
        self.table_count = self._tables
        return out

    # ─── document scaffolding ──────────────────────────────────────

    def _page_setup(self, doc: DocumentObject) -> None:
        for section in doc.sections:
            section.left_margin = section.right_margin = Inches(0.8)
            section.top_margin = Inches(0.9)
            section.bottom_margin = Inches(0.8)
            section.header_distance = Inches(0.4)
            section.footer_distance = Inches(0.35)

    def _styles(self, doc: DocumentObject) -> None:
        styles = doc.styles
        normal = styles["Normal"]
        _font(normal, T.OFFICE_FONT, 10.5, _INK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        _lang(normal, "en-US")

        h1 = styles["Heading 1"]
        _font(h1, T.OFFICE_FONT, 18, _INK, bold=True)
        h1.paragraph_format.space_before = Pt(0)
        h1.paragraph_format.space_after = Pt(12)
        h1.paragraph_format.page_break_before = True
        h1.paragraph_format.keep_with_next = True
        _border(h1.element.get_or_add_pPr(), "left", T.IRIS, size=24, space=8)

        h2 = styles["Heading 2"]
        _font(h2, T.OFFICE_FONT, 13, _INK, bold=True)
        h2.paragraph_format.space_before = Pt(16)
        h2.paragraph_format.space_after = Pt(6)
        h2.paragraph_format.keep_with_next = True

        h3 = styles["Heading 3"]
        _font(h3, T.OFFICE_FONT, 11, _INK_2, bold=True)

        title = styles["Title"]
        _font(title, T.OFFICE_FONT, 30, _INK)
        _clear_border(title.element.get_or_add_pPr())
        _clear_spacing(title.element)
        title.paragraph_format.space_after = Pt(6)
        subtitle = styles["Subtitle"]
        _font(subtitle, T.OFFICE_FONT, 13, _INK_2, italic=False)
        _clear_spacing(subtitle.element)
        subtitle.paragraph_format.space_after = Pt(4)

        caption = styles["Caption"]
        _font(caption, T.OFFICE_FONT, 8.5, _INK_2, italic=False, bold=False)
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(3)

        bullet = styles["List Bullet"]
        _font(bullet, T.OFFICE_FONT, 10.5, _INK)
        bullet.paragraph_format.space_after = Pt(3)

    def _properties(self, doc: DocumentObject, report: Report, env: Environment,
                    opts: Any) -> None:
        props = doc.core_properties
        target = (f"{len(env.source_environments):,} Prism Centrals" if env.is_combined
                  else env.prism_central.host or "Nutanix environment")
        props.title = f"Architecture Design Document: {target}"
        props.subject = report.subject
        props.author = opts.author or "NTNXVision"
        props.last_modified_by = opts.author or "NTNXVision"
        props.category = "Architecture Design Document"
        props.comments = ("Generated by NTNXVision from a read-only inventory snapshot "
                          f"collected {T.collected_label(env)}.")
        props.keywords = report.keywords
        props.language = "en-US"
        props.created = props.modified = datetime.now()
        props.revision = 1

    # ─── cover & contents ──────────────────────────────────────────

    def _cover(self, doc: DocumentObject, report: Report) -> None:
        cover = report.cover
        band = doc.add_paragraph()
        _border(band._p.get_or_add_pPr(), "top", T.IRIS, size=36, space=10)
        band.paragraph_format.space_after = Pt(14)
        word = band.add_run("NTNXVision")
        _run(word, 10.5, _INK, bold=True)

        eyebrow = doc.add_paragraph()
        eyebrow.paragraph_format.space_before = Pt(96)
        eyebrow.paragraph_format.space_after = Pt(6)
        r = eyebrow.add_run(cover.eyebrow.upper())
        _run(r, 8.5, _INK_3, bold=True)
        _letter_spacing(r, 20)

        doc.add_paragraph(T.clean(cover.title), style="Title")
        sub = doc.add_paragraph(T.clean(cover.subtitle), style="Subtitle")
        sub.paragraph_format.space_after = Pt(10)
        rule = doc.add_paragraph()
        rule.paragraph_format.right_indent = Inches(5.9)
        rule.paragraph_format.space_after = Pt(40)
        _border(rule._p.get_or_add_pPr(), "bottom", T.IRIS, size=12, space=1)

        for label, value in cover.meta:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(1.55))
            p.paragraph_format.left_indent = Inches(1.55)
            p.paragraph_format.first_line_indent = -Inches(1.55)
            _run(p.add_run(f"{label}\t"), 9, _INK_3)
            _run(p.add_run(T.clean(value) or NA), 10, _INK)

        note = doc.add_paragraph()
        note.paragraph_format.space_before = Pt(60)
        _border(note._p.get_or_add_pPr(), "top", T.RULE_2, size=4, space=6)
        _run(note.add_run(T.clean(cover.note)), 8.5, _INK_3, italic=True)
        doc.add_page_break()

    def _toc(self, doc: DocumentObject, report: Report) -> None:
        """Contents: a real TOC field (Heading 1 and 2, hyperlinked), pre-filled with the
        headings so it reads correctly before Word refreshes the page numbers."""
        title = doc.add_paragraph("Contents", style="Contents Heading")
        title.paragraph_format.keep_with_next = True
        outline = report.outline()
        first = doc.add_paragraph(style="TOC 1" if outline and outline[0][0] == 0 else None)
        _field_begin(first, 'TOC \\o "1-2" \\h \\z \\u')
        paragraphs = [first]
        for i, (level, text) in enumerate(outline):
            p = first if i == 0 else doc.add_paragraph(style=f"TOC {level + 1}")
            p.add_run(T.clean(text))
            paragraphs.append(p)
        _field_end(paragraphs[-1])
        hint = doc.add_paragraph()
        hint.paragraph_format.space_before = Pt(10)
        _run(hint.add_run("Page numbers are filled in when Word opens the document; to "
                          "refresh them later, right-click the contents and choose Update "
                          "Field."), 8.5, _INK_3, italic=True)

    # ─── chapters ──────────────────────────────────────────────────

    def _chapter(self, doc: DocumentObject, ch: Chapter) -> None:
        """A chapter; sections holding a figure wider than the portrait text width go
        in a landscape section, so architecture figures are 1:1 on the page. When the
        first section is wide the whole chapter opening goes landscape with it."""
        wide = [self._is_wide(doc, s.blocks) for s in ch.sections]
        self._orientation(doc, bool(wide) and wide[0])
        h = doc.add_heading(T.clean(ch.heading), level=1)
        h.paragraph_format.page_break_before = not self._fresh
        h.paragraph_format.keep_with_next = True
        self._fresh = False
        self._blocks(doc, ch.blocks)
        for section, is_wide in zip(ch.sections, wide, strict=True):
            if is_wide != self._landscape:
                self._orientation(doc, is_wide)
            head = doc.add_heading(T.clean(section.heading), level=2)
            head.paragraph_format.keep_with_next = True
            if self._fresh:
                head.paragraph_format.space_before = Pt(0)
            self._fresh = False
            self._blocks(doc, section.blocks)

    @staticmethod
    def _portrait_width(doc: DocumentObject) -> int:
        s = doc.sections[-1]
        w, h = int(s.page_width or Inches(8.5)), int(s.page_height or Inches(11))
        return min(w, h) - int(s.left_margin or 0) - int(s.right_margin or 0)

    def _is_wide(self, doc: DocumentObject, blocks: list[Any]) -> bool:
        """A figure (with its summary column when it sits beside it) wider than the
        portrait text width needs a landscape page."""
        limit = self._portrait_width(doc)
        return any(isinstance(b, SceneFigure)
                   and int(Pt(b.scene.width + (SUMMARY_MIN_W if b.summary_beside else 0)))
                   > limit for b in blocks)

    def _orientation(self, doc: DocumentObject, landscape: bool) -> None:
        """Start a new section (on a new page) in the wanted orientation, if needed."""
        if landscape == self._landscape:
            return
        cur = doc.sections[-1]
        w, h = int(cur.page_width or Inches(8.5)), int(cur.page_height or Inches(11))
        new = doc.add_section(WD_SECTION.NEW_PAGE)
        new.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        new.page_width = Emu(max(w, h) if landscape else min(w, h))
        new.page_height = Emu(min(w, h) if landscape else max(w, h))
        self._landscape = landscape
        self._fresh = True

    def _blocks(self, doc: DocumentObject, blocks: list[Any]) -> None:
        for blk in blocks:
            if isinstance(blk, Paragraph):
                p = doc.add_paragraph()
                if blk.style == "note":
                    _run(p.add_run(T.clean(blk.text)), 9.5, _INK_3, italic=True)
                else:
                    p.add_run(T.clean(blk.text))
            elif isinstance(blk, Bullets):
                if blk.intro:
                    intro = doc.add_paragraph()
                    intro.paragraph_format.keep_with_next = True
                    intro.paragraph_format.space_before = Pt(6)
                    _run(intro.add_run(T.clean(blk.intro)), None, _INK, bold=True)
                for item in blk.items:
                    doc.add_paragraph(T.clean(item), style="List Bullet")
            elif isinstance(blk, FigureStrip):
                self._strip(doc, blk)
            elif isinstance(blk, Table):
                self._table(doc, blk)
            elif isinstance(blk, Figure):
                self._figure(doc, blk)
            elif isinstance(blk, SceneFigure):
                self._scene_figure(doc, blk)
            elif isinstance(blk, FigureRow):
                self._figure_row(doc, blk)

    def _strip(self, doc: DocumentObject, strip: FigureStrip) -> None:
        """Headline figures in one ruled row: small label, light numeral, qualifier."""
        if not strip.items:
            return
        width = _text_width(doc)
        table = doc.add_table(rows=1, cols=len(strip.items))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _table_borders(table, top=T.RULE, bottom=T.RULE, inside_v=T.RULE_2)
        _fixed_layout(table, [width / len(strip.items)] * len(strip.items))
        for cell, item in zip(table.rows[0].cells, strip.items, strict=True):
            label = cell.paragraphs[0]
            label.paragraph_format.space_before = Pt(5)
            label.paragraph_format.space_after = Pt(0)
            _run(label.add_run(T.clean(item.label)), 8.5, _INK_3)
            value = cell.add_paragraph()
            value.paragraph_format.space_after = Pt(0)
            r = value.add_run(T.clean(item.value))
            _run(r, 22, _INK)
            q = cell.add_paragraph()
            q.paragraph_format.space_after = Pt(5)
            if item.qualifier:
                color = RGBColor.from_string(T.STATUS_HEX.get(item.tone or "", T.INK_3))
                _run(q.add_run(T.clean(item.qualifier)), 8, color)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _table(self, doc: DocumentObject, t: Table) -> None:
        cap = doc.add_paragraph(style="Caption")
        cap.paragraph_format.keep_with_next = True
        _run(cap.add_run(f"{t.label} "), None, _INK_2, bold=True)
        _run(cap.add_run(T.clean(t.caption)), None, _INK_2)

        width = _text_width(doc)
        table = doc.add_table(rows=1, cols=len(t.headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _table_borders(table, bottom=T.RULE, inside_h=T.RULE_2)
        _cell_margins(table, top=45, bottom=45, left=85, right=85)
        _fixed_layout(table, _column_widths(t, width))

        header = table.rows[0]
        _repeat_header(header)
        # Never a header alone at a page foot: it keeps with the first row, and a small
        # table (up to SMALL_TABLE rows) keeps all its rows on one page.
        small = len(t.rows) <= SMALL_TABLE
        for cell, text, align in zip(header.cells, t.headers, t.align, strict=True):
            _shade(cell, T.HEADER_FILL)
            _cell_border(cell, "bottom", T.IRIS, size=8)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == "R" else WD_ALIGN_PARAGRAPH.LEFT
            _run(p.add_run(text), 8.5, _INK_2, bold=True)

        for ri, row in enumerate(t.rows):
            tr = table.add_row()
            _cant_split(tr)
            for cell, value, head, align in zip(tr.cells, row, t.headers, t.align,
                                                strict=True):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if small and ri < len(t.rows) - 1:
                    p.paragraph_format.keep_with_next = True
                p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if align == "R"
                               else WD_ALIGN_PARAGRAPH.LEFT)
                text = cell_text(value, head)
                if isinstance(value, Status):
                    glyph = p.add_run(f"{GLYPH} ")
                    _run(glyph, 8.5, RGBColor.from_string(T.STATUS_HEX[value.tone]))
                _run(_nb_run(p, text), 8.5, _INK_3 if text == NA else _INK)
        self._tables += 1
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    def _figure(self, doc: DocumentObject, f: Figure) -> None:
        width = min(_text_width(doc), int(Inches(f.diagram.width_in)))
        # A picture taller than the page body would be clipped and push its caption to
        # an otherwise empty page: scale it to fit with its caption.
        max_h = _body_height(doc) - int(Pt(48))
        if f.diagram.width and width * f.diagram.height / f.diagram.width > max_h:
            width = int(max_h * f.diagram.width / f.diagram.height)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.add_run().add_picture(io.BytesIO(f.diagram.png), width=Emu(width))
        _alt_text(p, T.clean(f.title))
        self._figure_caption(doc.add_paragraph(style="Caption"), f)

    # ─── architecture diagrams ─────────────────────────────────────
    # Heading, summary, figure and caption are chained with keep-with-next, and part 1
    # of a diagram is laid out to fit under them (content.first_part_budget), so they
    # share a page.

    @staticmethod
    def _eyebrow(p: Any) -> None:
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run("DESIGN SUMMARY")
        _run(r, 7.5, _INK_3, bold=True)
        _letter_spacing(r, 16)

    @staticmethod
    def _summary_items(container: Any, items: list[str]) -> None:
        """Bullets with the iris square marker the PDF and the mocks use."""
        for item in items:
            b = container.add_paragraph()
            b.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = b.paragraph_format
            pf.left_indent = Inches(0.17)
            pf.first_line_indent = -Inches(0.17)
            pf.space_after = Pt(3)
            pf.keep_with_next = True
            pf.tab_stops.add_tab_stop(Inches(0.17))
            marker = _run(b.add_run(GLYPH), 5.5, _IRIS)
            _rfonts(marker._r, "Arial")
            b.add_run("\t")
            text = _run(_nb_run(b, T.clean(item)), 9.5, _INK)
            _rfonts(text._r, T.OFFICE_FONT)

    def _scene_figure(self, doc: DocumentObject, f: SceneFigure) -> None:
        text_w = _text_width(doc)
        width = min(text_w, int(Pt(f.scene.width)))
        png = render_png(f.scene, dpi=220)
        if f.summary_beside:
            # Decision 3: beside the figure when it is 480 pt or narrower.
            if f.new_page:
                doc.add_paragraph().paragraph_format.page_break_before = True
            gap = int(Inches(0.2))
            table = doc.add_table(rows=1, cols=2)
            _fixed_layout(table, [width + gap, max(text_w - width - gap, int(Inches(1)))])
            _cell_margins(table, top=0, bottom=0, left=0, right=0)
            left, right = table.rows[0].cells
            pic = left.paragraphs[0]
            pic.paragraph_format.space_after = Pt(0)
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(io.BytesIO(png), width=Emu(width))
            _alt_text(pic, T.clean(f.title))
            self._eyebrow(right.paragraphs[0])
            self._summary_items(right, f.summary)
        else:
            if f.summary:
                self._eyebrow(doc.add_paragraph())
                if len(f.summary) > 3 and text_w >= int(Inches(8)):
                    half = (len(f.summary) + 1) // 2
                    table = doc.add_table(rows=1, cols=2)
                    # Within the text width, with Word's normal cell margins.
                    _fixed_layout(table, [text_w // 2, text_w - text_w // 2])
                    _cell_margins(table, top=0, bottom=0, left=108, right=108)
                    for cell, items in zip(table.rows[0].cells,
                                           (f.summary[:half], f.summary[half:]), strict=True):
                        cell._tc.remove(cell.paragraphs[0]._p)
                        self._summary_items(cell, items)
                else:
                    self._summary_items(doc, f.summary)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(4)
            # Too tall for the room under the heading and summary: the figure and its
            # caption start the next page (the summary stays with its heading).
            p.paragraph_format.page_break_before = f.new_page
            p.add_run().add_picture(io.BytesIO(png), width=Emu(width))
            _alt_text(p, T.clean(f.title))
        cap = doc.add_paragraph(style="Caption")
        self._figure_caption(cap, f)
        if f.note:
            cap.paragraph_format.space_after = Pt(2)
            cap.paragraph_format.keep_with_next = True
            note = doc.add_paragraph()
            note.paragraph_format.space_after = Pt(10)
            _run(_nb_run(note, T.clean(f.note)), 8, _INK_3)

    @staticmethod
    def _figure_caption(cap: Any, f: Figure | SceneFigure) -> None:
        cap.paragraph_format.space_after = Pt(10)
        _run(_nb_run(cap, T.clean(f.title)), None, _INK_3, italic=True)

    def _figure_row(self, doc: DocumentObject, row: FigureRow) -> None:
        if not row.figures:
            return
        width = _text_width(doc)
        n = len(row.figures)
        table = doc.add_table(rows=1, cols=n)
        _fixed_layout(table, [width / n] * n)
        _cell_margins(table, top=0, bottom=0, left=0, right=110)
        for cell, f in zip(table.rows[0].cells, row.figures, strict=True):
            p = cell.paragraphs[0]
            col = int(width / n) - int(Inches(0.1))
            p.add_run().add_picture(io.BytesIO(f.diagram.png),
                                    width=Emu(min(col, int(Inches(f.diagram.width_in)))))
            _alt_text(p, T.clean(f.title))
            cap = cell.add_paragraph(style="Caption")
            self._figure_caption(cap, f)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ─── running header / footer ───────────────────────────────────

    def _header_footer(self, doc: DocumentObject, report: Report) -> None:
        """The running header and footer on every section; each gets its own copy so
        the right-aligned text sits at that section's margin (landscape sections are
        wider)."""
        for i, section in enumerate(doc.sections):
            if i:
                section.header.is_linked_to_previous = False
                section.footer.is_linked_to_previous = False
                section.different_first_page_header_footer = False
            else:
                section.different_first_page_header_footer = True
            self._section_header_footer(section, report)

    @staticmethod
    def _section_header_footer(section: Any, report: Report) -> None:
        right = Emu(int(section.page_width) - int(section.left_margin or 0)
                    - int(section.right_margin or 0))
        hp = section.header.paragraphs[0]
        hp.text = ""
        _right_tab_only(hp, right)
        _run(hp.add_run("NTNXVision"), 8, _INK, bold=True)
        _run(hp.add_run(f"  ·  {report.running_title}\t{T.clean(report.scope)}"), 8, _INK_3)
        _border(hp._p.get_or_add_pPr(), "bottom", T.RULE_2, size=4, space=4)

        fp = section.footer.paragraphs[0]
        fp.text = ""
        _right_tab_only(fp, right)
        _border(fp._p.get_or_add_pPr(), "top", T.RULE_2, size=4, space=4)
        left = "Confidential"
        if report.prepared_for:
            left += f" · Prepared for {T.clean(report.prepared_for)}"
        _run(fp.add_run(f"{left}\tPage "), 8, _INK_3)
        _simple_field(fp, "PAGE")
        _run(fp.add_run(" of "), 8, _INK_3)
        _simple_field(fp, "NUMPAGES")


# ─── helpers: layout ──────────────────────────────────────────────────────


#: Customer Word template ceilings (a real template is well under a few MB).
MAX_TEMPLATE_BYTES = 25 * 1024**2
MAX_TEMPLATE_UNPACKED = 200 * 1024**2
_TEMPLATE_MAX_MEMBERS = 2000
_TEMPLATE_MAX_RATIO = 200


def _text_width(doc: DocumentObject) -> int:
    s = doc.sections[-1]
    return int(s.page_width or Inches(8.5)) - int(s.left_margin or 0) - int(s.right_margin or 0)


def _column_widths(t: Table, width: int) -> list[float]:
    """Column widths (EMU) from the shared weights: content-sized when the table is
    narrow, spread across the text width otherwise."""
    per_char = Inches(0.07)
    pad = Inches(0.2)
    glyph = Inches(0.13)
    natural, minimum = [], []
    for i, header in enumerate(t.headers):
        texts = [cell_text(r[i], header) for r in t.rows[:250]]
        extra = glyph if any(isinstance(r[i], Status) for r in t.rows[:250]) else 0
        widest = max((len(x) for x in texts), default=0)
        natural.append(min(max(widest, len(header) * 0.9), 60) * per_char + pad + extra)
        # Never narrower than the longest word (a name, a date, a status), so the table
        # wraps between words instead of breaking "CRITICAL" across two lines.
        longest = max((len(w) for x in [*texts, header] for w in x.split(" ")), default=4)
        minimum.append(min(longest, 18) * per_char + pad + extra)
    total = sum(natural)
    if total <= width:
        return natural if total < width * 0.62 else [w * width / total for w in natural]
    floor = sum(minimum)
    if floor >= width:
        return [m * width / floor for m in minimum]
    spare, flex = width - floor, total - floor
    return [m + (n - m) * spare / flex for n, m in zip(natural, minimum, strict=True)]


def _fixed_layout(table: Any, widths: list[float]) -> None:
    tbl = table._tbl
    pr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    pr.append(layout)
    tbl_w = pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) / 635)))
    table.autofit = False
    grid = tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths, strict=False):
        col.set(qn("w:w"), str(int(w / 635)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths, strict=False):
            cell.width = Emu(int(w))


def _table_borders(table: Any, top: str | None = None, bottom: str | None = None,
                   inside_h: str | None = None, inside_v: str | None = None) -> None:
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, color in (("top", top), ("left", None), ("bottom", bottom), ("right", None),
                        ("insideH", inside_h), ("insideV", inside_v)):
        el = OxmlElement(f"w:{edge}")
        if color:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    pr.append(borders)


def _cell_margins(table: Any, top: int, bottom: int, left: int, right: int) -> None:
    pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    pr.append(mar)


def _repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)
    _cant_split(row)


def _cant_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _shade(cell: Any, hex_color: str) -> None:
    """Set a table cell's background fill (docx has no direct API for this)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _cell_border(cell: Any, edge: str, color: str, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    borders.append(el)


# ─── helpers: text ────────────────────────────────────────────────────────


def _rfonts(el: Any, name: str) -> None:
    """Set every rFonts slot (theme fonts in the default template otherwise win)."""
    r_pr = el.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def _font(style: Any, name: str, size: float, color: RGBColor, bold: bool | None = None,
          italic: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    _rfonts(style.element, name)


def _lang(style: Any, tag: str) -> None:
    r_pr = style.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), tag)
    r_pr.append(lang)


#: A hyphenated token this short (``Table 5-2``, ``X-Play``, ``en-pc02``,
#: ``md-pc01.demo.local``) never breaks at its hyphen; longer ones (cluster names) may.
_NB_TOKEN_MAX = 24
_TOKEN = re.compile(r"\S+")
_INNER_HYPHEN = re.compile(r"\w-\w")


def _nb_hyphens(text: str) -> str:
    """`text` with the hyphens of short hyphenated tokens as U+2011 markers."""
    def keep(m: re.Match[str]) -> str:
        tok = m.group(0)
        if len(tok) <= _NB_TOKEN_MAX and _INNER_HYPHEN.search(tok):
            return tok.replace("-", "\u2011")
        return tok
    return _TOKEN.sub(keep, text)


def _nb_run(paragraph: Any, text: str) -> Any:
    """A run with `text`, its short hyphenated tokens joined by the non-breaking hyphen
    character (U+2011, which Calibri and Figtree carry), so "Table 5-2" or "en-pc02"
    never split at a line end. A literal character rather than ``w:noBreakHyphen``:
    Word measured that element too narrow in table cells and squeezed the next words."""
    return paragraph.add_run(_nb_hyphens(text))


def _body_height(doc: DocumentObject) -> int:
    s = doc.sections[-1]
    return (int(s.page_height or Inches(11)) - int(s.top_margin or 0)
            - int(s.bottom_margin or 0))


def _run(run: Any, size: float | None, color: RGBColor, bold: bool = False,
         italic: bool = False) -> Any:
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold:
        run.font.bold = True
    if italic:
        run.font.italic = True
    return run


def _letter_spacing(run: Any, twentieths: int) -> None:
    r_pr = run._r.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(twentieths))
    r_pr.append(sp)


def _border(p_pr: Any, edge: str, color: str, size: int = 4, space: int = 1) -> None:
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    borders.append(el)


def _right_tab_only(paragraph: Any, position: Emu) -> None:
    """One right-aligned tab at the text edge. The Header / Footer styles carry centre and
    right stops for Letter (3.25" / 6.5"); left in place, the centre stop catches the tab
    and runs the right-hand text into the left."""
    stops = paragraph.paragraph_format.tab_stops
    for inches in (3.25, 6.5):
        stops.add_tab_stop(Inches(inches), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(position, WD_TAB_ALIGNMENT.RIGHT)


def _clear_spacing(style_el: Any) -> None:
    """Drop the character spacing (tracking) the default Title / Subtitle styles carry."""
    r_pr = style_el.get_or_add_rPr()
    for el in r_pr.findall(qn("w:spacing")):
        r_pr.remove(el)


def _clear_border(p_pr: Any) -> None:
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def _alt_text(paragraph: Any, text: str) -> None:
    """Alt text (title and description) on the picture just added, for screen readers."""
    for doc_pr in paragraph._p.iter(qn("wp:docPr")):
        doc_pr.set("descr", text)
        doc_pr.set("title", text)


def _field_begin(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    for el in (begin, instr, sep):
        run._r.append(el)


def _field_end(paragraph: Any) -> None:
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _simple_field(paragraph: Any, instruction: str) -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), T.INK_3)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    r_pr.append(color)
    r_pr.append(size)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def _update_fields_on_open(doc: DocumentObject) -> None:
    """Ask Word to refresh fields (the contents' page numbers) when the file opens."""
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def _ensure_styles(doc: DocumentObject) -> None:
    """Styles the renderer uses that python-docx's default template (or a customer
    template) may lack: the contents heading and the TOC levels."""
    existing = {s.name for s in doc.styles}
    if "Contents Heading" not in existing:
        st = doc.styles.add_style("Contents Heading", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        _font(st, T.OFFICE_FONT, 18, _INK, bold=True)
        st.paragraph_format.space_after = Pt(12)
        _border(st.element.get_or_add_pPr(), "left", T.IRIS, size=24, space=8)
    for level, (size, color, indent, bold) in {
        1: (10.5, _INK, 0.0, True), 2: (10, _INK_2, 0.3, False),
    }.items():
        name = f"TOC {level}"
        if name in existing:
            continue
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        _font(st, T.OFFICE_FONT, size, color, bold=bold)
        st.paragraph_format.left_indent = Inches(indent)
        st.paragraph_format.space_before = Pt(5 if level == 1 else 1)
        st.paragraph_format.space_after = Pt(1)
    for name in ("Title", "Subtitle", "Caption"):
        if name not in existing:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH).base_style = \
                doc.styles["Normal"]


_REQUIRED_STYLES = (
    ("Heading 1", "paragraph"), ("Heading 2", "paragraph"), ("Heading 3", "paragraph"),
    ("List Bullet", "paragraph"), ("Caption", "paragraph"),
)


def load_template(template: str) -> DocumentObject:
    """Open a customer ``.dotx`` / ``.docx`` as the base document, body emptied.

    python-docx refuses ``.dotx`` because its main part has the *template* content
    type; the package is otherwise identical to a document, so a copy with that one
    content type rewritten opens normally. The template's body is removed (keeping the
    final section properties: page size, margins, headers and footers), its styles,
    theme, numbering and headers/footers are kept, and any style this exporter relies on
    that the template does not define is added so rendering never fails on a missing
    style.
    """
    import tempfile
    import zipfile

    src = Path(template).expanduser()
    if not src.is_file():
        raise FileNotFoundError(f"Word template not found: {src}")
    if src.suffix.lower() not in (".dotx", ".docx"):
        raise ValueError("A Word template must be a .dotx or .docx file")
    if src.stat().st_size > MAX_TEMPLATE_BYTES:
        raise ValueError(f"The Word template is too large (over "
                         f"{MAX_TEMPLATE_BYTES // 1024**2} MB).")
    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp) / "template.docx"
        try:
            zin_ctx = zipfile.ZipFile(src)
        except zipfile.BadZipFile as exc:
            raise ValueError("The Word template is not a valid .dotx or .docx file") from exc
        with zin_ctx as zin, zipfile.ZipFile(work, "w", zipfile.ZIP_DEFLATED) as zout:
            infos = zin.infolist()
            if len(infos) > _TEMPLATE_MAX_MEMBERS or \
                    sum(i.file_size for i in infos) > MAX_TEMPLATE_UNPACKED:
                raise ValueError("The Word template is too large or has too many parts to open.")
            for item in infos:
                if item.file_size > 1024**2 and item.file_size > _TEMPLATE_MAX_RATIO * max(
                        item.compress_size, 1):
                    raise ValueError("The Word template looks like a compression bomb; "
                                     "not opened.")
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"wordprocessingml.template.main+xml",
                        b"wordprocessingml.document.main+xml")
                zout.writestr(item, data)
        doc = Document(str(work))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    existing = {s.name for s in doc.styles}
    for name, kind in _REQUIRED_STYLES:
        if name not in existing:
            style = doc.styles.add_style(
                name, WD_STYLE_TYPE.PARAGRAPH if kind == "paragraph" else WD_STYLE_TYPE.CHARACTER)
            if name.startswith("Heading"):
                style.font.bold = True
                style.font.size = Pt({"1": 16, "2": 13, "3": 12}[name[-1]])
    return doc


__all__ = ["WordExporter", "load_template"]
