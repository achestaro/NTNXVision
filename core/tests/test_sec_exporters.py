"""Security hardening: Excel formula injection, snapshot string caps, input size limits."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
import requests
from openpyxl import load_workbook

from ntnxvision_core.api import transport
from ntnxvision_core.api.errors import NutanixApiError
from ntnxvision_core.compare import export_comparison
from ntnxvision_core.exporters import ExcelExporter
from ntnxvision_core.models import Cluster, Environment, VirtualMachine
from ntnxvision_core.store.transfer import MAX_TEXT_CHARS, cap_strings

HOSTILE = ["=1+1", "+SUM(A1)", "-2+3", "@SUM(1)", "\t=1", "\r=1", "\n=1",
           "＝1+1", "＋1", "－1", "＠SUM(1)",
           '=HYPERLINK("http://198.51.100.1","x")']


def _env(names: list[str]) -> Environment:
    env = Environment(collected_at="2026-06-30T00:00:00Z")
    env.clusters.append(Cluster(uuid="c1", name="=cluster", aos_version="7.0",
                                hypervisor_types=["AHV"], num_nodes=3))
    for i, n in enumerate(names):
        env.vms.append(VirtualMachine(uuid=f"v{i}", name=n, description=n, cluster_name="=c",
                                      categories={"=k": n}, power_state="ON", num_vcpus=2))
    env.collection_notes["=x"] = "=cmd|' /C calc'!A0"
    return env


def _no_formulas(path: Path) -> int:
    wb = load_workbook(path)
    strings = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
                strings += isinstance(c.value, str)
    return strings


def test_inventory_workbook_has_no_formula_cells(tmp_path: Path) -> None:
    out = ExcelExporter().export(_env(HOSTILE), tmp_path / "a.xlsx")
    assert _no_formulas(out) > 0
    wb = load_workbook(out)
    values = {c.value for ws in wb.worksheets for r in ws.iter_rows() for c in r}
    assert "=1+1" in values  # kept verbatim, as text


def test_compare_export_has_no_formula_cells(tmp_path: Path) -> None:
    out = export_comparison(_env(["ok"]), _env(HOSTILE), "=base", "@cur", "excel", tmp_path)
    assert _no_formulas(out) > 0


def test_numbers_stay_numeric(tmp_path: Path) -> None:
    out = ExcelExporter().export(_env(["vm"]), tmp_path / "n.xlsx")
    wb = load_workbook(out)
    assert any(c.data_type == "n" for ws in wb.worksheets for r in ws.iter_rows() for c in r)


def test_cap_strings_truncates_only_long_text() -> None:
    long = "x" * (MAX_TEXT_CHARS + 500)
    out = cap_strings({"a": [long, "short"], "n": 5})
    assert len(out["a"][0]) == MAX_TEXT_CHARS and out["a"][0].endswith("…")
    assert out["a"][1] == "short" and out["n"] == 5


def test_word_template_size_limits(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from ntnxvision_core.exporters import word

    big = tmp_path / "t.dotx"
    big.write_bytes(b"x")
    with monkeypatch.context() as m:
        m.setattr(word, "MAX_TEMPLATE_BYTES", 0)
        with pytest.raises(ValueError, match="too large"):
            word.load_template(str(big))
    bomb = tmp_path / "bomb.dotx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (50 * 1024**2))
    with pytest.raises(ValueError, match="compression bomb|too large"):
        word.load_template(str(bomb))


def test_http_response_cap() -> None:
    class Adapter(requests.adapters.BaseAdapter):
        def send(self, request: Any, **kw: Any) -> requests.Response:
            r = requests.Response()
            r.status_code = 200
            r.url = request.url
            r.raw = io.BytesIO(b"[" + b"1," * 600 + b"1]")
            return r

        def close(self) -> None:
            pass

    s = transport.new_session()
    s.mount("https://", Adapter())
    assert len(s.get("https://example.lab/a").json()) == 601
    s.max_bytes = 100  # type: ignore[attr-defined]
    with pytest.raises(NutanixApiError, match="limit"):
        s.get("https://example.lab/a")
