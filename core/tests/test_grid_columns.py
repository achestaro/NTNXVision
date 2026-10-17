"""Where cluster figures live: grid, detail pane, Clusters tables and the capacity tables.

* `Col(grid=False)` keeps a column in the detail pane and the reports but out of the
  Inventory grid (and its Columns menu): the cluster's Data Services IP.
* The capacity and oversubscription figures are not Clusters columns at all: the Excel
  Capacity sheet, the PDF/Word "Capacity & Oversubscription" section and the cluster
  detail pane carry them. These tests make sure nothing was lost on the way.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.engine import collect_environment
from ntnxvision_core.exporters import ExportOptions, export_environment
from ntnxvision_core.exporters import tables as T
from ntnxvision_core.exporters.content import Table, build_report
from ntnxvision_core.inventory import InventoryService

#: Columns the Clusters grid leaves out (detail pane only).
# The vCenter an ESXi cluster is registered to is a detail too (empty on AHV).
DETAIL_ONLY_KEYS = {"data_services_ip", "vcenter"}
#: The capacity columns the Excel Clusters sheet used to carry, by header.
MOVED_TO_CAPACITY = {
    "vCPU:pCore (all VMs)", "vCPU:pCore N-1", "vCPU Total", "vCPU Total N-1",
    "Storage Provisioned:Usable", "Storage Provisioned:Usable N-1", "Storage Total (GiB)",
    "Storage Total N-1 (GiB)",
}
#: Any header naming a capacity figure (none may appear in a Clusters table).
CAPACITY_WORDS = ("vCPU", "pCore", "Storage Total", "Provisioned", "Usable", "Oversub",
                  "Raw", "Cores")


@pytest.fixture
def env(demo: Any) -> Any:
    return collect_environment(PrismCredentials(host="pc01.demo.local", username="admin",
                                                password="demo"))


def _tables(report: Any) -> list[Table]:
    return [b for ch in report.chapters for blocks in
            [ch.blocks, *[s.blocks for s in ch.sections]] for b in blocks
            if isinstance(b, Table)]


# ─── the flag ──────────────────────────────────────────────────────────


def test_flag_defaults_to_the_grid() -> None:
    getter: Any = lambda o, _c: None  # noqa: E731
    assert T.Col("a", "A", getter, label="A").grid is True


def test_grid_false_drops_a_column_from_ui_cols_only() -> None:
    getter: Any = lambda o, _c: None  # noqa: E731
    spec = T.TableSpec(
        "x", "X", "X", "X", "x", lambda e: [],
        (T.Col("cluster", "Cluster", getter, label="Cluster"),
         T.Col("shown", "Shown", getter, label="Shown"),
         T.Col("detail", "Detail", getter, label="Detail", grid=False)),
        id=lambda o: "", name=lambda o: "")
    assert [c.key for c in spec.ui_cols] == ["cluster", "shown"]
    assert [c.key for c in spec.excel_cols] == ["cluster", "shown", "detail"]


# ─── the Clusters spec (grid + Excel sheet) ────────────────────────────


def test_clusters_spec() -> None:
    grid = {c.key for c in T.CLUSTERS.ui_cols}
    headers = [c.header for c in T.CLUSTERS.excel_cols]
    assert {c.key for c in T.CLUSTERS.cols if not c.grid} == DETAIL_ONLY_KEYS
    assert not grid & DETAIL_ONLY_KEYS
    assert "Data Services IP" in headers
    assert not set(headers) & MOVED_TO_CAPACITY
    assert not [h for h in headers if h and any(w in h for w in CAPACITY_WORDS)]
    # Cluster stays the first column of the grid and the sheet.
    assert T.CLUSTERS.ui_cols[0].label == "Cluster" and headers[0] == "Cluster"


def test_every_moved_figure_is_on_the_capacity_sheet_spec() -> None:
    capacity = {c.header for c in T.CAPACITY.excel_cols}
    assert capacity >= MOVED_TO_CAPACITY


# ─── the app ───────────────────────────────────────────────────────────


def test_grid_and_detail_pane(env: Any) -> None:
    svc = InventoryService()
    res = svc.query(env, "clusters", page={"limit": 50})
    keys = {c["key"] for c in res["columns"]}
    assert not keys & DETAIL_ONLY_KEYS
    assert not keys & {"cpu_ratio", "cpu_ratio_n1", "storage_ratio", "storage_ratio_n1"}
    # Still searchable: a Data Services IP finds its cluster.
    dsip = next(c.data_services_ip for c in env.clusters if c.data_services_ip)
    assert svc.query(env, "clusters", filter_={"text": dsip}, page={"limit": 5})["matching"]

    cluster = next(c for c in env.clusters if c.data_services_ip)
    row = next(r for r in res["rows"] if r["cells"]["name"] == cluster.name)
    fields = {f["label"]: f["value"] for f in svc.detail(env, "clusters", row["id"])["fields"]}
    labels = list(fields)
    assert fields["Data Services IP"] == cluster.data_services_ip
    assert labels.index("Data Services IP") == labels.index("External IP") + 1
    # Every capacity figure, N-1 after a "·".
    for label in ("vCPU Total", "vCPU:pCore", "Storage ratio", "Target ratio",
                  "vCPU allocated", "Raw storage", "Provisioned disk", "Usable cores"):
        assert label in fields, label
    for label in ("vCPU Total", "vCPU:pCore", "Storage ratio"):
        assert "· N-1 " in str(fields[label]), label
    assert any(lbl.startswith("Storage Total") and "· N-1 " in str(v)
               for lbl, v in fields.items())


# ─── the reports ───────────────────────────────────────────────────────


def test_excel_clusters_sheet_and_capacity_sheet(env: Any, tmp_path: Path) -> None:
    from ntnxvision_core.exporters.tables import Ctx

    wb = load_workbook(export_environment(env, "excel", tmp_path, "x"))
    clusters = [c.value for c in wb["Clusters"][1]]
    assert "Data Services IP" in clusters
    assert not set(clusters) & MOVED_TO_CAPACITY
    sheet = wb["Capacity"]
    head = [c.value for c in sheet[1]]
    assert set(head) >= MOVED_TO_CAPACITY
    rows = {r[0]: dict(zip(head, r, strict=False))
            for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] and r[1] is not None}
    # The figures are the ones the capacity math produces for each cluster.
    ctx = Ctx(env)
    for c in env.clusters:
        cap = ctx.capacity(c)
        assert cap is not None
        row = rows[c.name]
        assert row["VMs"] == cap.vm_count
        expected = "N/A" if cap.vcpu_total is None else int(cap.vcpu_total)
        assert row["vCPU Total"] == expected


def test_pdf_word_clusters_table_and_capacity_section(env: Any) -> None:
    report = build_report(env, ExportOptions(), eyebrow="x")
    tables = _tables(report)
    overview = next(t for t in tables if t.caption.startswith("Registered clusters"))
    assert "Data Services IP" in overview.headers
    assert not [h for h in overview.headers if any(w in h for w in CAPACITY_WORDS)]
    cpu = next(t for t in tables if t.caption.startswith("CPU capacity"))
    storage = next(t for t in tables if t.caption.startswith("Storage capacity"))
    assert {"vCPU Total", "vCPU Total N-1", "vCPU Allocated", "vCPU:pCore",
            "vCPU:pCore N-1"} <= set(cpu.headers)
    assert {"Storage Total (TiB)", "Storage Total N-1 (TiB)", "Provisioned (TiB)",
            "Provisioned:Usable", "Provisioned:Usable N-1"} <= set(storage.headers)
    assert len(cpu.rows) == len(storage.rows) == len(env.clusters)
