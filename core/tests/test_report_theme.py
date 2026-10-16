"""The three report exports agree with each other and follow the report rules: one shared
outline for PDF and Word, N/A for missing values, no em dashes, comma grouping, the app's
collection wording, correct combined and Prism Element handling, and the light Runway
theme's structural pieces (document properties, contents, heading styles, status cells).
"""

from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from ntnxvision_core.analysis import analyze_feature_usage
from ntnxvision_core.engine import combine_environments
from ntnxvision_core.exporters import ExportOptions, export_environment
from ntnxvision_core.exporters import theme as TH
from ntnxvision_core.exporters.content import Status, Table, build_report
from ntnxvision_core.exporters.excel import ExcelExporter
from ntnxvision_core.exporters.pdf import PdfExporter
from ntnxvision_core.exporters.word import GLYPH, WordExporter
from ntnxvision_core.models import (
    Alert,
    Cluster,
    Environment,
    Host,
    LicenseInfo,
    PrismCentral,
    StorageContainer,
    Subnet,
    VirtualMachine,
)
from ntnxvision_core.runinfo import api_routing, sunset_areas

EM = chr(0x2014)
GIB = 1024**3


def _pc_env(host: str = "pc.lab.local", version: str = "pc.2024.3") -> Environment:
    env = Environment(collected_at="2026-09-05T08:30:00Z",
                      prism_central=PrismCentral(host=host, version=version))
    env.clusters.append(Cluster(uuid=f"c-{host}", name="cluster-01", aos_version="7.0",
                                hypervisor_types=["AHV"], num_nodes=3, redundancy_factor=2))
    env.hosts.append(Host(uuid=f"h-{host}", name="node-a", cluster_uuid=f"c-{host}",
                          cluster_name="cluster-01", num_cpu_cores=32,
                          memory_bytes=512 * GIB, num_vms=1500))
    env.vms += [VirtualMachine(uuid=f"v{i}-{host}", name=f"vm-{i}", cluster_uuid=f"c-{host}",
                               cluster_name="cluster-01", power_state="ON" if i % 2 else "on",
                               num_vcpus=2, memory_bytes=4 * GIB) for i in range(12)]
    env.storage_containers.append(StorageContainer(
        uuid=f"s-{host}", name="default", cluster_uuid=f"c-{host}", replication_factor=2,
        compression_enabled=True, erasure_coding="off", deduplication="NONE",
        max_capacity_bytes=12_000 * GIB, used_bytes=3_000 * GIB))
    env.licenses.append(LicenseInfo(product="NCI", edition="Pro", quantity=12000,
                                    expiry="2027-11-05T00:00:00Z", compliance="EXPIRED"))
    env.alerts.append(Alert(uuid=f"a-{host}", title="Disk failure", severity="CRITICAL",
                            resolved=False, cluster_name="cluster-01",
                            created_at="2026-09-01T10:15:00Z"))
    env.collection_notes = {
        "clusters": "ok", "hosts": "ok", "vms": f"ok: 3 VM(s) skipped {EM} internal error",
        "storage_containers": "ok", "licenses": "ok", "alerts": "ok",
        "files": "skipped: Nutanix Files not deployed",
    }
    env.warnings.append(f"Prism Central {version} {EM} outside the tested range")
    return env


def _pe_env() -> Environment:
    env = Environment(collected_at="2026-09-05T08:30:00Z", source_type="prism_element",
                      prism_central=PrismCentral(host="pe.lab.local", version="7.0"))
    env.clusters.append(Cluster(uuid="c1", name="pe-cluster", num_nodes=1))
    env.storage_containers.append(StorageContainer(uuid="s1", name="ctr", cluster_uuid="c1",
                                                   max_capacity_bytes=100 * GIB))
    env.subnets.append(Subnet(uuid="n1", name="vlan-10", vlan_id=10))  # no cluster ref
    env.collection_notes = {k: "ok" for k in
                            ("clusters", "hosts", "storage_containers", "subnets", "vms")}
    return env


# ─── shared outline ────────────────────────────────────────────────────


def test_pdf_and_word_render_the_same_outline(tmp_path: Path) -> None:
    env = _pc_env()
    report = build_report(env, ExportOptions(), "x")
    headings = [h for level, h in report.outline() if level == 0]
    assert headings[0] == "1. Executive Summary"
    assert headings[-1] == "Appendix A. Collection Coverage & API Routing"
    numbers = [h.split(".", 1)[0] for h in headings if h[0].isdigit()]
    assert numbers == [str(i) for i in range(1, len(numbers) + 1)]  # no gaps

    doc = Document(str(WordExporter().export(env, tmp_path / "a.docx")))
    word_h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert word_h1 == headings
    word_h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert word_h2 == [h for level, h in report.outline() if level == 1]


def test_sections_filter_both_documents_identically() -> None:
    env = _pc_env()
    opts = ExportOptions(sections=frozenset({"licensing_lcm", "collection_notes"}))
    titles = [c.title for c in build_report(env, opts, "x").chapters]
    assert titles == ["Licensing & Feature Usage", "Assumptions & Constraints",
                      "Collection Coverage & API Routing"]


def test_pdf_contents_fit_their_reserved_pages(tmp_path: Path) -> None:
    """The contents pages are reserved up front; an overflow would shift every page
    number printed after it (fpdf2 raises instead of silently renumbering)."""
    envs = [_pc_env(f"pc{i}.lab.local") for i in range(3)]
    combined = combine_environments(envs)
    pdf = PdfExporter()
    pdf.export(combined, tmp_path / "c.pdf")
    assert pdf.page_count > 5


# ─── values ────────────────────────────────────────────────────────────


def test_missing_values_read_na_and_numbers_are_grouped() -> None:
    report = build_report(_pc_env(), ExportOptions(), "x")
    tables = [b for ch in report.chapters for s in ch.sections for b in s.blocks
              if isinstance(b, Table)]
    hosts = next(t for t in tables if t.headers[0] == "Host")
    assert hosts.align[hosts.headers.index("VMs")] == "R"
    from ntnxvision_core.exporters.content import cell_text

    assert cell_text(1500, "VMs") == "1,500"
    assert cell_text(None, "VMs") == "N/A"
    assert cell_text(512.0, "Memory (GiB)") == "512.0"
    alerts = next(t for t in tables if t.headers[0] == "Severity")
    sev = alerts.rows[0][0]
    assert isinstance(sev, Status) and sev.tone == "error"
    assert alerts.rows[0][alerts.headers.index("Created (UTC)")] == "2026-09-01 10:15"


def test_no_em_dash_reaches_any_export(tmp_path: Path) -> None:
    env = _pc_env()
    report = build_report(env, ExportOptions(prepared_for="Acme", author="Ops"), "x")
    blob = repr(report)
    assert EM not in blob
    for fmt in ("excel", "word"):
        path = export_environment(env, fmt, tmp_path, "r")
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if name.endswith(".xml"):
                    assert EM not in z.read(name).decode("utf-8"), f"{fmt}:{name}"
    assert TH.clean(f"a {EM} b") == "a, b"


def test_erasure_coding_off_is_not_counted_as_in_use() -> None:
    rows = {r.feature: r for r in analyze_feature_usage(_pc_env())}
    assert rows["Erasure Coding (EC-X)"].status == "No"


def test_prism_element_features_read_not_collected_rather_than_no() -> None:
    rows = {r.feature: r for r in analyze_feature_usage(_pe_env())}
    assert rows["Async DR replication"].status == "Not collected"
    assert "Prism Element" in rows["Async DR replication"].detail
    assert rows["Redundancy Factor 3"].status == "No"  # clusters were collected


# ─── combined and Prism Element ────────────────────────────────────────


def test_combined_routing_judges_each_pc_against_its_own_version() -> None:
    new, old = _pc_env("new.lab", "pc.7.3"), _pc_env("old.lab", "pc.2024.1")
    old.prism_central.negotiated_api_versions = {"clusters": "legacy: v4 not GA on this PC"}
    combined = combine_environments([new, old])
    areas = [
        {"id": "new.lab: clusters", "label": "new.lab · Clusters", "status": "collected",
         "api": "v4"},
        {"id": "old.lab: clusters", "label": "old.lab · Clusters", "status": "collected",
         "api": "v3"},
        {"id": "old.lab: xplay", "label": "old.lab · X-Play playbooks",
         "status": "collected", "api": "v3"},
    ]
    rows = api_routing(combined, areas)
    by = {(r["source"], r["api"], r["reason"]): r for r in rows}
    old_clusters = by[("old.lab", "v3", "v4 not GA on pc.2024.1")]
    assert old_clusters["areas"] == "old.lab: Clusters"
    assert old_clusters["after_sunset"]["label"] == "Moves to v4 after a PC upgrade"
    assert by[("new.lab", "v4", "v4 SDK served by this Prism Central")]
    assert sunset_areas(rows) == ["X-Play playbooks"]


def test_prism_element_report_has_no_prism_central_chapter(tmp_path: Path) -> None:
    env = _pe_env()
    report = build_report(env, ExportOptions(), "x")
    titles = [c.title for c in report.chapters]
    assert not any("Prism Central" in t for t in titles)
    assert "Prism Element pe.lab.local (AOS 7.0)" in repr(report.chapters[0].blocks[0])
    # The single PE cluster owns the networks and containers its v2.0 payloads leave bare.
    nets = next(b for c in report.chapters for s in c.sections for b in s.blocks
                if isinstance(b, Table) and b.headers[0] == "Subnet")
    assert nets.rows[0][1] == "pe-cluster"
    wb = load_workbook(ExcelExporter().export(env, tmp_path / "pe.xlsx"))
    assert "Prism Central" not in wb.sheetnames
    overview = {r[0]: r[1] for r in wb["Overview"].iter_rows(values_only=True) if r[0]}
    assert overview["Prism Element"] == "pe.lab.local"
    assert wb["Storage"]["B2"].value == "pe-cluster"


# ─── Excel ─────────────────────────────────────────────────────────────


def test_excel_overview_notes_routing_and_values(tmp_path: Path) -> None:
    env = _pc_env()
    env.run = {"areas": [
        {"id": "clusters", "label": "Clusters", "status": "collected", "api": "v4"},
        {"id": "vms", "label": "Virtual machines", "status": "notes", "api": "v4",
         "detail": f"3 VM(s) skipped {EM} internal error"},
        {"id": "xplay", "label": "X-Play playbooks", "status": "collected", "api": "v3"},
    ]}
    out = ExcelExporter().export(env, tmp_path / "w.xlsx",
                                 {"prepared_for": "Acme", "author": "Ops"})
    wb = load_workbook(out)
    assert wb.properties.title == "Nutanix Inventory Workbook: pc.lab.local"
    assert wb.properties.creator == "Ops"
    assert "API Routing" in wb.sheetnames
    ov = wb["Overview"]
    links = {c.value: c.hyperlink.location for row in ov.iter_rows() for c in row
             if c.hyperlink}
    assert links["Clusters"] == "'Clusters'!A1"
    labels = {r[0]: r[1] for r in ov.iter_rows(values_only=True) if r[0]}
    assert labels["Prepared For"] == "Acme"

    notes = list(wb["Collection Notes"].iter_rows(values_only=True))
    assert notes[0] == ("Area", "Outcome", "API", "Detail")
    vm_row = next(r for r in notes if r[0] == "Virtual machines")
    # A note is a successful collection with a caveat, not an error.
    assert vm_row[1:3] == ("Collected with notes", "v4")
    assert EM not in vm_row[3]
    routing = list(wb["API Routing"].iter_rows(values_only=True))
    assert routing[0] == ("Areas", "API", "Reason", "After Q4 2026 v3 sunset")
    assert ("X-Play playbooks", "v3", "No v4 equivalent yet", "Stops collecting") in routing


def test_excel_without_run_metadata_reads_ok_notes_as_collected(tmp_path: Path) -> None:
    """Older snapshots have only collection notes; an "ok: ..." note used to land in the
    Collection Notes sheet as "Error" with the raw note text."""
    wb = load_workbook(ExcelExporter().export(_pc_env(), tmp_path / "old.xlsx"))
    notes = {r[0]: r for r in wb["Collection Notes"].iter_rows(values_only=True)}
    assert notes["Virtual machines"][1] == "Collected with notes"
    assert notes["Nutanix Files"][1] == "Skipped"

    lic = wb["Licensing"]
    headers = [c.value for c in lic[1]]
    expiry = lic.cell(row=2, column=headers.index("Expiry") + 1)
    assert expiry.value == datetime(2027, 11, 5) and expiry.number_format == "yyyy-mm-dd"
    qty = lic.cell(row=2, column=headers.index("Quantity") + 1)
    assert qty.number_format == "#,##0"
    comp = lic.cell(row=2, column=headers.index("Compliance") + 1)
    assert comp.font.color.rgb.endswith(TH.ERROR)
    vms = wb["VMs"]
    power = [c.value for c in vms["B"][1:]]
    assert set(power) == {"On"}  # ON / on normalised
    assert vms.freeze_panes == "A2" and vms.auto_filter.ref


# ─── Word ──────────────────────────────────────────────────────────────


def test_word_structure_contents_properties_and_status(tmp_path: Path) -> None:
    out = WordExporter().export(_pc_env(), tmp_path / "a.docx",
                                {"prepared_for": "Acme Corp", "author": "Ops"})
    doc = Document(str(out))
    props = doc.core_properties
    assert props.title.startswith("Architecture Design Document")
    assert props.author == "Ops" and props.category == "Architecture Design Document"
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode()
        settings = z.read("word/settings.xml").decode()
    assert 'TOC \\o "1-2" \\h \\z \\u' in body
    assert "updateFields" in settings
    assert any(p.text.startswith("Prepared for\tAcme Corp") for p in doc.paragraphs)
    header = doc.sections[0].header.paragraphs[0].text
    assert header.startswith("NTNXVision") and "pc.lab.local" in header
    assert doc.sections[0].different_first_page_header_footer
    lic = next(t for t in doc.tables if t.rows[0].cells[0].text == "Product")
    assert lic.rows[1].cells[-1].text == f"{GLYPH} EXPIRED"
    assert lic.rows[1].cells[3].text == "12,000"
