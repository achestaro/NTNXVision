"""Host Network Configuration (A4): facts, layout, determinism, golden scenes, reports.

Estates come from the session fixture `design_envs` (conftest): the scenario estates and
demo PCs collected through the mock API with the real collectors.
"""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pytest
from fpdf import FPDF

from ntnxvision_core.diagrams import TYPES, check_scene, model_envelope, model_from_envelope
from ntnxvision_core.diagrams import network as N
from ntnxvision_core.diagrams import scene as S
from ntnxvision_core.diagrams.layout.conceptual import Page
from ntnxvision_core.diagrams.layout.network import layout
from ntnxvision_core.diagrams.render.pdf import draw_scene
from ntnxvision_core.diagrams.render.svg import render_svg
from ntnxvision_core.models import Environment

from .diagram_helpers import shuffled as _shuffled

GOLDEN = Path(__file__).parent / "golden" / "diagrams"
UPDATE = os.environ.get("NTNXVISION_UPDATE_GOLDEN") == "1"
EM_DASH = chr(0x2014)
ESTATES = ["small", "medium", "enterprise", "xlarge"]


def _scenes(env: Environment, page: Page | None = None) -> list[S.Scene]:
    return layout(N.network_model(env), page)


def _texts(sc: S.Scene) -> list[str]:
    return [t.text for n in sc.nodes for t in n.lines] + \
        [t.text for e in sc.edges for lb in e.labels for t in lb.lines]


def _pdf_bytes(sc: S.Scene) -> bytes:
    pdf = FPDF(orientation="L", unit="pt", format=(sc.height + 20, sc.width + 20))
    pdf.set_creation_date(datetime(2026, 10, 1, tzinfo=timezone.utc))
    pdf.add_page()
    draw_scene(pdf, sc, 10, 10)
    return bytes(pdf.output())




# ─── facts ─────────────────────────────────────────────────────────────


def test_medium_profiles_bonds_lacp_mtu_and_lldp(design_envs: dict[str, Any]) -> None:
    m = N.network_model(design_envs["medium"])
    assert [p.letter for p in m.profiles] == ["A", "B"]
    a = m.profiles[0]
    assert a.hosts == 16 and a.clusters == ["PROD-E01", "PROD-W01"]
    vs0, vs1 = a.switches
    assert (vs0.name, vs0.default, vs0.bond_mode, vs0.lacp, vs0.mtu) == \
        ("vs0", True, "ACTIVE_BACKUP", False, 1500)
    assert (vs1.name, vs1.bond_mode, vs1.lacp, vs1.mtu) == ("vs1-storage", "BALANCE_TCP",
                                                           True, 9000)
    assert [u.name for u in vs0.uplinks] == ["eth0", "eth1"]
    assert N.uplinks_text(vs0) == "2 x 25 GbE: eth0, eth1"
    assert N.bond_text(vs1.bond_mode, vs1.lacp) == "Balance-tcp with LACP"
    assert [t.label for t in vs0.traffic] == ["AHV host", "CVM", "Guest VMs"]
    assert not vs1.traffic  # no subnet on the storage switch: nothing invented
    assert [nb.key for nb in a.neighbours] == ["A", "B"]
    assert a.neighbours[0].model == "Arista DCS-7050SX3-48YC8"
    assert a.neighbours[0].seen_on == ["eth0", "eth2"]
    assert [u.switch for u in vs0.uplinks] == ["A", "B"]
    # ESXi: stated, never drawn.
    assert m.esxi is not None and m.esxi.clusters == ["ESX-W01", "VDI-E01"]
    lines = N.esxi_lines(m.esxi)
    assert lines[0].startswith("ESXi hosts (ESX-W01 and VDI-E01, 7 hosts)")
    assert "vmnic0 to vmnic3 per host (2 x 10 GbE, 2 x 25 GbE link speed)" in lines[1]
    assert lines[2].startswith("Source of truth for this network: vCenter ")


def test_no_lldp_estate_shows_no_neighbour(design_envs: dict[str, Any]) -> None:
    m = N.network_model(design_envs["large3"])
    assert all(not p.neighbours for p in m.profiles)
    texts = [t for sc in layout(m) for t in _texts(sc)]
    assert "LLDP NEIGHBOUR" not in texts
    assert not any(t.startswith("LLDP: switch") for t in texts)
    assert "No LLDP neighbour is reported on profile A's uplinks." in N.bullets(m)
    # EDGE-A has no vCenter address: the note says nothing about one.
    assert m.esxi is not None and not any("vCenter" in ln for ln in N.esxi_lines(m.esxi))


def test_text_helpers() -> None:
    assert N.switch_model("Arista Networks EOS version 4.31.2F running on an Arista "
                          "DCS-7050SX3-48YC8") == "Arista DCS-7050SX3-48YC8"
    assert N.switch_model("Cisco NX-OS") == "Cisco NX-OS"
    assert N.nic_short("Mellanox Technologies MT2894 Family [ConnectX-6 Lx]") == \
        "Mellanox ConnectX-6 Lx"
    assert N.nic_short("Intel Corporation Ethernet Controller X710 for 10GbE SFP+") == \
        "Intel X710"
    assert N.speed_text(25000) == "25 GbE" and N.speed_text(2500) == "2.5 GbE"
    assert N.speed_text(100) == "100 Mb/s" and N.speed_text(None) is None
    assert N.mtu_text(9000) == "MTU 9,000" and N.mtu_text(None) == "MTU N/A"
    assert N.bond_text(None, False) == "Bond mode N/A"
    assert N.name_list(["A", "B", "C", "D", "E"]) == "A, B, C and 2 more clusters"


def test_applies_only_when_host_networking_was_collected(design_envs: dict[str, Any]
                                                         ) -> None:
    t = TYPES["network"]
    assert (t.chapter, t.section, t.slug) == ("physical_design", "Host Network Configuration",
                                              "host-network-profile")
    assert t.applies(design_envs["medium"])
    assert not t.applies(Environment())
    env = design_envs["medium"].model_copy(deep=True)
    env.virtual_switches = []
    assert not t.applies(env)


# ─── determinism and documents ─────────────────────────────────────────


@pytest.mark.parametrize("name", ESTATES)
def test_render_twice_identical(design_envs: dict[str, Any], name: str) -> None:
    a, b = _scenes(design_envs[name]), _scenes(design_envs[name])
    assert [S.scene_json(x) for x in a] == [S.scene_json(x) for x in b]
    assert [render_svg(x) for x in a] == [render_svg(x) for x in b]
    assert [_pdf_bytes(x) for x in a] == [_pdf_bytes(x) for x in b]


@pytest.mark.parametrize("name", ["medium", "enterprise", "multi3"])
def test_shuffled_api_order_gives_identical_scenes(design_envs: dict[str, Any],
                                                   name: str) -> None:
    base = [S.scene_json(x) for x in _scenes(design_envs[name])]
    for seed in (1, 2, 3):
        assert [S.scene_json(x) for x in _scenes(_shuffled(design_envs[name], seed))] == base


def test_model_envelope_round_trip(design_envs: dict[str, Any]) -> None:
    m = N.network_model(design_envs["enterprise"])
    doc = json.loads(json.dumps(model_envelope("network", m)))
    type_id, again = model_from_envelope(doc)
    assert type_id == "network" and again == m
    assert [S.scene_json(x) for x in layout(again)] == [S.scene_json(x) for x in layout(m)]


def _pretty(sc: S.Scene) -> bytes:
    return (json.dumps(S.to_doc(sc), sort_keys=True, indent=1, ensure_ascii=False)
            + "\n").encode()


@pytest.mark.parametrize("name", ["small", "medium"])
def test_golden_scene_and_svg(design_envs: dict[str, Any], name: str) -> None:
    scenes = _scenes(design_envs[name])
    for suffix, got in (("scene.json", b"".join(_pretty(x) for x in scenes)),
                        ("svg", b"".join(render_svg(x) for x in scenes))):
        path = GOLDEN / f"network-{name}.{suffix}"
        if UPDATE or not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(got)
        assert got == path.read_bytes(), (
            f"{path.name} changed; rerun with NTNXVISION_UPDATE_GOLDEN=1 after reviewing")


# ─── overflow and scaling ──────────────────────────────────────────────


@pytest.mark.parametrize("name", [*ESTATES, "multi3", "large1", "large3", "pc02"])
def test_no_text_exceeds_its_box(design_envs: dict[str, Any], name: str) -> None:
    page = Page()
    scenes = _scenes(design_envs[name], page)
    assert scenes
    for sc in scenes:
        assert check_scene(sc) == [], sc.id
        assert sc.width <= page.width and sc.height <= page.max_h
        assert all(EM_DASH not in t for t in _texts(sc))
        assert sc.meta["group_name"].startswith("Profile ") and sc.type == "network"
    assert len({sc.meta["group"] for sc in scenes}) <= N.MAX_DRAWN


def _stress_model() -> N.NetworkModel:
    switches = []
    for i in range(7):
        vid = f"profile:stress:vs:{i}"
        ups = [N.Uplink(id=f"{vid}:nic:{k}", name=f"ens{i}f{k}np{k}", speed_mbps=100000,
                        up=k % 3 != 0, switch="AB"[k % 2] if i % 2 else None)
               for k in range(8)]
        traffic = [N.Traffic(f"{vid}:t:{j}", label, sub) for j, (label, sub) in enumerate(
            [("AHV host", "management"), ("CVM", "external"),
             ("Guest VMs", "1,234 subnets per cluster")][: 3 - i % 3])]
        switches.append(N.VSwitch(
            id=vid, name=f"vs{i}-with-an-unusually-long-virtual-switch-name", default=i == 0,
            bridge=f"br{i}-long-bridge", bond=f"br{i}-up-with-a-long-bond-name",
            bond_mode="BALANCE_TCP", lacp=True, mtu=9216, uplinks=ups, traffic=traffic))
    nbs = [N.Neighbour(id=f"profile:stress:lldp:{k}", key=k,
                       model="Vendor Extremely-Long-Switch-Model-Name-48x100G-8x400G",
                       chassis="00:00:5E:00:53:FF", port=None,
                       seen_on=[f"ens{i}f0np0" for i in range(7)]) for k in "AB"]
    p = N.Profile(id="profile:stress", letter="A", hosts=1234,
                  clusters=[f"CLUSTER-{i:02d}-WITH-A-LONG-NAME" for i in range(30)],
                  switches=switches, neighbours=nbs, host="stress-host-n1",
                  nic_model="Vendor Very Long NIC Family Name [SuperNIC-9000 Dual Port 100GbE]",
                  nic_driver="supernic_core", nic_firmware="99.99.9999")
    esxi = N.EsxiNote(clusters=[f"ESX-CLUSTER-{i}" for i in range(12)], hosts=96,
                      nics=[f"vmnic{i}" for i in range(8)], speeds=[10000] * 4 + [25000] * 4,
                      port_groups=1234, vcenters=["vcenter.example.lab"])
    return N.NetworkModel(profiles=[p], esxi=esxi, ahv_clusters=30)


def test_stress_model_never_overflows_and_splits() -> None:
    scenes = layout(_stress_model())
    assert len(scenes) > 1 and scenes[0].part == (1, len(scenes))
    for sc in scenes:
        assert check_scene(sc) == [], sc.id
        assert sc.width == Page().width and sc.height <= Page().max_h
        assert min(t.size for n in sc.nodes for t in n.lines) >= 6.0
    names = [n for sc in scenes for n in sc.meta["switches"]]
    assert len(names) == 7  # every switch drawn once
    assert sum("NOT REPORTED BY NUTANIX APIS" in _texts(sc) for sc in scenes) == 1


# ─── reports ───────────────────────────────────────────────────────────


@pytest.fixture
def only_network(monkeypatch: pytest.MonkeyPatch) -> None:
    """The report with this type alone (keeps the test about A4)."""
    for k in list(TYPES):
        if k != "network":
            monkeypatch.delitem(TYPES, k)


def test_word_and_pdf_reports_contain_the_section(design_envs: dict[str, Any],
                                                  tmp_path: Path, only_network: None,
                                                  monkeypatch: pytest.MonkeyPatch) -> None:
    from docx import Document

    from ntnxvision_core.exporters import pdf as pdf_mod
    from ntnxvision_core.exporters.word import WordExporter

    env = design_envs["medium"]
    doc = Document(str(WordExporter().export(env, tmp_path / "add.docx")))
    headings = [p.text for p in doc.paragraphs
                if p.style is not None and p.style.name == "Heading 2"]
    assert any(h.endswith("Host Network Configuration") for h in headings)
    descr = [el.get("descr") or "" for el in doc.element.body.iter()
             if el.tag.endswith("}docPr")]
    assert sum("Host network profile" in d for d in descr) == 2
    text = "\n".join([p.text for p in doc.paragraphs]
                     + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "DESIGN SUMMARY" in text and "Host network profiles" in text
    assert "Balance-tcp with LACP, 2 x 25 GbE, 9,000" in text
    assert "Not reported (see vCenter)" in text

    drawn: list[str] = []
    real = pdf_mod.draw_scene

    def spy(pdf: Any, sc: S.Scene, *a: Any, **kw: Any) -> float:
        drawn.append(sc.caption)
        return real(pdf, sc, *a, **kw)

    monkeypatch.setattr(pdf_mod, "draw_scene", spy)
    pdf_mod.PdfExporter().export(env, tmp_path / "r.pdf")
    assert [c.split(":")[0] for c in drawn] == ["Host network profile A",
                                                "Host network profile B"]


def test_profiles_beyond_the_drawn_ones_are_noted(design_envs: dict[str, Any],
                                                  only_network: None,
                                                  monkeypatch: pytest.MonkeyPatch) -> None:
    from ntnxvision_core.diagrams.layout import network as L
    from ntnxvision_core.exporters import ExportOptions
    from ntnxvision_core.exporters.content import SceneFigure, Table, build_report

    monkeypatch.setattr(N, "MAX_DRAWN", 1)
    monkeypatch.setattr(L, "MAX_DRAWN", 1)
    report = build_report(design_envs["medium"], ExportOptions(), "x")
    section = next(s for ch in report.chapters for s in ch.sections
                   if s.title == "Host Network Configuration")
    figs = [b for b in section.blocks if isinstance(b, SceneFigure)]
    table = next(b for b in section.blocks if isinstance(b, Table))
    assert len(figs) == 1 and 2 <= len(figs[0].summary) <= 5
    assert figs[0].note == f"Profile B is listed in {table.label.rstrip('.')}."
    assert [r[0] for r in table.rows] == ["A", "B", "ESXi"]
