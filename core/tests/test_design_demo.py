"""Design data end to end against the mock Nutanix API (demo mode).

Real `PrismConnection`, real collectors, synthetic payloads: every new endpoint and its
fallback is exercised the way a lab run would exercise it.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.engine import collect_environment, combine_environments
from ntnxvision_core.exporters import export_environment
from ntnxvision_core.models import Environment
from ntnxvision_core.runinfo import RunTracker

DESIGN_AREAS = ["clusters", "hosts", "storage_containers", "host_networking",
                "data_protection", "files", "objects", "pc_config"]


def _collect(host: str, areas: list[str] | None = None, **kw: Any
             ) -> tuple[Environment, dict[str, dict[str, Any]]]:
    tracker = RunTracker(lambda _e: None, run_id="t", host=host)
    env = collect_environment(PrismCredentials(host=host, username="admin", password="demo"),
                              areas=areas, tracker=tracker, **kw)
    tracker.finish()
    return env, {a["id"]: a for a in tracker.final_areas()}


def _design(env: Environment) -> dict[str, Any]:
    return env.model_dump(include={"blocks", "cvms", "virtual_switches", "host_nics",
                                   "external_storages", "pc_vms", "pc_registrations",
                                   "pc_config", "protection_policies", "recovery_plans"})


def test_pc02_scale_out_flasharray_lacp_and_lldp(demo: Any, tmp_path: Path) -> None:
    env, areas = _collect("pc02.demo.local", DESIGN_AREAS)
    for area in ("hosts", "storage_containers", "host_networking", "pc_config",
                 "data_protection"):
        assert areas[area]["status"] == "collected", areas[area]
        assert areas[area]["api"] == "v4"

    cfg = env.pc_config
    assert (cfg.size, cfg.node_count, cfg.is_scale_out) == ("LARGE", 3, True)
    assert cfg.virtual_ip and cfg.hosting_cluster_name == "SEA-PROD-01"
    assert len(env.pc_vms) == 3 and all(v.hosting_cluster_name == "SEA-PROD-01"
                                        for v in env.pc_vms)
    assert "NUTANIX_DISASTER_RECOVERY" in cfg.enabled_products
    assert {r.remote_type for r in env.pc_registrations} == {"AOS", "DOMAIN_MANAGER"}

    (array,) = env.external_storages
    assert (array.product, array.api_source, array.cluster_name) == (
        "Pure FlashArray", "v4", "SEA-PROD-02")
    ctr = next(c for c in env.storage_containers if c.external_storage_uuid)
    assert ctr.external_storage_name == array.name and ctr.name in array.container_names

    bonds = {(s.name, s.bond_mode) for s in env.virtual_switches}
    assert ("vs0", "ACTIVE_BACKUP") in bonds and ("vs1-storage", "BALANCE_TCP") in bonds
    storage = next(s for s in env.virtual_switches if s.name == "vs1-storage")
    assert storage.lacp and storage.mtu == 9000
    assert all(h.lacp_status == "NEGOTIATED" for h in storage.hosts)

    first = [n for n in env.host_nics if n.cluster_name == "SEA-PROD-01"]
    assert first and all(n.switch_vendor and n.switch_port_id for n in first)
    second = [n for n in env.host_nics if n.cluster_name == "SEA-PROD-02"]
    assert second and all(n.switch_vendor and not n.switch_port_id for n in second)
    rest = [n for n in env.host_nics if n.cluster_name in ("SEA-VDI-01", "SEA-ROBO-01")]
    assert rest and not any(n.has_neighbour for n in rest)

    assert {c.categories.get("POD", [""])[0] for c in env.clusters} == {"SEA-1", "SEA-2"}
    robo = next(c for c in env.clusters if c.name == "SEA-ROBO-01")
    assert (robo.fault_domain, robo.cluster_functions) == ("DISK", ["AOS", "ONE_NODE"])
    assert all(c.nearsync_capable is not None for c in env.clusters)
    assert len(env.cvms) == len(env.hosts) == len(env.blocks)

    silver = next(p for p in env.protection_policies if p.name == "Silver-1h")
    assert silver.is_bidirectional and silver.source_clusters == ["SEA-PROD-02"]
    assert env.recovery_plans[0].api_source == "v4"

    xlsx = export_environment(env, "excel", tmp_path, "pc02")
    sheets = load_workbook(xlsx, read_only=True).sheetnames
    for title in ("Virtual Switches", "Host NICs", "External Storage", "PC VMs"):
        assert title in sheets
    header = next(load_workbook(xlsx, read_only=True)["Host NICs"].iter_rows(values_only=True))
    assert header[0] == "Cluster"


def test_pc01_esxi_hosts_have_nics_but_no_switches(demo: Any) -> None:
    env, areas = _collect("pc01.demo.local", DESIGN_AREAS)
    esx = [n for n in env.host_nics if n.cluster_name == "AMS-EDGE-02"]
    assert esx and all(n.speed_mbps and n.model is None and n.virtual_switch_uuid is None
                       for n in esx)
    assert not any(s.cluster_name == "AMS-EDGE-02" for s in env.virtual_switches)
    slb = {s.cluster_name for s in env.virtual_switches if s.bond_mode == "BALANCE_SLB"}
    assert slb == {"PHX-DR-01"}
    status = {s.name: s.status for s in env.pc_config.services}
    assert status == {"FILES": "NOT_DEPLOYED", "OBJECTS": "DEPLOYED"}
    assert env.pc_config.size == "SMALL" and env.pc_config.virtual_ip is None
    assert areas["objects"]["api"] == "v3"
    # PROD-02 (12 nodes) is rack aware: its blocks carry rack ids.
    racks = {h.rack_id for h in env.hosts if h.cluster_name == "PHX-PROD-02"}
    assert racks == {1, 2, 3}


def test_pc03_below_v4_uses_v3_and_prism_element(demo: Any) -> None:
    # Below v4 the CVMs come from the CVM system VMs: the VM area is needed.
    env, areas = _collect("pc03.demo.local", [*DESIGN_AREAS, "vms"])
    assert areas["hosts"]["status"] == "collected" and areas["hosts"]["api"] == "v3"
    assert areas["host_networking"]["status"] == "collected"
    assert areas["host_networking"]["api"] == "v2"
    assert env.virtual_switches == []
    assert env.host_nics and {n.api_source for n in env.host_nics} == {"v2.0"}
    assert env.blocks and {b.api_source for b in env.blocks} == {"v2.0"}
    assert all(h.block_position == 1 for h in env.hosts)
    assert env.cvms and all(c.cluster_uuid and c.host_uuid and c.hypervisor_type == "AHV"
                            for c in env.cvms)
    assert all(c.fault_domain for c in env.clusters)
    assert {p.api_source for p in env.pc_config.products} == {"v3"}
    assert env.pc_vms and env.pc_config.size == "SMALL"
    assert env.recovery_plans and env.recovery_plans[0].api_source == "v3"


def test_external_storage_from_groups_and_no_lldp(demo: Any) -> None:
    env2, _ = _collect("large2.demo.local", ["clusters", "hosts", "storage_containers"])
    (array,) = env2.external_storages
    assert (array.product, array.api_source, array.cluster_name) == (
        "Dell PowerFlex", "v3-groups", "PROD-A")
    assert array.container_names
    env3, _ = _collect("large3.demo.local", ["clusters", "hosts", "host_networking"])
    assert env3.host_nics and not any(n.has_neighbour for n in env3.host_nics)
    assert env3.external_storages == []


def test_scoped_run_keeps_only_the_selected_clusters_design_data(demo: Any) -> None:
    full, _ = _collect("pc01.demo.local", ["clusters"])
    keep = next(c.uuid for c in full.clusters if c.name == "PHX-PROD-01")
    env, _ = _collect("pc01.demo.local", ["clusters", "hosts", "host_networking",
                                          "storage_containers"], clusters=[keep])
    for name in ("blocks", "cvms", "virtual_switches", "host_nics"):
        items = getattr(env, name)
        assert items and {o.cluster_uuid for o in items} == {keep}, name


def test_scoped_run_keeps_only_the_selected_clusters_registrations(demo: Any) -> None:
    full, _ = _collect("pc02.demo.local", ["clusters"])
    keep = next(c.uuid for c in full.clusters if c.name == "SEA-PROD-01")
    env, _ = _collect("pc02.demo.local", ["pc_config", "storage_containers"],
                      clusters=[keep])
    aos = [r for r in env.pc_registrations if r.remote_type == "AOS"]
    assert [r.remote_uuid for r in aos] == [keep]
    # The paired Prism Centrals are PC-wide and stay.
    assert {r.name for r in env.pc_registrations if r.is_remote_pc} == {
        "pc-dr", "pc01.demo.local"}
    # The FlashArray belongs to SEA-PROD-02, outside the scope.
    assert env.external_storages == []


def test_object_counts_are_the_areas_main_objects(demo: Any) -> None:
    env, areas = _collect("pc02.demo.local", ["hosts", "host_networking"])
    assert areas["hosts"]["objects"] == len(env.hosts)
    assert areas["host_networking"]["objects"] == len(env.virtual_switches)


def test_large1_has_four_node_blocks(demo: Any) -> None:
    env, _ = _collect("large1.demo.local", ["clusters", "hosts"])
    assert {len(b.nodes) for b in env.blocks} >= {4}
    assert {h.block_position for h in env.hosts} == {1, 2, 3, 4}
    edge = next(c for c in env.clusters if c.name == "EDGE-A")  # 6 nodes, 2 blocks
    assert edge.fault_domain == "NODE"
    vdi = next(c for c in env.clusters if c.name == "VDI-A")  # 12 nodes: rack aware
    assert vdi.fault_domain == "RACK"
    racks = {b.rack_id for b in env.blocks if b.cluster_name == "VDI-A"}
    assert racks == {1, 2, 3}


def test_design_data_is_deterministic_across_collections(demo: Any) -> None:
    a, _ = _collect("pc02.demo.local", DESIGN_AREAS)
    b, _ = _collect("pc02.demo.local", DESIGN_AREAS)
    assert _design(a) == _design(b)


def test_pc01_sync_and_nearsync_policies(demo: Any) -> None:
    env, _ = _collect("pc01.demo.local", ["clusters", "data_protection", "pc_config"])
    metro = next(p for p in env.protection_policies if p.name == "Metro-PHX-Sync")
    assert (metro.replication_type, metro.rpo_seconds) == ("SYNC", 0)
    assert (metro.source_clusters, metro.target_clusters) == (["PHX-PROD-01"], ["PHX-DR-01"])
    assert metro.replications[0].sync_auto_suspend_timeout_secs == 10
    near = next(p for p in env.protection_policies if p.name == "NearSync-PHX-SEA")
    assert (near.replication_type, near.target_az) == ("NEARSYNC", "pc02.demo.local")
    peer = next(r for r in env.pc_registrations if r.name == "pc02.demo.local")
    assert peer.is_remote_pc


def test_combined_run_names_the_peer_clusters(demo: Any) -> None:
    one, _ = _collect("pc01.demo.local", ["clusters", "data_protection", "pc_config"])
    two, _ = _collect("pc02.demo.local", ["clusters", "data_protection", "pc_config"])
    near = next(p for p in one.protection_policies if p.name == "NearSync-PHX-SEA")
    assert near.target_clusters != ["SEA-PROD-01"]  # pc02's cluster: unknown alone
    combined = combine_environments([one, two])
    near = next(p for p in combined.protection_policies if p.name == "NearSync-PHX-SEA")
    assert near.target_clusters == ["SEA-PROD-01"]
    # Each PC's registration of the other carries the other's own PC uuid.
    pcs = {s.prism_central.host: s.prism_central.uuid for s in combined.source_environments}
    regs = {(r.source_pc, r.remote_uuid) for r in combined.pc_registrations if r.is_remote_pc}
    assert ("pc01.demo.local", pcs["pc02.demo.local"]) in regs
    assert ("pc02.demo.local", pcs["pc01.demo.local"]) in regs


def test_combined_run_keeps_design_data_per_pc(demo: Any) -> None:
    one, _ = _collect("pc01.demo.local", DESIGN_AREAS)
    two, _ = _collect("pc02.demo.local", DESIGN_AREAS)
    combined = combine_environments([one, two])
    assert len(combined.pc_vms) == 4
    assert {v.source_pc for v in combined.pc_vms} == {"pc01.demo.local", "pc02.demo.local"}
    assert {s.pc_config.size for s in combined.source_environments} == {"SMALL", "LARGE"}
    assert combined.pc_config.products == []  # no single management plane
    assert len(combined.host_nics) == len(one.host_nics) + len(two.host_nics)
