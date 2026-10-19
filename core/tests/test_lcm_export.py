"""LCM (Life Cycle Manager) rendering in the Excel, PDF and Word exporters.

Covers the three cases the data actually produces: an unknown update flag (which must
never read as "up to date"), a component LCM offers nothing for, and one with more
upgrade targets than the constrained formats can show.
"""

from __future__ import annotations

from datetime import datetime

from docx import Document
from openpyxl import load_workbook

from ntnxvision_core.engine import combine_environments
from ntnxvision_core.exporters import export_environment
from ntnxvision_core.exporters.excel import ExcelExporter
from ntnxvision_core.exporters.pdf import PdfExporter
from ntnxvision_core.exporters.word import GLYPH, WordExporter
from ntnxvision_core.models import Cluster, Environment, LcmEntity, PrismCentral

_SHEET = "LCM (Life Cycle Manager)"

# LCM hands back up to a dozen upgrade targets for a component like AOS; the PDF and
# the ADD trim this list, the Excel sheet carries it whole.
_MANY_VERSIONS = [f"7.0.{i}" for i in range(10, 0, -1)]


def _lcm_env() -> Environment:
    env = Environment(collected_at="2026-09-05T00:00:00Z")
    env.prism_central.host = "pc.lab.local"
    env.prism_central.version = "pc.2024.3"
    env.clusters += [
        Cluster(uuid="c1", name="cluster-01", aos_version="7.0", hypervisor_types=["AHV"],
                num_nodes=3, redundancy_factor=2),
        Cluster(uuid="c2", name="cluster-02", aos_version="6.8", hypervisor_types=["AHV"],
                num_nodes=4, redundancy_factor=2),
    ]
    env.lcm_entities += [
        LcmEntity(
            uuid="lcm-1", name="AHV", entity_class="Hypervisor", entity_type="software",
            current_version="10.0", available_versions=["10.1", "10.0.1"],
            update_available=True, cluster_uuid="c1", cluster_name="cluster-01",
            last_updated="2026-09-01T10:00:00Z",
        ),
        LcmEntity(
            uuid="lcm-2", name="AOS", entity_class="Core Cluster", current_version="7.0",
            target_version="7.0.10", available_versions=_MANY_VERSIONS,
            update_available=True, cluster_uuid="c1", cluster_name="cluster-01",
        ),
        # Nothing offered: LCM knows the component and has no newer build for it.
        LcmEntity(
            uuid="lcm-3", name="Foundation", entity_class="Core Cluster",
            current_version="5.7", update_available=False, cluster_uuid="c2",
            cluster_name="cluster-02",
        ),
        # Unknown: LCM reported the component without an upgrade verdict.
        LcmEntity(
            uuid="lcm-4", name="SATADOM SL 3ME4", entity_class="Core Cluster",
            current_version="S670330N", cluster_uuid="c2", cluster_name="cluster-02",
        ),
        LcmEntity(
            uuid="lcm-5", name="Files", entity_class="Files", current_version="5.0.0",
            available_versions=["5.1.0"], update_available=True,
        ),
    ]
    env.collection_notes["lcm"] = "ok"
    return env


def _empty_env() -> Environment:
    env = Environment(collected_at="2026-09-05T00:00:00Z")
    env.prism_central.host = "pc.lab.local"
    env.clusters.append(Cluster(uuid="c1", name="cluster-01", num_nodes=1))
    env.collection_notes["lcm"] = "skipped: LCM inventory not available"
    return env


# ─── Excel ─────────────────────────────────────────────────────────────


def test_excel_lcm_sheet_headers_and_rows(tmp_path) -> None:
    out = ExcelExporter().export(_lcm_env(), tmp_path / "lcm.xlsx")
    ws = load_workbook(out)[_SHEET]

    headers = [c.value for c in ws[1]]
    assert headers == [  # Cluster first: the house column-order rule
        "Cluster", "Component", "Class", "Current Version",
        "Update Available", "Available Versions", "Last Updated", "API",
    ]
    # `target_version` is intentionally absent: LCM reports it equal to the running
    # version unless an upgrade is staged (34/34 on the lab PC), so a column for it
    # duplicates Current Version.
    assert "Target Version" not in headers
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:H6"

    col = {name: i for i, name in enumerate(headers)}
    rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["AHV"][:4] == ("cluster-01", "AHV", "Hypervisor", "10.0")
    assert rows["AHV"][col["Update Available"]] == "Yes"
    assert rows["AHV"][col["Available Versions"]] == "10.1, 10.0.1"
    # A real Excel timestamp (UTC), shown as yyyy-mm-dd hh:mm, so it sorts and filters.
    assert rows["AHV"][col["Last Updated"]] == datetime(2026, 9, 1, 10, 0)
    assert rows["AHV"][col["API"]] == "v4"
    assert rows["Foundation"][col["Update Available"]] == "No"
    # Excel is the full-fidelity format: no trimming of the version list here.
    assert rows["AOS"][col["Available Versions"]] == ", ".join(_MANY_VERSIONS)


def test_excel_unknown_update_flag_is_not_no(tmp_path) -> None:
    """`update_available=None` means LCM did not say — reporting "No" would claim the
    component is current, the opposite of what is known.
    """
    out = ExcelExporter().export(_lcm_env(), tmp_path / "lcm.xlsx")
    ws = load_workbook(out)[_SHEET]
    col = {name: i for i, name in enumerate(c.value for c in ws[1])}
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[1] == "SATADOM SL 3ME4")
    assert row[col["Update Available"]] == "Unknown"
    # empty available_versions renders as the N/A placeholder, never a blank cell
    assert row[col["Available Versions"]] == "N/A"


def test_excel_omits_lcm_sheet_when_not_collected(tmp_path) -> None:
    out = ExcelExporter().export(_empty_env(), tmp_path / "empty.xlsx")
    wb = load_workbook(out)
    assert _SHEET not in wb.sheetnames
    assert "Collection Notes" in wb.sheetnames  # the "why" still lands somewhere


def test_excel_combined_run_attributes_lcm_rows_to_a_pc(tmp_path) -> None:
    """Combined reports add a "Prism Central" column to every inventory sheet (right
    after Cluster);
    LCM component names (AHV, AOS) repeat across estates, so without it the merged
    rows are indistinguishable.
    """
    a, b = _lcm_env(), _lcm_env()
    a.prism_central = PrismCentral(host="pc-a.lab.local")
    b.prism_central = PrismCentral(host="pc-b.lab.local")
    out = ExcelExporter().export(combine_environments([a, b]), tmp_path / "combined.xlsx")

    ws = load_workbook(out)[_SHEET]
    assert [c.value for c in ws[1]][:3] == ["Cluster", "Prism Central", "Component"]
    sources = {r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert sources == {"pc-a.lab.local", "pc-b.lab.local"}
    assert ws.max_row == 11  # header + 5 components from each PC


def test_combined_run_attributes_lcm_rows_in_pdf_and_word(tmp_path) -> None:
    a, b = _lcm_env(), _lcm_env()
    a.prism_central = PrismCentral(host="pc-a.lab.local")
    b.prism_central = PrismCentral(host="pc-b.lab.local")
    combined = combine_environments([a, b])

    doc = Document(str(WordExporter().export(combined, tmp_path / "combined.docx")))
    table = next(t for t in doc.tables
                 if len(t.rows[0].cells) > 2 and t.rows[0].cells[2].text == "Component")
    assert [c.text for c in table.rows[0].cells[:2]] == ["Cluster", "Prism Central"]
    assert {r.cells[1].text for r in table.rows[1:]} == {"pc-a.lab.local", "pc-b.lab.local"}

    # The PDF's column widths are rescaled around the extra column; assert it renders.
    assert PdfExporter().export(combined, tmp_path / "combined.pdf").stat().st_size > 1000


# ─── PDF ───────────────────────────────────────────────────────────────


def test_pdf_export_includes_lcm(tmp_path) -> None:
    out = PdfExporter().export(_lcm_env(), tmp_path / "lcm.pdf")
    assert out.exists()
    assert out.stat().st_size > 1000


def test_pdf_export_without_lcm_still_renders(tmp_path) -> None:
    out = PdfExporter().export(_empty_env(), tmp_path / "empty.pdf")
    assert out.exists()
    assert out.stat().st_size > 1000


def test_pdf_version_trimming() -> None:
    """The PDF column is too narrow for a full list; trimming keeps whole version
    strings rather than letting the cell truncate one mid-number.
    """
    from ntnxvision_core.exporters.pdf import _versions

    assert _versions([], 2) == "N/A"
    assert _versions(["7.1"], 2) == "7.1"
    assert _versions(["7.1", "7.0"], 2) == "7.1, 7.0"
    assert _versions(_MANY_VERSIONS, 2) == "7.0.10, 7.0.9 (+8 more)"


# ─── Word (Architecture Design Document) ───────────────────────────────


def test_word_lifecycle_section_and_counts(tmp_path) -> None:
    out = WordExporter().export(_lcm_env(), tmp_path / "add.docx")
    doc = Document(str(out))

    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Platform Lifecycle (LCM)" in h for h in headings)

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Of the 5 component(s) inventoried, 3 have a newer version available" in text
    assert "1 component(s) LCM did not report an upgrade target" in text

    table = next(t for t in doc.tables if t.rows[0].cells[1].text == "Component")
    assert [c.text for c in table.rows[0].cells] == [
        "Cluster", "Component", "Class", "Current Version", "Update Available",
        "Available Versions",
    ]
    by_component = {r.cells[1].text: [c.text for c in r.cells] for r in table.rows[1:]}
    assert by_component["SATADOM SL 3ME4"][4] == "Unknown"
    assert by_component["SATADOM SL 3ME4"][5] == "N/A"
    # "No update" is a healthy status: an ok glyph precedes the ink label.
    assert by_component["Foundation"][4] == f"{GLYPH}\u00a0No"
    assert by_component["AOS"][5] == "7.0.10, 7.0.9, 7.0.8 (+7 more)"


def test_word_omits_lifecycle_section_when_not_collected(tmp_path) -> None:
    out = WordExporter().export(_empty_env(), tmp_path / "empty.docx")
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert not any("Lifecycle" in h for h in headings)
    # Section numbering closes over the gap rather than skipping a number.
    h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1" and p.text[:1].isdigit()]
    numbers = [h.split(".", 1)[0] for h in h1s]
    assert numbers == [str(i) for i in range(1, len(numbers) + 1)]


# ─── dispatcher ────────────────────────────────────────────────────────


def test_all_formats_export_with_lcm(tmp_path) -> None:
    env = _lcm_env()
    xlsx = export_environment(env, "excel", tmp_path, "lcm")
    pdf = export_environment(env, "pdf", tmp_path, "lcm")
    docx = export_environment(env, "word", tmp_path, "lcm")
    assert _SHEET in load_workbook(xlsx).sheetnames
    assert pdf.stat().st_size > 1000
    assert docx.stat().st_size > 1000
