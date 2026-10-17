"""End to end without a lab: mock server → real collectors → snapshot → exports → compare.

Every layer except the network is the production code: the real `PrismConnection`,
`ElementConnection` and `NdbConnection` speak HTTPS to the local mock, the real
collectors parse the synthetic payloads, and the real exporters write the files.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.api.element_client import ElementCredentials
from ntnxvision_core.api.ndb_client import NdbCredentials
from ntnxvision_core.compare import compare, export_comparison
from ntnxvision_core.engine import (
    collect_element_environment,
    collect_environment,
    combine_environments,
)
from ntnxvision_core.exporters import export_environment
from ntnxvision_core.runinfo import RunTracker, build_run_metadata
from ntnxvision_core.store import SnapshotStore
from ntnxvision_core.synthetic import demo as demo_mode


def _creds(host: str) -> PrismCredentials:
    return PrismCredentials(host=host, username="admin", password="demo")


def test_collect_snapshot_export_compare(demo: Any, tmp_path: Path) -> None:
    events: list[dict[str, Any]] = []
    tracker = RunTracker(events.append, run_id="r1", host="pc01.demo.local")
    env = collect_environment(
        _creds("pc01.demo.local"), tracker=tracker,
        ndb_creds=NdbCredentials(host="ndb.demo.local", username="admin", password="x"))
    tracker.finish()

    # The realistic edge cases all surfaced.
    by_id = {a["id"]: a for a in tracker.final_areas()}
    assert by_id["vms"]["status"] == "notes"
    assert "3 ESXi VM(s) read from Prism Element" in by_id["vms"]["detail"]
    assert by_id["files"]["status"] == "skipped"
    assert by_id["alerts"]["status"] == "notes"  # one page retried after a 429
    assert by_id["objects"]["api"] == "v3"  # v4 objects service down (503)
    assert by_id["clusters"]["api"] == "v4"
    assert by_id["ndb"]["status"] == "collected" and by_id["ndb"]["api"] == "era"
    assert len(env.databases) == 12
    assert env.pc_config.pc_vm_names and env.system_vms
    # Cross references resolved across APIs (v4 VMs only carry uuids).
    vm = env.vms[0]
    assert vm.host_name and vm.cluster_name and vm.disks[0].storage_container_name
    assert vm.nics[0].subnet_name and vm.nics[0].vlan_id is not None

    env.run = build_run_metadata(env, tracker)
    path = SnapshotStore(tmp_path).save(env, "week1")
    loaded = SnapshotStore(tmp_path).load(path)
    assert loaded.run and loaded.run["areas"] and loaded.run["api_routing"]

    out = tmp_path / "out"
    xlsx, detail = export_environment(loaded, "excel", out, "week1", with_detail=True)
    assert detail and detail.endswith("sheets")
    assert "VMs" in load_workbook(xlsx).sheetnames
    pdf = export_environment(loaded, "pdf", out, "week1",
                             options={"prepared_for": "Acme", "author": "Tester"})
    assert pdf.read_bytes().startswith(b"%PDF")
    docx = export_environment(loaded, "word", out, "week1",
                              options={"prepared_for": "Acme Corp", "author": "Tester"})
    doc = Document(str(docx))
    assert doc.core_properties.author == "Tester"
    assert any("Acme Corp" in p.text for p in doc.paragraphs)

    # A week later: the same Prism Central has drifted.
    demo_mode.advance("pc01.demo.local")
    env2 = collect_environment(_creds("pc01.demo.local"))
    result = compare(loaded, env2)
    assert result["summary"] == {"added": 36, "removed": 7, "changed": 10, "total": 53}
    rows = {(r["kind"], r["object_type"], r["object"]) for r in result["changes"]}
    assert ("added", "VM", "k8s-worker-168 … 199") in rows
    assert ("removed", "VM", "legacy-app-01 … 06") in rows
    aos = next(r for r in result["changes"] if r["field"] == "AOS")
    assert (aos["before"], aos["after"]) == ("6.10.1", "7.0")
    xl = export_comparison(loaded, env2, "week1", "week2", "excel", out)
    sheet = load_workbook(xl)["Changes"]
    assert sheet.max_row - 1 >= 53  # ungrouped: one row per object / field


def test_v3_only_prism_central(demo: Any) -> None:
    env = collect_environment(_creds("pc03.demo.local"))
    sources = {o.api_source for o in [*env.clusters, *env.hosts, *env.vms]}
    assert sources == {"v3"}
    assert env.collection_notes["licenses"].startswith("skipped")
    assert env.collection_notes["lcm_entities"].startswith("skipped")
    assert len(env.security_policies) > 0  # v3 still lists policies on this PC


def test_prism_element_mode(demo: Any) -> None:
    tracker = RunTracker(None, mode="prism_element")
    env = collect_element_environment(
        ElementCredentials(host="pe01.demo.local", username="admin", password="x"),
        tracker=tracker)
    assert env.source_type == "prism_element"
    assert len(env.clusters) == 1 and len(env.hosts) == 6 and len(env.vms) == 212
    assert {a["api"] for a in tracker.final_areas()} == {"v2"}
    assert env.vms[0].host_name and env.vms[0].disks


def test_area_selection_limits_collectors(demo: Any) -> None:
    tracker = RunTracker(None)
    env = collect_environment(_creds("tiny.demo.local"), areas=["clusters", "vms"],
                              tracker=tracker)
    assert [a["id"] for a in tracker.final_areas()] == ["clusters", "vms"]
    assert set(env.collection_notes) == {"clusters", "vms"}
    assert env.vms and not env.hosts and not env.alerts


def test_twenty_thousand_vm_estate_stays_fast(demo: Any, tmp_path: Path) -> None:
    import time

    from ntnxvision_core.inventory import InventoryService
    from ntnxvision_core.store.snapshot import EnvironmentCache

    envs = [collect_environment(_creds(f"large{i}.demo.local"), areas=["clusters", "vms"])
            for i in (1, 2, 3)]
    combined = combine_environments(envs)
    assert len(combined.vms) >= 20_000
    assert {v.source_pc for v in combined.vms} == {f"large{i}.demo.local" for i in (1, 2, 3)}

    path = SnapshotStore(tmp_path).save(combined, "large")
    t0 = time.perf_counter()
    env = EnvironmentCache().get(path)
    svc = InventoryService()
    first = svc.query(env, "vms", {}, None, {"offset": 0, "limit": 25})
    cold = time.perf_counter() - t0
    assert first["total"] >= 20_000
    timings = []
    for filt, sort in (({"text": "k8s"}, {"key": "memory", "dir": "desc"}),
                       ({"facets": {"power": "Off"}}, {"key": "name", "dir": "asc"}),
                       ({"text": "web 10."}, None)):
        t0 = time.perf_counter()
        res = svc.query(env, "vms", filt, sort, {"offset": 500, "limit": 25})
        timings.append(time.perf_counter() - t0)
        assert res["matching"] > 0
    # Generous bounds for slow CI runners; locally ~1 s cold and ~5 ms per query.
    assert cold < 15
    assert max(timings) < 1.0
