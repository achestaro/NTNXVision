"""No report shows a full Prism Central / AOS build string.

The synthetic Prism Centrals answer their version fields with a realistic full build
string (``el8.5-release-fraser-<version>-stable-<40-hex>``, a fake commit), as a real
PC does. Every export (Word, PDF, Excel, compare exports, combined multi-PC reports and
Prism Element mode) must show the short version instead. The single exception is the
"Build" row of the Excel Prism Central sheet, which keeps the full string on purpose.
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.api.element_client import ElementCredentials
from ntnxvision_core.compare import export_comparison
from ntnxvision_core.compat import check_aos, check_prism_central, parse_version
from ntnxvision_core.engine import (
    collect_element_environment,
    collect_environment,
    combine_environments,
)
from ntnxvision_core.exporters import export_environment
from ntnxvision_core.models import Environment, PrismCentral, PrismCentralConfig
from ntnxvision_core.store import SnapshotStore
from ntnxvision_core.synthetic import demo as demo_mode
from ntnxvision_core.versions import pc_version

#: A full build string. Loose on purpose: one with a short suffix ("-stable-pc-0")
#: instead of a 40-hex commit is caught too.
LONG_BUILD = re.compile(r"el\d[\w.]*-release-", re.I)
FAKE = "el8.5-release-fraser-7.5.1.1-stable-" + "ab" * 20


def _creds(host: str) -> PrismCredentials:
    return PrismCredentials(host=host, username="admin", password="demo")


# ─── report text extraction ──────────────────────────────────────────────


def docx_text(path: Path) -> str:
    """Every XML part of the package: body, headers, footers, core properties."""
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist()
                         if n.endswith((".xml", ".rels")))


@pytest.fixture
def pdf_text(monkeypatch: pytest.MonkeyPatch) -> Iterator[list[str]]:
    """Records every string fpdf2 lays out (cells, multi-cells, free text, metadata).

    The PDF's own text is glyph-encoded font subsets, so it is captured on the way in.
    """
    import fpdf

    seen: list[str] = []
    normalize = fpdf.FPDF.normalize_text

    def record(self: Any, text: str) -> str:
        seen.append(text)
        return normalize(self, text)

    monkeypatch.setattr(fpdf.FPDF, "normalize_text", record)
    for name in ("set_title", "set_subject", "set_keywords", "set_author"):
        original = getattr(fpdf.FPDF, name)

        def meta(self: Any, value: str, *a: Any, _o: Any = original, **k: Any) -> Any:
            seen.append(value)
            return _o(self, value, *a, **k)

        monkeypatch.setattr(fpdf.FPDF, name, meta)
    yield seen


def xlsx_cells(path: Path) -> Iterator[tuple[str, str, str]]:
    """(sheet, coordinate, text) for every non-empty cell."""
    wb = load_workbook(path, read_only=False)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    yield ws.title, cell.coordinate, str(cell.value)


def assert_no_long_build(text: str, where: str) -> None:
    hit = LONG_BUILD.search(text)
    assert hit is None, f"{where} shows a full build string: {hit.group(0) if hit else ''}"


def check_reports(env: Environment, out: Path, name: str, pdf_seen: list[str],
                  build_cell: bool) -> str:
    """Write Word, PDF and Excel for `env`; return the docx text for further checks."""
    docx = export_environment(env, "word", out, name)
    text = docx_text(docx)
    assert_no_long_build(text, f"{name}.docx")

    pdf_seen.clear()
    export_environment(env, "pdf", out, name)
    assert pdf_seen, "the PDF recorder saw no text"
    assert_no_long_build("\n".join(pdf_seen), f"{name}.pdf")

    xlsx = export_environment(env, "excel", out, name)
    builds = []
    for sheet, coord, value in xlsx_cells(xlsx):
        if LONG_BUILD.search(value):
            builds.append((sheet, coord, value))
    if build_cell:
        # Exactly one: the Build row of the Prism Central sheet.
        assert len(builds) == 1, builds
        sheet, coord, _ = builds[0]
        ws = load_workbook(xlsx)[sheet]
        assert sheet == "Prism Central"
        assert ws[f"A{ws[coord].row}"].value == "Build"
    else:
        assert builds == []
    return text


# ─── model and compat ─────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("7.5.1.1", "pc.7.5.1.1"),
        ("7.5.1.6", "pc.7.5.1.6"),
        ("pc.2024.3", "pc.2024.3"),
        ("PC.2024.3", "pc.2024.3"),
        ("2024.3", "pc.2024.3"),
        (FAKE, "pc.7.5.1.1"),
        ("el8.5-release-fraser-pc.2024.3-stable-pc-0", "pc.2024.3"),
        ("unknown", "unknown"),
        (None, None),
        ("", None),
    ],
)
def test_pc_version_keeps_every_part_with_the_pc_prefix(raw: str | None,
                                                        expected: str | None) -> None:
    assert pc_version(raw) == expected


def test_assignment_shortens_versions_and_build_keeps_the_full_string() -> None:
    pc = PrismCentral()
    pc.version = FAKE
    assert pc.version == "7.5.1.1"
    cfg = PrismCentralConfig()
    cfg.version = FAKE
    cfg.ncc_version = "el8.5-release-fraser-5.1.0-stable-" + "cd" * 20
    cfg.build = FAKE
    assert (cfg.version, cfg.ncc_version, cfg.build) == ("7.5.1.1", "5.1.0", FAKE)


def test_compat_reads_the_version_inside_a_build_string() -> None:
    assert parse_version(FAKE) == ((7, 5, 1, 1), "semantic")  # not the "8.5" OS prefix
    assert check_prism_central(FAKE).in_range
    newer = check_prism_central("el8.5-release-fraser-7.6-stable-" + "ef" * 20)
    assert newer.raw == "7.6" and newer.warning
    assert "release-" not in newer.warning
    old = check_aos("el8.5-release-fraser-6.1-stable-" + "01" * 20)
    assert not old.in_range and old.warning and old.warning.startswith("AOS 6.1 ")


# ─── reports ─────────────────────────────────────────────────────────────


def test_prism_central_reports_show_the_short_version(
        demo: Any, tmp_path: Path, pdf_text: list[str]) -> None:
    env = collect_environment(_creds("pc01.demo.local"))
    # The mock answers like a real PC: the full build string is collected...
    assert env.pc_config.build and LONG_BUILD.search(env.pc_config.build)
    # ...and every version field holds the short form, even before a save and load.
    assert env.prism_central.version == env.pc_config.version == "pc.2024.3"
    assert all(not LONG_BUILD.search(c.aos_version or "") for c in env.clusters)

    out = tmp_path / "out"
    text = check_reports(env, out, "fresh", pdf_text, build_cell=True)
    # The cover: subtitle, source row and compatibility row.
    assert "Prism Central pc01.demo.local · pc.2024.3" in text
    assert "pc01.demo.local, pc.2024.3" in text
    assert ("Prism Central pc.2024.3 is within the supported range "
            "(pc.2024.1 to pc.7.5.1.6).") in text

    path = SnapshotStore(tmp_path).save(env, "saved")
    loaded = SnapshotStore(tmp_path).load(path)
    assert loaded.pc_config.build == env.pc_config.build
    check_reports(loaded, out, "loaded", pdf_text, build_cell=True)

    # Compare exports between two collections a week apart.
    demo_mode.advance("pc01.demo.local")
    later = collect_environment(_creds("pc01.demo.local"))
    for fmt in ("excel", "csv"):
        written = export_comparison(env, later, "week1", "week2", fmt, out)
        if fmt == "excel":
            for sheet, coord, value in xlsx_cells(written):
                assert_no_long_build(value, f"compare {sheet}!{coord}")
        else:
            assert_no_long_build(written.read_text(encoding="utf-8"), "compare csv")
    pdf_text.clear()
    export_comparison(env, later, "week1", "week2", "pdf", out)
    assert pdf_text
    assert_no_long_build("\n".join(pdf_text), "compare pdf")


def test_combined_reports_show_short_versions(
        demo: Any, tmp_path: Path, pdf_text: list[str]) -> None:
    envs = [collect_environment(_creds(h)) for h in ("pc01.demo.local", "pc02.demo.local")]
    combined = combine_environments(envs)
    assert [s.prism_central.version for s in combined.source_environments] == [
        "pc.2024.3", "pc.7.3"]
    # A combined workbook has no single Prism Central sheet, so no Build row either.
    text = check_reports(combined, tmp_path, "combined", pdf_text, build_cell=False)
    assert "pc01.demo.local (pc.2024.3)" in text
    assert "pc01.demo.local runs pc.2024.3" in text


def test_prism_element_reports_show_the_short_aos_version(
        demo: Any, tmp_path: Path, pdf_text: list[str]) -> None:
    env = collect_element_environment(
        ElementCredentials(host="pe01.demo.local", username="admin", password="x"))
    assert env.prism_central.version == "7.0"
    text = check_reports(env, tmp_path, "pe", pdf_text, build_cell=False)
    assert "AOS 7.0 is within the supported range (AOS 6.5 to 7.5)." in text


# ─── one format everywhere ────────────────────────────────────────────────

#: The numeric version pc04 reports, anywhere it is not preceded by "pc.".
BARE = re.compile(r"(?<![\w.-])(?<!pc\.)7\.5\.1\.1(?![\w.])")
SHOWN = "pc.7.5.1.1"


def test_a_numeric_version_shows_as_pc_everywhere(
        demo: Any, tmp_path: Path, pdf_text: list[str]) -> None:
    """pc04 reports "7.5.1.1": the cover, the tables, the prose, the diagram, the
    workbook, the snapshot index and the connection test all show "pc.7.5.1.1"."""
    from docx import Document

    from ntnxvision_core.connectivity import test_connection
    from ntnxvision_core.diagrams import conceptual_layout, conceptual_model
    from ntnxvision_core.diagrams.render.svg import render_svg
    from ntnxvision_core.runinfo import compat_info
    from ntnxvision_core.store.snapshot import read_meta

    env = collect_environment(_creds("pc04.demo.local"))
    assert env.prism_central.version == "7.5.1.1"  # the model keeps the plain value
    assert compat_info(env)["version"] == SHOWN

    text = check_reports(env, tmp_path, "pc04", pdf_text, build_cell=True)
    assert f"Prism Central pc04.demo.local · {SHOWN}" in text
    assert f"pc04.demo.local, {SHOWN}" in text
    assert (f"Prism Central {SHOWN} is within the supported range "
            "(pc.2024.1 to pc.7.5.1.6).") in text
    assert not BARE.search(text), BARE.findall(text)
    assert not BARE.search("\n".join(pdf_text))
    assert any(SHOWN in t for t in pdf_text)

    # Tables: the Prism Central property table and the availability zones table.
    doc = Document(str(tmp_path / "pc04.docx"))
    cells = {(row.cells[0].text, row.cells[1].text) for t in doc.tables for row in t.rows
             if len(row.cells) > 1}
    assert ("Version", SHOWN) in cells
    az_rows = [[c.text for c in row.cells] for t in doc.tables for row in t.rows
               if row.cells and "PC version" in [c.text for c in t.rows[0].cells]]
    assert az_rows and any(SHOWN in r for r in az_rows[1:])

    # Excel: the summary, the Prism Central sheet (Version; Build keeps the full string).
    wb = load_workbook(tmp_path / "pc04.xlsx")
    pc_rows = {r[0]: r[1] for r in wb["Prism Central"].iter_rows(values_only=True)}
    assert pc_rows["Version"] == SHOWN
    assert pc_rows["Build"] and "7.5.1.1-stable-" in pc_rows["Build"]
    for _sheet, _coord, value in xlsx_cells(tmp_path / "pc04.xlsx"):
        if not LONG_BUILD.search(value):
            assert not BARE.search(value), value

    # The diagram: every part, with the prefix.
    svgs = [render_svg(sc) for sc in conceptual_layout(conceptual_model(env))]
    svg = "".join(s if isinstance(s, str) else s.decode() for s in svgs)
    assert SHOWN in svg and not BARE.search(svg)

    # The snapshot index and the connection test.
    path = SnapshotStore(tmp_path).save(env, "pc04")
    assert read_meta(path)["version"] == SHOWN
    r = test_connection({"mode": "prism_central", "host": "pc04.demo.local", "port": 9440,
                         "username": "admin", "password": "demo", "insecure": False})
    assert r["version"] == SHOWN and r["message"].endswith(SHOWN)


def test_old_index_without_the_prefix_reads_with_it(tmp_path: Path) -> None:
    import json

    from ntnxvision_core.store.snapshot import meta_path, read_meta

    env = Environment.model_validate({"prism_central": {"host": "pc.example.lab",
                                                        "version": "7.5.1.1"}})
    path = SnapshotStore(tmp_path).save(env, "old")
    assert read_meta(path)["version"] == SHOWN
    meta = read_meta(path)
    meta.update(version="7.5.1.1", snapshot_mtime=path.stat().st_mtime)
    meta_path(path).write_text(json.dumps(meta), encoding="utf-8")
    assert read_meta(path)["version"] == SHOWN
