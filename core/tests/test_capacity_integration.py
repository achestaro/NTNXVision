"""Capacity figures end to end: collectors → model → inventory grid / detail → reports."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

import ntnxvision_core.sidecar as sc
from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.api.element_client import ElementCredentials
from ntnxvision_core.capacity import compute_environment
from ntnxvision_core.collectors.element import _cvms, _hosts
from ntnxvision_core.collectors.hosts import HostCollector
from ntnxvision_core.engine import (
    collect_element_environment,
    collect_environment,
    combine_environments,
)
from ntnxvision_core.exporters import ExportOptions, export_environment
from ntnxvision_core.inventory import InventoryService
from ntnxvision_core.models import Environment

TB = 10**12


# ─── collectors ───────────────────────────────────────────────────────


class _HostConn:
    def __init__(self, v4: bool, groups: list[dict[str, Any]] | None = None) -> None:
        self._v4 = v4
        self.groups = groups
        self.group_calls = 0

    def v4_ready(self, area: str) -> bool:
        return self._v4

    def v4_list(self, path: str) -> list[dict[str, Any]]:
        return [{"extId": "h1", "hostName": "n1", "numberOfCpuCores": 32,
                 "disk": [{"sizeInBytes": 2 * TB}, {"sizeInBytes": 2 * TB},
                          {"sizeInBytes": None}]},
                {"extId": "h2", "hostName": "n2", "numberOfCpuCores": 32}]

    def v3_list(self, kind: str) -> list[dict[str, Any]]:
        return [{"metadata": {"uuid": "h1"},
                 "status": {"name": "n1", "resources": {"num_cpu_cores": 32}}}]

    def v3_groups(self, entity: str, attrs: list[str], **_kw: Any) -> list[dict[str, Any]]:
        self.group_calls += 1
        assert entity == "host" and "storage.capacity_bytes" in attrs
        if self.groups is None:
            raise RuntimeError("groups unsupported")
        return self.groups


def test_v4_host_raw_storage_is_the_sum_of_its_disks() -> None:
    env = Environment()
    conn = _HostConn(v4=True, groups=[{"entity_id": "h2", "host_name": "n2",
                                       "storage.capacity_bytes": str(5 * TB)}])
    HostCollector(conn).collect(env)  # type: ignore[arg-type]
    by = {h.uuid: h for h in env.hosts}
    assert by["h1"].storage_capacity_bytes == 4 * TB and by["h1"].disk_count == 3
    # A v4 host without a disk list is backfilled from groups.
    assert by["h2"].storage_capacity_bytes == 5 * TB


def test_v3_host_raw_storage_comes_from_groups_and_tolerates_failure() -> None:
    env = Environment()
    HostCollector(_HostConn(v4=False, groups=[{"entity_id": "h1", "host_name": "n1",  # type: ignore[arg-type]
                                               "storage.capacity_bytes": "7000"}])).collect(env)
    assert env.hosts[0].storage_capacity_bytes == 7000
    env = Environment()
    HostCollector(_HostConn(v4=False, groups=None)).collect(env)  # type: ignore[arg-type]
    assert env.hosts[0].storage_capacity_bytes is None
    assert env.collection_notes["hosts"] == "ok"


class _PeConn:
    def __init__(self, v1_fails: bool = False) -> None:
        self.v1_fails = v1_fails

    def v2_list(self, path: str, **_kw: Any) -> list[dict[str, Any]]:
        return [{"uuid": "h1", "name": "n1", "num_cpu_cores": 24,
                 "usage_stats": {"storage.capacity_bytes": "123456"}}]

    def v1_list(self, path: str) -> list[dict[str, Any]]:
        if self.v1_fails:
            raise RuntimeError("v1 disabled")
        return [{"uuid": "c1", "vmName": "NTNX-A-CVM", "controllerVm": True, "numVCpus": 10,
                 "hostUuid": "h1", "ipAddresses": ["10.0.0.5"]},
                {"uuid": "u1", "vmName": "app", "controllerVm": False, "numVCpus": 4}]


def test_element_raw_storage_and_cvms() -> None:
    env = Environment(source_type="prism_element")
    _hosts(_PeConn(), env)  # type: ignore[arg-type]
    _cvms(_PeConn(), env)  # type: ignore[arg-type]
    assert env.hosts[0].storage_capacity_bytes == 123456
    assert [(v.name, v.vm_type, v.num_vcpus) for v in env.system_vms] == [
        ("NTNX-A-CVM", "CVM", 10)]
    env = Environment(source_type="prism_element")
    _cvms(_PeConn(v1_fails=True), env)  # type: ignore[arg-type]
    assert env.system_vms == []


# ─── synthetic estate, every API path ─────────────────────────────────


def _caps(env: Environment) -> dict[str, Any]:
    caps = compute_environment(env)
    return {c.name: caps[id(c)] for c in env.clusters}


def test_demo_v4_estate_has_heterogeneous_and_two_node_clusters(demo: Any) -> None:
    env = collect_environment(PrismCredentials(host="pc01.demo.local", username="admin",
                                               password="demo"))
    caps = _caps(env)
    mixed = caps["PHX-PROD-02"]
    assert {n.cores for n in mixed.nodes} == {64, 96}
    assert {n.cvm_vcpus for n in mixed.nodes} == {12, 16}
    # N-1 drops one of the large (96-core, 16 CVM vCPU) nodes.
    assert mixed.usable_cores_n1 == mixed.usable_cores - 80
    assert mixed.rf == 3 and mixed.cpu_ratio and mixed.storage_ratio
    edge = caps["AMS-EDGE-02"]
    assert edge.node_count == 2 and edge.cpu_ratio is not None and edge.cpu_ratio_n1 is None
    assert all(not n.cvm_estimated for k in caps.values() for n in k.nodes)


def test_demo_v3_and_single_node_paths(demo: Any) -> None:
    v3 = collect_environment(PrismCredentials(host="pc03.demo.local", username="admin",
                                              password="demo"))
    assert all(k.storage_total_bytes and k.usable_cores for k in _caps(v3).values())
    pc02 = collect_environment(PrismCredentials(host="pc02.demo.local", username="admin",
                                                password="demo"))
    robo = _caps(pc02)["SEA-ROBO-01"]
    assert robo.node_count == 1 and robo.usable_cores == 24
    assert robo.vcpu_total_n1 is None and robo.storage_total_n1_bytes is None
    combined = combine_environments([pc02, v3])
    names = {k.cluster_name: k.source_pc for k in compute_environment(combined).values()}
    assert names["SEA-ROBO-01"] == "pc02.demo.local"
    assert names["DAL-PROD-01"] == "pc03.demo.local"


def test_demo_prism_element_mode(demo: Any) -> None:
    env = collect_element_environment(ElementCredentials(host="pe01.demo.local",
                                                         username="admin", password="demo"))
    (cap,) = _caps(env).values()
    assert cap.node_count == 6 and cap.cvm_vcpus == 72
    assert cap.usable_cores and cap.storage_total_bytes and cap.cpu_ratio_n1


# ─── grid, detail, reports ────────────────────────────────────────────


def test_ratios_live_in_the_detail_pane_not_the_grid(demo: Any) -> None:
    env = collect_environment(PrismCredentials(host="pc02.demo.local", username="admin",
                                               password="demo"))
    svc = InventoryService()
    result = svc.query(env, "clusters", page={"limit": 50})
    cols = {c["key"] for c in result["columns"]}
    # The four ratios are detail-pane figures (grid=False); the grid and its Columns menu
    # (built from the same column list) never offer them.
    assert not cols & {"cpu_ratio", "cpu_ratio_n1", "storage_ratio", "storage_ratio_n1"}
    robo = next(r for r in result["rows"] if r["cells"]["name"] == "SEA-ROBO-01")
    assert "cpu_ratio" not in robo["cells"]
    detail = svc.detail(env, "clusters", robo["id"], vcpu_ratio=6)
    fields = {f["label"]: f["value"] for f in detail["fields"]}
    # No target recorded (collected outside the sidecar): the passed ratio, as default.
    assert fields["Target ratio"] == "6.0 : 1 (default)"
    assert fields["vCPU Total"] == "144 · N-1 N/A"
    assert fields["vCPU:pCore"].endswith(": 1 · N-1 N/A")
    assert fields["Storage ratio"].endswith(" · N-1 N/A")
    assert "vCPU:pCore N-1" not in fields and "Storage ratio N-1" not in fields
    assert any("N-1 is not applicable" in n["text"] for n in detail["notes"])


def test_reports_render_the_capacity_figures(demo: Any, tmp_path: Path) -> None:
    env = collect_environment(PrismCredentials(host="pc02.demo.local", username="admin",
                                               password="demo"))
    opts = ExportOptions(vcpu_ratio=5)
    xlsx = export_environment(env, "excel", tmp_path, "cap", opts)
    wb = load_workbook(xlsx)
    ws = wb["Capacity"]
    headers = [c.value for c in ws[1]]
    rows = {r[0]: dict(zip(headers, r, strict=False))
            for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1] is not None}
    robo = rows["SEA-ROBO-01"]
    assert robo["Target vCPU:pCore"] == "5.0 : 1 (default)"
    assert robo["vCPU Total"] == 120 and robo["vCPU Total N-1"] == "N/A"
    assert str(robo["vCPU:pCore (all VMs)"]).endswith(" : 1")
    legend = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert any("5.0 : 1 vCPU:pCPU target" in str(t) for t in legend)
    clusters = wb["Clusters"]
    cheads = [c.value for c in clusters[1]]
    # Capacity figures are on the Capacity sheet only (test_grid_columns.py).
    assert "vCPU:pCore (all VMs)" not in cheads and "Storage Total N-1 (GiB)" not in cheads
    overview = [str(c.value) for row in wb["Overview"].iter_rows() for c in row if c.value]
    assert any(v.startswith("5.0 : 1") for v in overview)

    docx = export_environment(env, "word", tmp_path, "cap", opts)
    text = "\n".join(p.text for p in Document(str(docx)).paragraphs)
    assert "Capacity & Oversubscription" in text and "5.0 : 1 vCPU:pCPU target" in text
    tables = Document(str(docx)).tables
    assert any(t.rows[0].cells[0].text == "Cluster" and
               any(c.text == "vCPU:pCore N-1" for c in t.rows[0].cells) for t in tables)
    pdf = export_environment(env, "pdf", tmp_path, "cap", opts)
    assert pdf.stat().st_size > 10_000


def test_sidecar_passes_the_ratio_through(demo: Any, data_dir: Path) -> None:
    """A collection records each cluster's target; the detail pane and re-exports use the
    recorded one whatever ratio they are passed later (the current setting)."""
    run = sc._m_inventory(None, {
        "host": "pc02.demo.local", "username": "admin", "password": "demo", "port": 9440,
        "insecure": False, "name": "cap", "formats": [], "export_dir": str(data_dir)})
    path = run["snapshot_path"]
    base = {"snapshot": path, "object_type": "clusters", "page": {"limit": 10}}
    rid = sc._m_query_inventory(None, base)["rows"][0]["id"]
    d4 = sc._m_get_object_detail(None, {**base, "id": rid})
    d8 = sc._m_get_object_detail(None, {**base, "id": rid, "vcpu_ratio": 8})
    target = {"label": "Target ratio", "value": "4.0 : 1 (default)", "mono": True}
    assert target in d4["fields"] and target in d8["fields"]
    out = sc._m_export(None, {"snapshot": path, "formats": ["excel"], "vcpu_ratio": 3,
                              "export_dir": str(data_dir / "x")})
    wb = load_workbook(out["paths"]["excel"])
    assert wb["Capacity"]["H2"].value == "4.0 : 1 (default)"


def _robo_uuid() -> str:
    env = collect_environment(PrismCredentials(host="pc02.demo.local", username="admin",
                                               password="demo"))
    return next(c.uuid for c in env.clusters if c.name == "SEA-ROBO-01")


def test_per_cluster_targets_are_recorded_and_used(demo: Any, data_dir: Path) -> None:
    robo_uuid = _robo_uuid()
    run = sc._m_inventory(None, {
        "host": "pc02.demo.local", "username": "admin", "password": "demo", "port": 9440,
        "insecure": False, "name": "custom", "formats": ["excel", "word"],
        "export_dir": str(data_dir / "out"), "vcpu_ratio": 5,  # Settings value
        "vcpu_ratio_overrides": {robo_uuid: 6.5}})
    path = run["snapshot_path"]
    env = sc._cache.get(path)
    assert env.vcpu_ratio_default == 5.0
    by_name = {c.name: c for c in env.clusters}
    assert by_name["SEA-ROBO-01"].vcpu_ratio_target == 6.5
    assert by_name["SEA-ROBO-01"].vcpu_ratio_custom is True
    others = [c for c in env.clusters if c.name != "SEA-ROBO-01"]
    assert others and all(c.vcpu_ratio_target == 5.0 and c.vcpu_ratio_custom is False
                          for c in others)

    # Capacity math: the custom target drives vCPU Total (24 usable cores x 6.5).
    caps = {k.cluster_name: k for k in compute_environment(env, 9).values()}
    robo = caps["SEA-ROBO-01"]
    assert robo.target_ratio == 6.5 and robo.target_custom
    assert robo.vcpu_total == robo.usable_cores * 6.5
    assert all(caps[c.name].target_ratio == 5.0 for c in others)

    # Detail pane, whatever the current setting.
    base = {"snapshot": path, "object_type": "clusters", "page": {"limit": 50}}
    rows = sc._m_query_inventory(None, base)["rows"]
    rid = next(r["id"] for r in rows if r["cells"]["name"] == "SEA-ROBO-01")
    detail = sc._m_get_object_detail(None, {**base, "id": rid, "vcpu_ratio": 8})
    fields = {f["label"]: f["value"] for f in detail["fields"]}
    assert fields["Target ratio"] == "6.5 : 1 (custom)"

    # Reports written by the run: Capacity sheet, legend, overview, Word.
    wb = load_workbook(run["exports"]["excel"])
    ws = wb["Capacity"]
    headers = [c.value for c in ws[1]]
    sheet = {r[0]: dict(zip(headers, r, strict=False))
             for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1] is not None}
    assert sheet["SEA-ROBO-01"]["Target vCPU:pCore"] == "6.5 : 1 (custom)"
    assert all(sheet[c.name]["Target vCPU:pCore"] == "5.0 : 1 (default)" for c in others)
    overview = [str(c.value) for row in wb["Overview"].iter_rows() for c in row if c.value]
    assert any(v.startswith("5.0 : 1 vCPU:pCPU target by default and a custom target on "
                            "1 cluster: SEA-ROBO-01 6.5 : 1") for v in overview)
    text = "\n".join(p.text for p in Document(run["exports"]["word"]).paragraphs)
    assert "SEA-ROBO-01 6.5 : 1" in text
    tables = Document(run["exports"]["word"]).tables
    cpu = next(t for t in tables if any(c.text == "vCPU:pCore N-1" for c in t.rows[0].cells))
    heads = [c.text for c in cpu.rows[0].cells]
    assert heads[0] == "Cluster" and "Target" in heads
    assert any(r.cells[0].text.replace("\u2011", "-") == "SEA-ROBO-01" and
               r.cells[heads.index("Target")].text == "6.5 : 1 (custom)" for r in cpu.rows)

    # A re-export passed another ratio keeps the recorded targets.
    out = sc._m_export(None, {"snapshot": path, "formats": ["excel"], "vcpu_ratio": 3,
                              "export_dir": str(data_dir / "x")})
    ws = load_workbook(out["paths"]["excel"])["Capacity"]
    targets = {r[0]: r[headers.index("Target vCPU:pCore")]
               for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1] is not None}
    assert targets["SEA-ROBO-01"] == "6.5 : 1 (custom)"


def test_bad_overrides_are_refused_before_collecting(demo: Any, data_dir: Path) -> None:
    import pytest

    from ntnxvision_core.api.errors import NutanixApiError

    for bad in ({"x": 0}, {"x": "six"}, {"x": 100}, ["x"]):
        with pytest.raises(NutanixApiError):
            sc._m_inventory(None, {
                "host": "pc02.demo.local", "username": "admin", "password": "demo",
                "name": f"bad-{len(str(bad))}", "formats": [],
                "export_dir": str(data_dir), "vcpu_ratio_overrides": bad})


def test_old_snapshots_fall_back_to_the_setting(demo: Any) -> None:
    """A snapshot without recorded targets uses the ratio passed in (the setting)."""
    env = collect_environment(PrismCredentials(host="pc02.demo.local", username="admin",
                                               password="demo"))
    assert env.vcpu_ratio_default is None
    assert all(c.vcpu_ratio_target is None for c in env.clusters)
    caps = compute_environment(env, 7).values()
    assert all(k.target_ratio == 7.0 and not k.target_custom for k in caps)


def test_prism_element_override_matches_the_listed_cluster(demo: Any, data_dir: Path) -> None:
    """Prism Element mode: the uuid New collection lists is the one the run records."""
    from ntnxvision_core.clusterlist import list_clusters

    listed = list_clusters({"mode": "prism_element", "host": "pe01.demo.local",
                            "username": "admin", "password": "demo", "port": 9440})
    (uuid,) = [c["uuid"] for c in listed["clusters"]]
    run = sc._m_inventory_cluster(None, {
        "host": "pe01.demo.local", "username": "admin", "password": "demo", "port": 9440,
        "name": "pe-custom", "formats": [], "export_dir": str(data_dir),
        "vcpu_ratio_overrides": {uuid: 3}})
    env = sc._cache.get(run["snapshot_path"])
    (cluster,) = env.clusters
    assert cluster.uuid == uuid
    assert (cluster.vcpu_ratio_target, cluster.vcpu_ratio_custom) == (3.0, True)


def test_multi_pc_targets_carry_their_own_overrides(demo: Any, data_dir: Path) -> None:
    robo_uuid = _robo_uuid()
    out = sc._m_inventory_multi(None, {
        "targets": [
            {"host": "pc01.demo.local", "username": "admin", "password": "demo",
             "name": "m-pc01"},
            {"host": "pc02.demo.local", "username": "admin", "password": "demo",
             "name": "m-pc02", "vcpu_ratio_overrides": {robo_uuid: 7}},
        ],
        "formats": [], "export_dir": str(data_dir), "combine": True,
        "combined_name": "m-both", "vcpu_ratio": 5, "parallel": False})
    combined = sc._cache.get(out["combined"]["snapshot_path"])
    assert combined.vcpu_ratio_default == 5.0
    targets = {(c.source_pc, c.name): (c.vcpu_ratio_target, c.vcpu_ratio_custom)
               for c in combined.clusters}
    assert targets[("pc02.demo.local", "SEA-ROBO-01")] == (7.0, True)
    assert all(v == (5.0, False) for k, v in targets.items()
               if k != ("pc02.demo.local", "SEA-ROBO-01"))
