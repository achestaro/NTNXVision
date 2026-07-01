"""Smoke tests that exercise the model → snapshot → Excel pipeline without a live PC."""

from __future__ import annotations

from openpyxl import load_workbook

from ntnxvision_core.exporters import ExcelExporter
from ntnxvision_core.models import (
    Cluster,
    Environment,
    Host,
    Nic,
    StorageContainer,
    Subnet,
    VirtualDisk,
    VirtualMachine,
)
from ntnxvision_core.store import SnapshotStore

GIB = 1024**3


def _sample_env() -> Environment:
    env = Environment(collected_at="2026-06-30T00:00:00Z")
    env.prism_central.host = "pc.lab.local"
    env.prism_central.version = "pc.2024.3"
    env.clusters.append(
        Cluster(uuid="c1", name="cluster-01", aos_version="7.0", hypervisor_types=["AHV"],
                num_nodes=3, redundancy_factor=2)
    )
    env.hosts.append(
        Host(uuid="h1", name="node-a", cluster_uuid="c1", cluster_name="cluster-01",
             hypervisor_type="AHV", num_cpu_cores=32, memory_bytes=256 * GIB)
    )
    vm = VirtualMachine(uuid="v1", name="app-01", cluster_name="cluster-01", power_state="ON",
                        num_vcpus=4, num_cores_per_vcpu=1, memory_bytes=16 * GIB)
    vm.disks.append(VirtualDisk(label="scsi.0", disk_size_bytes=100 * GIB, bus_type="SCSI"))
    vm.nics.append(Nic(mac_address="50:6b:00:00:00:01", ip_addresses=["10.0.0.10"]))
    env.vms.append(vm)
    env.storage_containers.append(
        StorageContainer(uuid="s1", name="default", cluster_name="cluster-01",
                         replication_factor=2, compression_enabled=True,
                         max_capacity_bytes=10 * 1024 * GIB)
    )
    env.subnets.append(
        Subnet(uuid="net1", name="vlan-100", cluster_name="cluster-01",
               subnet_type="VLAN", vlan_id=100)
    )
    env.collection_notes["vms"] = "ok"
    return env


def test_summary_counts() -> None:
    env = _sample_env()
    s = env.summary()
    assert s["clusters"] == 1
    assert s["hosts"] == 1
    assert s["vms"] == 1
    assert s["storage_containers"] == 1
    assert s["subnets"] == 1


def test_vm_derived_fields() -> None:
    vm = _sample_env().vms[0]
    assert vm.vcpu_total == 4
    assert vm.total_disk_bytes == 100 * GIB


def test_excel_export_creates_expected_sheets(tmp_path) -> None:
    out = tmp_path / "env.xlsx"
    ExcelExporter().export(_sample_env(), out)
    assert out.exists()
    wb = load_workbook(out)
    for sheet in ["Overview", "Clusters", "Hosts", "VMs", "Storage", "Subnets"]:
        assert sheet in wb.sheetnames
    vms = wb["VMs"]
    assert vms["A1"].value == "Name"
    assert vms.freeze_panes == "A2"


def test_snapshot_roundtrip(tmp_path) -> None:
    store = SnapshotStore(tmp_path)
    path = store.save(_sample_env(), "test-env")
    loaded = store.load(path)
    assert loaded.summary() == _sample_env().summary()
    assert loaded.vms[0].name == "app-01"


# ─── Phase 2 coverage ──────────────────────────────────────────────────


def _phase2_env() -> Environment:
    from ntnxvision_core.models import (
        Database,
        FileServer,
        ProtectionPolicy,
        RecoveryPlan,
        SecurityPolicy,
        User,
        Vpc,
    )

    env = _sample_env()
    env.protection_policies.append(
        ProtectionPolicy(uuid="pp1", name="gold", rpo_seconds=3600, replication_type="ASYNC")
    )
    env.recovery_plans.append(RecoveryPlan(uuid="rp1", name="dr-plan", stage_count=2))
    env.security_policies.append(
        SecurityPolicy(uuid="sp1", name="web-tier", policy_type="APPLICATION", state="ENFORCE")
    )
    env.vpcs.append(Vpc(uuid="vpc1", name="prod-vpc"))
    env.file_servers.append(FileServer(uuid="fs1", name="files-01", num_fsvm=3))
    env.databases.append(Database(id="db1", name="ORCL", db_engine="oracle", status="READY"))
    env.users.append(User(uuid="u1", username="admin", user_type="LOCAL"))
    return env


def test_phase2_summary_and_sheets(tmp_path) -> None:
    env = _phase2_env()
    s = env.summary()
    assert s["protection_policies"] == 1
    assert s["security_policies"] == 1
    assert s["vpcs"] == 1
    assert s["file_servers"] == 1
    assert s["databases"] == 1

    out = tmp_path / "p2.xlsx"
    ExcelExporter().export(env, out)
    wb = load_workbook(out)
    for sheet in [
        "Protection Policies", "Recovery Plans", "Flow Security Policies", "VPCs",
        "File Servers", "Databases", "Users", "Feature Usage",
    ]:
        assert sheet in wb.sheetnames


def test_feature_usage_analysis() -> None:
    from ntnxvision_core.analysis import analyze_feature_usage

    rows = analyze_feature_usage(_phase2_env())
    by_feature = {r.feature: r for r in rows}
    assert by_feature["Async DR replication"].in_use is True
    assert by_feature["Microsegmentation policies"].in_use is True
    assert by_feature["File servers"].in_use is True
    assert by_feature["Managed databases"].in_use is True
    # Something not configured in the sample should be flagged not-in-use.
    assert by_feature["Metro Availability"].in_use is False


# ─── Phase 3: PDF + Word exporters + dispatcher ────────────────────────


def test_export_dispatcher_all_formats(tmp_path) -> None:
    from ntnxvision_core.exporters import export_environment

    env = _phase2_env()
    xlsx = export_environment(env, "excel", tmp_path, "env")
    pdf = export_environment(env, "pdf", tmp_path, "env")
    docx = export_environment(env, "word", tmp_path, "env")

    assert xlsx.exists() and xlsx.suffix == ".xlsx"
    assert pdf.exists() and pdf.suffix == ".pdf"
    assert docx.exists() and docx.suffix == ".docx"
    # Non-trivial output sizes (files actually got written).
    assert pdf.stat().st_size > 1000
    assert docx.stat().st_size > 1000


def test_export_dispatcher_rejects_unknown_format(tmp_path) -> None:
    import pytest

    from ntnxvision_core.exporters import export_environment

    with pytest.raises(ValueError):
        export_environment(_sample_env(), "csv", tmp_path, "env")


def test_word_doc_has_headings(tmp_path) -> None:
    from docx import Document

    from ntnxvision_core.exporters.word import WordExporter

    out = WordExporter().export(_phase2_env(), tmp_path / "add.docx")
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Executive Summary" in h for h in headings)
    assert any("Data Protection" in h for h in headings)
    assert any("Licensing" in h for h in headings)
