"""Multi-Prism-Central collection: combine_environments() and the sidecar's
inventory_multi orchestration (per-target isolation + optional combine).
"""

from __future__ import annotations

from typing import Any
from unittest.mock import patch

import pytest

from ntnxvision_core.engine import combine_environments
from ntnxvision_core.models import Cluster, Environment, PrismCentral, VirtualMachine


def _env(host: str, cluster_uuid: str) -> Environment:
    env = Environment(prism_central=PrismCentral(host=host))
    env.clusters = [Cluster(uuid=cluster_uuid, name=f"clu-{host}")]
    env.vms = [VirtualMachine(uuid=f"vm-{host}", name=f"vm-{host}", cluster_name=f"clu-{host}")]
    env.warnings = [f"warning from {host}"]
    env.collection_notes = {"clusters": "ok"}
    env.flow.security_policy_count = 1
    return env


def test_combine_environments_merges_lists_and_sums_summaries() -> None:
    e1, e2 = _env("pc1.local", "c1"), _env("pc2.local", "c2")
    combined = combine_environments([e1, e2])

    assert [c.name for c in combined.clusters] == ["clu-pc1.local", "clu-pc2.local"]
    assert [v.name for v in combined.vms] == ["vm-pc1.local", "vm-pc2.local"]
    assert combined.warnings == ["warning from pc1.local", "warning from pc2.local"]
    assert combined.flow.security_policy_count == 2
    assert combined.collection_notes == {
        "pc1.local: clusters": "ok",
        "pc2.local: clusters": "ok",
    }
    assert [s.prism_central.host for s in combined.source_environments] == [
        "pc1.local", "pc2.local",
    ]
    assert combined.summary()["source_environments"] == 2


def test_combine_environments_single_is_identity() -> None:
    e1 = _env("pc1.local", "c1")
    assert combine_environments([e1]) is e1


def test_combine_environments_empty_raises() -> None:
    with pytest.raises(ValueError, match="at least one"):
        combine_environments([])


def test_inventory_multi_isolates_a_failing_target() -> None:
    """One target raising (bad creds, unreachable host, ...) must not prevent the
    other targets from collecting and being finalized normally.
    """
    from ntnxvision_core import sidecar

    def fake_collect(creds: Any, progress: Any = None, ndb_creds: Any = None,
                     **_kw: Any) -> Environment:
        if creds.host == "bad.local":
            raise RuntimeError("connection refused")
        return _env(creds.host, f"c-{creds.host}")

    with patch.object(sidecar, "collect_environment", side_effect=fake_collect), \
         patch.object(sidecar, "_finalize", side_effect=lambda env, name, params, **_kw: {
             "summary": env.summary(), "source_type": env.source_type,
             "prism_central": env.prism_central.model_dump(), "warnings": [],
             "collection_notes": {}, "snapshot_path": None, "exports": {}, "export_errors": {},
         }):
        result = sidecar._m_inventory_multi(1, {
            "targets": [
                {"host": "good1.local", "username": "admin", "password": "x"},
                {"host": "bad.local", "username": "admin", "password": "x"},
                {"host": "good2.local", "username": "admin", "password": "x"},
            ],
            "combine": False,
        })

    outcomes = {t["host"]: t["success"] for t in result["targets"]}
    assert outcomes == {"good1.local": True, "bad.local": False, "good2.local": True}
    bad = next(t for t in result["targets"] if t["host"] == "bad.local")
    assert "connection refused" in bad["error"]
    assert result["combined"] is None


def test_inventory_multi_combines_when_requested() -> None:
    from ntnxvision_core import sidecar

    def fake_collect(creds: Any, progress: Any = None, ndb_creds: Any = None,
                     **_kw: Any) -> Environment:
        return _env(creds.host, f"c-{creds.host}")

    finalize_calls: list[str] = []

    def fake_finalize(env: Environment, name: str, params: Any, **_kw: Any) -> dict[str, Any]:
        finalize_calls.append(name)
        return {
            "summary": env.summary(), "source_type": env.source_type,
            "prism_central": env.prism_central.model_dump(), "warnings": [],
            "collection_notes": {}, "snapshot_path": None, "exports": {}, "export_errors": {},
        }

    with patch.object(sidecar, "collect_environment", side_effect=fake_collect), \
         patch.object(sidecar, "_finalize", side_effect=fake_finalize):
        result = sidecar._m_inventory_multi(1, {
            "targets": [
                {"host": "pc1.local", "username": "admin", "password": "x"},
                {"host": "pc2.local", "username": "admin", "password": "x"},
            ],
            "combine": True,
            "combined_name": "acme-combined",
        })

    assert all(t["success"] for t in result["targets"])
    assert result["combined"] is not None
    assert result["combined"]["summary"]["source_environments"] == 2
    assert "acme-combined" in finalize_calls


def test_inventory_multi_skips_combine_with_fewer_than_two_successes() -> None:
    from ntnxvision_core import sidecar

    def fake_collect(creds: Any, progress: Any = None, ndb_creds: Any = None,
                     **_kw: Any) -> Environment:
        if creds.host == "bad.local":
            raise RuntimeError("nope")
        return _env(creds.host, "c1")

    with patch.object(sidecar, "collect_environment", side_effect=fake_collect), \
         patch.object(sidecar, "_finalize", side_effect=lambda env, name, params, **_kw: {
             "summary": env.summary(), "source_type": env.source_type,
             "prism_central": env.prism_central.model_dump(), "warnings": [],
             "collection_notes": {}, "snapshot_path": None, "exports": {}, "export_errors": {},
         }):
        result = sidecar._m_inventory_multi(1, {
            "targets": [
                {"host": "good.local", "username": "admin", "password": "x"},
                {"host": "bad.local", "username": "admin", "password": "x"},
            ],
            "combine": True,
        })

    assert result["combined"] is None


def test_inventory_multi_requires_at_least_one_target() -> None:
    from ntnxvision_core import sidecar
    from ntnxvision_core.api.errors import NutanixApiError

    with pytest.raises(NutanixApiError):
        sidecar._m_inventory_multi(1, {"targets": []})


# ─── combined exports ──────────────────────────────────────────────────
# A combined snapshot exported fine before these tests existed — it just described
# itself wrongly (one PC's sizing presented as the whole estate's, sources never
# named, identically-named clusters indistinguishable). These assert the report, not
# just the absence of an exception.


def _sizable_env(host: str, version: str, nodes: int, cluster_name: str) -> Environment:
    """A PC with its own sizing and a cluster named the same on every PC — the
    collision that makes source attribution necessary.
    """
    from ntnxvision_core.models import Host as HostModel
    from ntnxvision_core.models import PrismCentralConfig

    env = Environment(
        collected_at="2026-09-05T00:00:00Z",
        prism_central=PrismCentral(host=host, version=version),
    )
    env.pc_config = PrismCentralConfig(
        version=version, node_count=nodes, is_scale_out=nodes > 1,
        registered_cluster_count=1,
    )
    env.clusters = [Cluster(uuid=f"c-{host}", name=cluster_name, num_nodes=3,
                            hypervisor_types=["AHV"])]
    env.hosts = [HostModel(uuid=f"h-{host}", name=f"node-{host}", cluster_name=cluster_name)]
    env.vms = [VirtualMachine(uuid=f"v-{host}", name=f"vm-{host}", cluster_name=cluster_name)]
    return env


def _combined_pair() -> Environment:
    return combine_environments([
        _sizable_env("pc-emea.local", "pc.2024.3.1.8", 1, "Cluster01"),
        _sizable_env("pc-amer.local", "pc.2024.3.2", 3, "Cluster01"),
    ])


def test_combine_stamps_every_merged_object_with_its_source_pc() -> None:
    combined = _combined_pair()
    assert [c.source_pc for c in combined.clusters] == ["pc-emea.local", "pc-amer.local"]
    assert [h.source_pc for h in combined.hosts] == ["pc-emea.local", "pc-amer.local"]
    assert [v.source_pc for v in combined.vms] == ["pc-emea.local", "pc-amer.local"]
    assert combined.is_combined


def test_combined_keeps_each_pcs_sizing_instead_of_the_first_ones() -> None:
    """The combined environment must not adopt one PC's management-plane config —
    reports rendered it as the whole estate's sizing.
    """
    combined = _combined_pair()
    assert combined.pc_config.node_count is None
    sizing = [(s.prism_central.host, s.pc_config.node_count)
              for s in combined.source_environments]
    assert sizing == [("pc-emea.local", 1), ("pc-amer.local", 3)]
    assert combined.source_environments[0].summary["clusters"] == 1


def test_single_pc_environment_is_not_treated_as_combined() -> None:
    env = _sizable_env("pc1.local", "pc.2024.3", 1, "Cluster01")
    assert not env.is_combined
    assert env.clusters[0].source_pc is None


def test_combined_excel_names_the_sources_and_attributes_every_row(tmp_path) -> None:
    from openpyxl import load_workbook

    from ntnxvision_core.exporters import ExcelExporter

    path = ExcelExporter().export(_combined_pair(), tmp_path / "combined.xlsx")
    wb = load_workbook(path)

    overview = {r[0]: r[1] for r in wb["Overview"].iter_rows(values_only=True) if r[0]}
    assert overview["Prism Centrals"] == "2 (combined)"
    assert overview["Sources"] == "pc-emea.local, pc-amer.local"
    # The single-PC property sheet — which would have to pick one PC's sizing — is
    # replaced by a row per source.
    assert "Prism Central" not in wb.sheetnames
    pcs = list(wb["Prism Centrals"].iter_rows(values_only=True))
    assert pcs[0][:4] == ("Prism Central", "Version", "Deployment", "PC Nodes")
    assert [(r[0], r[1], r[3]) for r in pcs[1:]] == [
        ("pc-emea.local", "pc.2024.3.1.8", 1),
        ("pc-amer.local", "pc.2024.3.2", 3),
    ]

    for sheet, expected in (("Clusters", "Cluster01"), ("Hosts", "node-pc-emea.local")):
        rows = list(wb[sheet].iter_rows(values_only=True))
        assert rows[0][0] == "Prism Central", f"{sheet} lacks the attribution column"
        assert rows[1][0] == "pc-emea.local"
        assert rows[2][0] == "pc-amer.local"
        assert rows[1][1] == expected
    assert list(wb["VMs"].iter_rows(values_only=True))[0][0] == "Prism Central"


def test_single_pc_excel_keeps_its_original_shape(tmp_path) -> None:
    """The attribution column and the source sheet are combined-only — a normal run's
    workbook must be byte-for-byte the same shape it has always been.
    """
    from openpyxl import load_workbook

    from ntnxvision_core.exporters import ExcelExporter

    env = _sizable_env("pc1.local", "pc.2024.3", 1, "Cluster01")
    wb = load_workbook(ExcelExporter().export(env, tmp_path / "solo.xlsx"))

    assert "Prism Centrals" not in wb.sheetnames
    assert "Prism Central" in wb.sheetnames          # the single-PC property sheet
    overview = {r[0]: r[1] for r in wb["Overview"].iter_rows(values_only=True) if r[0]}
    assert overview["Prism Central"] == "pc1.local"
    assert overview["PC Version"] == "pc.2024.3"
    assert list(wb["Clusters"].iter_rows(values_only=True))[0][0] == "Name"


def test_combined_word_names_the_sources_and_attributes_clusters(tmp_path) -> None:
    from docx import Document

    from ntnxvision_core.exporters.word import WordExporter

    path = WordExporter().export(_combined_pair(), tmp_path / "combined.docx")
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)

    assert "spanning 2 Prism Centrals (pc-emea.local, pc-amer.local)" in text
    assert "Prism Centrals — Management & Sizing" in text
    # The misleading single-PC sizing table must be gone, not merely supplemented.
    assert "Prism Central management-plane configuration" not in text

    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    source_table = next(t for t in tables if t[0][:3] == ["Prism Central", "Version", "Deployment"])
    assert [r[0] for r in source_table[1:]] == ["pc-emea.local", "pc-amer.local"]
    assert [r[3] for r in source_table[1:]] == ["1", "3"]        # per-PC node counts

    cluster_table = next(t for t in tables if t[0] == ["Prism Central", "Cluster", "AOS",
                                                       "Hypervisor", "Nodes", "RF", "Encryption"])
    assert [r[0] for r in cluster_table[1:]] == ["pc-emea.local", "pc-amer.local"]
    assert {r[1] for r in cluster_table[1:]} == {"Cluster01"}    # same name, distinguishable


def test_combined_pdf_exports_with_the_extra_column(tmp_path) -> None:
    """The PDF's column widths are tuned to the landscape page, so the attribution
    column has to fit rather than overflow it.
    """
    from ntnxvision_core.exporters.pdf import PdfExporter

    path = PdfExporter().export(_combined_pair(), tmp_path / "combined.pdf")
    assert path.exists() and path.stat().st_size > 10_000


def test_legacy_combined_snapshot_still_names_its_sources(tmp_path) -> None:
    """`source_environments` used to be a bare list of PrismCentral. Snapshots in that
    shape must still read back with their PCs named, not as "unknown".
    """
    import json

    from ntnxvision_core.store import SnapshotStore

    (tmp_path / "old.ntnxsnap.json").write_text(json.dumps({
        "collected_at": "2026-08-25T00:00:00Z",
        "prism_central": {"host": "2 Prism Centrals (combined)"},
        "source_environments": [
            {"host": "pc1.local", "version": "pc.2024.3"},
            {"host": "pc2.local", "version": "pc.2024.3"},
        ],
        "clusters": [{"uuid": "c1", "name": "Cluster01"}],
    }))
    env = SnapshotStore(tmp_path).load(tmp_path / "old.ntnxsnap.json")

    assert env.is_combined
    assert [s.prism_central.host for s in env.source_environments] == ["pc1.local", "pc2.local"]
    assert [s.prism_central.version for s in env.source_environments] == ["pc.2024.3"] * 2


# ─── figures ───────────────────────────────────────────────────────────


def _pc_env(host: str, hosts: int, vms: int, subnets: int = 0) -> Environment:
    """A PC whose cluster is called "Cluster01", like everyone else's."""
    from ntnxvision_core.models import Host as HostModel
    from ntnxvision_core.models import Subnet

    env = Environment(prism_central=PrismCentral(host=host))
    env.clusters = [Cluster(uuid=f"c-{host}", name="Cluster01", num_nodes=hosts,
                            hypervisor_types=["AHV"])]
    # Hosts carry the cluster UUID; VMs and subnets only the name — both resolution
    # paths have to stay inside their own Prism Central.
    env.hosts = [HostModel(uuid=f"h-{host}-{i}", name=f"node-{host}-{i}",
                           cluster_uuid=f"c-{host}", cluster_name="Cluster01")
                 for i in range(hosts)]
    env.vms = [VirtualMachine(uuid=f"v-{host}-{i}", name=f"vm-{host}-{i}",
                              cluster_name="Cluster01") for i in range(vms)]
    env.subnets = [Subnet(uuid=f"s-{host}-{i}", name=f"vlan-{host}-{i}",
                          cluster_name="Cluster01", vlan_id=100 + i)
                   for i in range(subnets)]
    return env


def test_identically_named_clusters_do_not_pool_their_objects() -> None:
    """Two PCs each running a "Cluster01" used to be matched by name alone, pooling one
    estate's hosts, VMs and subnets into the other's figures.
    """
    from ntnxvision_core.exporters.diagrams import _counts_per_cluster

    combined = combine_environments([_pc_env("pc1.local", hosts=2, vms=3, subnets=1),
                                     _pc_env("pc2.local", hosts=1, vms=1)])
    c1, c2 = combined.clusters

    assert _counts_per_cluster(combined, combined.vms) == {id(c1): 3, id(c2): 1}
    assert _counts_per_cluster(combined, combined.hosts) == {id(c1): 2, id(c2): 1}
    assert _counts_per_cluster(combined, combined.subnets) == {id(c1): 1}


def test_per_cluster_figures_only_see_their_own_clusters_objects() -> None:
    from ntnxvision_core.exporters.diagrams import cluster_compute, network_topology

    combined = combine_environments([_pc_env("pc1.local", hosts=2, vms=3, subnets=1),
                                     _pc_env("pc2.local", hosts=0, vms=1)])
    c1, c2 = combined.clusters

    assert cluster_compute(combined, c1) is not None
    # pc2's cluster has no nodes of its own — it must not borrow pc1's.
    assert cluster_compute(combined, c2) is None
    assert network_topology(combined, c1) is not None
    assert network_topology(combined, c2) is None


def test_cluster_labels_disambiguate_only_when_names_collide() -> None:
    from ntnxvision_core.exporters.diagrams import _cluster_labels

    combined = combine_environments([_pc_env("pc1.local", hosts=1, vms=1),
                                     _pc_env("pc2.local", hosts=1, vms=1)])
    assert [_cluster_labels(combined)[id(c)] for c in combined.clusters] == [
        "Cluster01 · pc1.local", "Cluster01 · pc2.local",
    ]

    solo = _pc_env("pc1.local", hosts=1, vms=1)
    assert [_cluster_labels(solo)[id(c)] for c in solo.clusters] == ["Cluster01"]


# ─── CLI ───────────────────────────────────────────────────────────────


def _run_cli(args: list[str], collect: Any, tmp_path: Any) -> Any:
    """Invoke the CLI with collection and the snapshot/export dirs stubbed out."""
    from typer.testing import CliRunner

    from ntnxvision_core import __main__ as cli

    with patch.object(cli, "collect_environment", collect), \
            patch.object(cli, "snapshots_dir", lambda: tmp_path), \
            patch.object(cli, "exports_dir", lambda: tmp_path):
        return CliRunner().invoke(cli.app, args)


def test_cli_inventory_multi_collects_every_target_and_combines(tmp_path) -> None:
    seen = []

    def collect(creds: Any, progress: Any = None, ndb_creds: Any = None) -> Environment:
        seen.append((creds.host, creds.username, creds.port))
        return _env(creds.host, f"c-{creds.host}")

    result = _run_cli(
        ["inventory-multi", "--host", "pc1.local", "--host", "pc2.local",
         "--username", "admin", "--password", "pw", "--combined-name", "estate"],
        collect, tmp_path,
    )

    assert result.exit_code == 0, result.output
    # One --username / --password applies to every --host.
    assert seen == [("pc1.local", "admin", 9440), ("pc2.local", "admin", 9440)]
    # Per-target snapshots plus the combined one.
    names = sorted(p.name for p in tmp_path.glob("*.ntnxsnap.json"))
    assert names == ["estate.ntnxsnap.json", "pc1.local.ntnxsnap.json", "pc2.local.ntnxsnap.json"]


def test_cli_inventory_multi_isolates_a_failing_target(tmp_path) -> None:
    def collect(creds: Any, progress: Any = None, ndb_creds: Any = None) -> Environment:
        if creds.host == "bad.local":
            raise RuntimeError("connection refused")
        return _env(creds.host, f"c-{creds.host}")

    result = _run_cli(
        ["inventory-multi", "--host", "bad.local", "--host", "good.local",
         "--username", "admin", "--password", "pw"],
        collect, tmp_path,
    )

    # The good target is still collected and exported; the run reports the failure.
    assert "connection refused" in result.output
    assert (tmp_path / "good.local.ntnxsnap.json").exists()
    assert result.exit_code == 1
    # One success is not enough to combine.
    assert not (tmp_path / "combined-environment.ntnxsnap.json").exists()


def test_cli_inventory_multi_rejects_a_mismatched_option_count(tmp_path) -> None:
    """A miscounted --username is a typo worth catching before a long run, not after."""
    def collect(creds: Any, progress: Any = None, ndb_creds: Any = None) -> Environment:
        raise AssertionError("must not start collecting")

    result = _run_cli(
        ["inventory-multi", "--host", "a.local", "--host", "b.local", "--host", "c.local",
         "--username", "admin", "--username", "root", "--password", "pw"],
        collect, tmp_path,
    )
    assert result.exit_code != 0
    assert "once per host" in result.output.replace("\n", " ")


def test_cli_inventory_multi_passes_per_target_ndb_credentials(tmp_path) -> None:
    ndb_seen = []

    def collect(creds: Any, progress: Any = None, ndb_creds: Any = None) -> Environment:
        ndb_seen.append(None if ndb_creds is None else (ndb_creds.host, ndb_creds.username))
        return _env(creds.host, f"c-{creds.host}")

    result = _run_cli(
        ["inventory-multi", "--host", "pc1.local", "--host", "pc2.local",
         "--username", "admin", "--password", "pw",
         "--ndb-host", "ndb1.local", "--ndb-host", "", "--ndb-username", "era",
         "--ndb-password", "era-pw"],
        collect, tmp_path,
    )
    assert result.exit_code == 0, result.output
    assert ndb_seen == [("ndb1.local", "era"), None]


def test_inventory_multi_uses_each_targets_own_ndb_credentials() -> None:
    """The sidecar has always read a per-target `ndb` block, but the IPC contract had no
    field for it, so the UI could never send one and the path was dead.
    """
    from ntnxvision_core import sidecar

    seen: list[tuple[str, Any]] = []

    def fake_collect(creds: Any, progress: Any = None, ndb_creds: Any = None,
                     **_kw: Any) -> Environment:
        seen.append((creds.host, None if ndb_creds is None
                     else (ndb_creds.host, ndb_creds.username, ndb_creds.verify_ssl)))
        return _env(creds.host, f"c-{creds.host}")

    with patch.object(sidecar, "collect_environment", side_effect=fake_collect), \
         patch.object(sidecar, "_finalize", side_effect=lambda env, name, params, **_kw: {
             "summary": env.summary(), "source_type": env.source_type,
             "prism_central": env.prism_central.model_dump(), "warnings": [],
             "collection_notes": {}, "snapshot_path": None, "exports": {}, "export_errors": {},
         }):
        sidecar._m_inventory_multi(1, {
            "targets": [
                {"host": "pc1.local", "username": "admin", "password": "x", "insecure": True,
                 "ndb": {"host": "era1.local", "username": "era-admin", "password": "y"}},
                {"host": "pc2.local", "username": "admin", "password": "x"},
            ],
            "combine": False,
        })

    # Targets collect in parallel, so compare without depending on call order.
    assert sorted(seen, key=lambda s: s[0]) == [
        # Each estate's own Era instance, and the target's TLS choice carried across.
        ("pc1.local", ("era1.local", "era-admin", False)),
        ("pc2.local", None),
    ]
