"""Comma thousands grouping for every number the core shows to users."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ntnxvision_core.exporters import ExcelExporter
from ntnxvision_core.exporters.tables import SPECS
from ntnxvision_core.inventory import _facts, display, fmt_bytes
from ntnxvision_core.models import Cluster, Environment, Host, Subnet, VirtualMachine
from ntnxvision_core.numfmt import cell_text, excel_number_format, fmt_count
from ntnxvision_core.runinfo import _join_labels


def test_fmt_count_uses_comma_grouping() -> None:
    assert fmt_count(0) == "0"
    assert fmt_count(999) == "999"
    assert fmt_count(1214) == "1,214"
    assert fmt_count(12000) == "12,000"
    assert fmt_count(1234567) == "1,234,567"
    assert fmt_count(1234.5) == "1,234.5"
    # Non-numbers pass through; None is a dash, never "None".
    assert fmt_count("—") == "—"
    assert fmt_count(None) == "—"
    assert fmt_count(True) == "True"


def test_cell_text_leaves_identifier_columns_alone() -> None:
    assert cell_text(4096, "VMs") == "4,096"
    assert cell_text(1024, "VLAN") == "1024"
    assert cell_text(None, "VMs") == ""
    assert excel_number_format(1214) == "#,##0"
    assert excel_number_format(12.5) == "#,##0.00"
    assert excel_number_format(True) is None and excel_number_format("x") is None


def test_core_detail_strings_use_commas() -> None:
    env = Environment()
    env.clusters.append(Cluster(uuid="c1", name="c1"))
    env.vms.extend(VirtualMachine(uuid=f"v{i}", name=f"vm-{i}", power_state="ON")
                   for i in range(1214))
    assert _facts(SPECS["vms"], env, len(env.vms)) == "1,214 VMs · 1 cluster · 1,214 powered on"
    assert _join_labels([f"area-{i}" for i in range(1010)]).endswith(" + 1,002 more")
    assert fmt_bytes(1023 * 1024**3) == "1,023 GB"
    vlan = next(c for c in SPECS["subnets"].cols if c.key == "vlan")
    assert display(vlan, 2048) == "2048"
    hosts = next(c for c in SPECS["clusters"].cols if c.key == "hosts")
    assert display(hosts, 1500) == "1,500"


def test_excel_quantity_cells_get_a_grouped_number_format(tmp_path: Path) -> None:
    env = Environment(collected_at="2026-06-30T00:00:00Z")
    env.clusters.append(Cluster(uuid="c1", name="c1", num_nodes=3))
    env.hosts.append(Host(uuid="h1", name="h1", cluster_uuid="c1", num_cpu_cores=32,
                          num_vms=1500))
    env.subnets.append(Subnet(uuid="s1", name="vlan-2048", vlan_id=2048))
    out = ExcelExporter().export(env, tmp_path / "inv.xlsx")
    wb = load_workbook(out)

    def cell(sheet: str, header: str) -> object:
        ws = wb[sheet]
        col = [c.value for c in ws[1]].index(header) + 1
        return ws.cell(row=2, column=col)

    vms = cell("Hosts", "VMs")
    assert vms.value == 1500 and vms.number_format == "#,##0"  # type: ignore[attr-defined]
    vlan = cell("Subnets", "VLAN")
    assert vlan.value == 2048 and vlan.number_format == "General"  # type: ignore[attr-defined]
