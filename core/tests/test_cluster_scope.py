"""Required snapshot names and cluster-scoped collection.

* A collection's snapshot name is typed by the user: required, sanitized like a file
  name, refused when another snapshot already uses it, and never given a date.
* `list_clusters` lists the clusters registered to a Prism Central (v4 and v3), without
  the PC's own pseudo-cluster; Prism Element answers its one cluster.
* A scoped run keeps nothing from unselected clusters, keeps Prism-Central-wide data,
  filters server-side where v4 allows, and records its scope for the app and reports.
* Compare does not report objects of clusters one snapshot left out as added/removed.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

import ntnxvision_core.sidecar as sc
from ntnxvision_core.api import PrismCredentials
from ntnxvision_core.api.errors import AuthenticationError, NutanixApiError
from ntnxvision_core.clusterlist import list_clusters
from ntnxvision_core.compare import compare, export_comparison
from ntnxvision_core.engine import collect_environment
from ntnxvision_core.models import ClusterScope, Environment, ScopeCluster
from ntnxvision_core.scope import CLUSTER_BOUND_FIELDS, describe
from ntnxvision_core.store import SnapshotStore

PW = {"username": "admin", "password": "demo"}


def _clusters(host: str) -> list[dict[str, Any]]:
    return list_clusters({"host": host, **PW})["clusters"]


def _collect(host: str, clusters: list[str] | None = None, **kw: Any) -> Environment:
    return collect_environment(PrismCredentials(host=host, **PW), clusters=clusters, **kw)


def _inventory(data_dir: Path, name: str, **extra: Any) -> dict[str, Any]:
    return sc._m_inventory(None, {"host": "pc01.demo.local", **PW, "port": 9440,
                                  "insecure": False, "name": name, "formats": [],
                                  "export_dir": str(data_dir / "exports"), **extra})


# ─── snapshot names ────────────────────────────────────────────────────


def test_snapshot_name_is_required(data_dir: Path) -> None:
    for raw in (None, "", "   "):
        with pytest.raises(sc.InvalidName, match="Enter a snapshot name"):
            sc.snapshot_name_for(raw)
    with pytest.raises(sc.InvalidName, match="no usable characters"):
        sc.snapshot_name_for(" ... ")


def test_snapshot_name_is_sanitized_like_a_file_name(data_dir: Path) -> None:
    assert sc.snapshot_name_for("  Q4 baseline  ") == "Q4 baseline"
    assert sc.snapshot_name_for('PHX: "prod"/dr?') == "PHX_ _prod__dr_"
    assert sc.snapshot_name_for("CON") == "CON_"
    long = sc.snapshot_name_for("x" * 400)
    assert 16 <= len(long) <= 120


def test_a_name_already_used_is_refused_before_collecting(demo: Any, data_dir: Path) -> None:
    first = _inventory(data_dir, "Q4 baseline")
    assert first["snapshot_name"] == "Q4 baseline"
    # Ignoring case (both file systems default to case-insensitive).
    with pytest.raises(sc.InvalidName, match="already exists"):
        _inventory(data_dir, "q4 BASELINE")
    with pytest.raises(sc.InvalidName):
        _inventory(data_dir, "")
    # The refused runs saved nothing; ids stay unique.
    listing = sc._m_list_snapshots(None, {})["snapshots"]
    assert [s["name"] for s in listing] == ["Q4 baseline"]


def test_collection_names_carry_no_date_and_reports_one(demo: Any, data_dir: Path) -> None:
    day = date.today().isoformat()
    result = _inventory(data_dir, "Q4 baseline", formats=["excel"])
    assert result["snapshot_name"] == "Q4 baseline"
    assert [Path(r["path"]).name for r in result["reports"]] == [f"Q4 baseline-{day}.xlsx"]
    assert len(re.findall(r"\d{4}-\d{2}-\d{2}", result["reports"][0]["name"])) == 1


def test_store_never_replaces_a_snapshot_saved_meanwhile(tmp_path: Path) -> None:
    store = SnapshotStore(tmp_path)
    a = store.save(Environment(), "estate", unique=True)
    b = store.save(Environment(), "ESTATE", unique=True)
    assert (a.name, b.name) == ("estate.ntnxsnap.json", "ESTATE-2.ntnxsnap.json")
    assert store.load(a).snapshot_id != store.load(b).snapshot_id
    assert store.name_taken("Estate") and not store.name_taken("other")


def test_multi_pc_target_names_are_checked_up_front(demo: Any, data_dir: Path) -> None:
    _inventory(data_dir, "Q4 (pc01)")
    with pytest.raises(sc.InvalidName, match="already exists"):
        sc._m_inventory_multi(None, {"targets": [
            {"host": "pc01.demo.local", **PW, "name": "Q4 (pc01)"},
            {"host": "pc02.demo.local", **PW, "name": "Q4 (pc02)"}], "formats": []})
    with pytest.raises(sc.InvalidName, match="already exists"):
        sc._m_inventory_multi(None, {"targets": [
            {"host": "pc02.demo.local", **PW, "name": "Q4 (pc02)"},
            {"host": "pc03.demo.local", **PW, "name": "Q4 (pc03)"}], "formats": [],
            "combine": True, "combined_name": "q4 (PC02)"})


# ─── list_clusters ─────────────────────────────────────────────────────


def test_list_clusters_v4_excludes_the_prism_central_pseudo_cluster(demo: Any) -> None:
    result = list_clusters({"host": "pc01.demo.local", **PW})
    assert result["api"] == "v4"
    names = [c["name"] for c in result["clusters"]]
    assert names == ["AMS-EDGE-02", "PHX-DR-01", "PHX-PROD-01", "PHX-PROD-02"]
    assert "pc01" not in names
    edge = result["clusters"][0]
    assert (edge["aos_version"], edge["num_nodes"], edge["hypervisors"]) == ("6.10", 2, ["ESXi"])


def test_list_clusters_v3_excludes_the_prism_central_pseudo_cluster(demo: Any) -> None:
    result = list_clusters({"host": "pc03.demo.local", **PW})
    assert result["api"] == "v3"
    assert [c["name"] for c in result["clusters"]] == ["DAL-EDGE-01", "DAL-PROD-01"]
    assert all(c["num_nodes"] and c["hypervisors"] == ["AHV"] for c in result["clusters"])


def test_list_clusters_single_cluster_and_prism_element(demo: Any) -> None:
    assert [c["name"] for c in _clusters("tiny.demo.local")] == ["LAB-01"]
    pe = list_clusters({"mode": "prism_element", "host": "pe01.demo.local", **PW})
    assert pe["api"] == "v2" and len(pe["clusters"]) == 1
    assert pe["clusters"][0]["name"] == "phx-dr-01"


def test_list_clusters_errors(demo: Any) -> None:
    with pytest.raises(AuthenticationError):
        list_clusters({"host": "pc01.demo.local", "username": "admin", "password": ""})
    with pytest.raises(AuthenticationError):
        list_clusters({"host": "pc01.demo.local", "username": "admin", "password": "wrong"})
    with pytest.raises(NutanixApiError):
        list_clusters({"host": "down.demo.local", **PW})


def test_list_clusters_through_the_sidecar(demo: Any, data_dir: Path) -> None:
    result = sc._METHODS["list_clusters"](None, {"host": "pc01.demo.local", **PW})
    assert len(result["clusters"]) == 4


# ─── scoped collection ─────────────────────────────────────────────────


def _owner(obj: Any) -> tuple[str | None, str | None]:
    return getattr(obj, "cluster_uuid", None), getattr(obj, "cluster_name", None)


@pytest.mark.parametrize("host", ["pc01.demo.local", "pc03.demo.local"])
def test_scoped_run_keeps_only_the_selected_clusters(demo: Any, host: str) -> None:
    registered = _clusters(host)
    chosen = registered[:-1]
    left_out = registered[-1]
    full = _collect(host)  # the demo estate does not drift outside the sidecar
    env = _collect(host, [c["uuid"] for c in chosen])

    assert {c.name for c in env.clusters} == {c["name"] for c in chosen}
    for field in ("clusters", *CLUSTER_BOUND_FIELDS):
        for obj in getattr(env, field):
            uuid, name = (obj.uuid, obj.name) if field == "clusters" else _owner(obj)
            assert uuid != left_out["uuid"], (field, obj)
            assert name not in (left_out["name"], left_out["uuid"]), (field, obj)
    for image in env.images:
        assert left_out["uuid"] not in image.cluster_uuids
    # Something was actually left out, and the selected clusters lost nothing.
    assert len(env.vms) < len(full.vms)
    kept = {c["uuid"] for c in chosen}
    assert len(env.vms) == sum(1 for v in full.vms if v.cluster_uuid in kept)
    assert len(env.hosts) == sum(1 for h in full.hosts if h.cluster_uuid in kept)

    # Prism-Central-wide data is collected whatever the scope.
    for field in ("users", "roles", "protection_policies", "recovery_plans",
                  "security_policies", "categories", "service_groups", "playbooks",
                  "templates"):
        assert len(getattr(env, field)) == len(getattr(full, field)), field

    scope = env.cluster_scope
    assert scope is not None and scope.registered_count == len(registered)
    assert [c.name for c in scope.excluded] == [left_out["name"]]
    assert sorted(c.uuid for c in scope.selected) == sorted(kept)
    assert describe(scope).startswith(f"{len(chosen)} of {len(registered)} clusters (")
    assert env.alerts_summary.total_count == len(env.alerts)


def test_v4_lists_are_filtered_server_side(demo: Any) -> None:
    registered = _clusters("pc01.demo.local")
    keep = registered[0]["uuid"]
    demo.api.filters.clear()
    env = _collect("pc01.demo.local", [keep])
    served = {path for (host, path) in demo.api.filters if host == "pc01.demo.local"}
    assert served == {"vmm/v4.0/ahv/config/vms", "vmm/v4.0/esxi/config/vms",
                      "clustermgmt/v4.0/config/storage-containers",
                      "monitoring/v4.0/serviceability/alerts"}
    assert all(v.cluster_uuid == keep for v in env.vms)


def test_a_rejected_filter_falls_back_to_client_side_pruning(
        demo: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    from ntnxvision_core.synthetic import mockapi

    monkeypatch.setattr(mockapi, "FILTERABLE", {})  # every $filter answers HTTP 400
    registered = _clusters("pc01.demo.local")
    keep = {c["uuid"] for c in registered[1:3]}
    env = _collect("pc01.demo.local", sorted(keep))
    assert env.vms and all(v.cluster_uuid in keep for v in env.vms)
    assert env.storage_containers and all(s.cluster_uuid in keep
                                          for s in env.storage_containers)
    assert all(a.cluster_uuid in keep for a in env.alerts)
    assert not env.collection_notes.get("vms", "ok").startswith("error")


def test_unscoped_runs_record_every_cluster(demo: Any) -> None:
    env = _collect("pc01.demo.local")
    assert env.cluster_scope is not None and not env.cluster_scope.is_partial
    assert describe(env.cluster_scope) == "All 4 clusters"
    tiny = _collect("tiny.demo.local")
    assert describe(tiny.cluster_scope) == "1 cluster"


def test_a_selection_matching_nothing_fails_before_collecting(demo: Any) -> None:
    with pytest.raises(NutanixApiError, match="None of the selected clusters"):
        _collect("pc01.demo.local", ["not-a-registered-cluster"])


def test_sidecar_run_records_the_scope(demo: Any, data_dir: Path) -> None:
    registered = _clusters("pc01.demo.local")
    with pytest.raises(NutanixApiError, match="at least one cluster"):
        _inventory(data_dir, "empty", clusters=[])
    result = _inventory(data_dir, "PHX prod only", formats=["excel", "pdf"],
                        clusters=[c["uuid"] for c in registered if "PROD" in c["name"]])
    assert result["scope"]["label"] == "2 of 4 clusters (PHX-PROD-01, PHX-PROD-02)"
    assert result["summary"]["clusters"] == 2
    detail = sc._m_get_snapshot(None, {"snapshot": result["snapshot_path"]})
    assert detail["scope"]["registered_count"] == 4
    assert detail["run"]["scope"]["label"].startswith("2 of 4")
    listing = sc._m_list_snapshots(None, {})["snapshots"]
    assert listing[0]["scope"]["label"].startswith("2 of 4")

    # The scope line is in the Excel Overview sheet.
    xlsx = next(r["path"] for r in result["reports"] if r["format"] == "excel")
    ws = load_workbook(xlsx)["Overview"]
    rows = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1)}
    assert rows["Cluster Scope"] == "2 of 4 clusters (PHX-PROD-01, PHX-PROD-02)"


def test_reports_state_the_scope(demo: Any) -> None:
    from ntnxvision_core.exporters import ExportOptions
    from ntnxvision_core.exporters.content import build_report

    registered = _clusters("pc01.demo.local")
    env = _collect("pc01.demo.local", [registered[2]["uuid"]])
    report = build_report(env, ExportOptions(), "Inventory report")
    meta = dict(report.cover.meta)
    assert meta["Cluster scope"] == "1 of 4 clusters (PHX-PROD-01)"
    summary = next(c for c in report.chapters if c.key == "executive_summary")
    text = " ".join(getattr(b, "text", "") for b in summary.blocks)
    assert "Scope: 1 of 4 clusters (PHX-PROD-01)" in text
    assumptions = next(c for c in report.chapters if c.key == "assumptions")
    bullets = [i for b in assumptions.blocks for i in getattr(b, "items", [])]
    assert any(i.startswith("Scope: 1 of 4 clusters") for i in bullets)


# ─── compare ───────────────────────────────────────────────────────────


def _scoped(env: Environment, keep: set[str]) -> Environment:
    """A copy of `env` as if its run had selected only `keep` (by cluster name)."""
    from ntnxvision_core.scope import RegisteredCluster, Scope

    copy = env.model_copy(deep=True)
    scope = Scope([RegisteredCluster(uuid=c.uuid, name=c.name) for c in env.clusters],
                  [c.uuid for c in env.clusters if c.name in keep])
    scope.prune(copy)
    copy.cluster_scope = scope.record()
    return copy


def test_compare_does_not_report_left_out_clusters_as_removed(demo: Any, tmp_path: Path
                                                              ) -> None:
    base = _collect("pc01.demo.local")
    cur = _scoped(base, {"PHX-PROD-01", "PHX-DR-01"})
    result = compare(base, cur)
    # Same data, smaller scope: nothing added, removed or changed (image placements on
    # the left-out clusters included).
    assert result["summary"] == {"added": 0, "removed": 0, "changed": 0, "total": 0}
    notice = result["scope_notice"]
    assert notice is not None
    assert notice["baseline"] == "All 4 clusters"
    assert notice["current"] == "2 of 4 clusters (PHX-DR-01, PHX-PROD-01)"
    assert notice["clusters"] == ["AMS-EDGE-02", "PHX-PROD-02"]
    assert notice["out_of_scope"] > 0
    assert "only one snapshot collected" in notice["detail"]
    vms = next(f for f in result["figures"] if f["label"] == "VMs")
    assert vms["value"] == len(cur.vms) and vms["qualifier"].endswith(f"was {len(cur.vms):,}")

    # ...nor added, the other way round.
    back = compare(cur, base)
    assert back["summary"]["added"] == 0 and back["scope_notice"] is not None

    # The comparison exports carry the notice.
    xlsx = export_comparison(base, cur, "full", "prod", "excel", tmp_path)
    ws = load_workbook(xlsx)["Summary"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert "Current scope" in cells
    assert any("only one snapshot collected" in c for c in cells)
    from ntnxvision_core.exporters.compare_pdf import build_comparison_report

    report = build_comparison_report(base, cur, "full", "prod")
    sections = [s.title for ch in report.chapters for s in ch.sections]
    assert "Cluster Scope" in sections


def test_same_scope_has_no_notice(demo: Any) -> None:
    base = _collect("pc01.demo.local")
    assert compare(base, base)["scope_notice"] is None
    legacy = base.model_copy(deep=True)
    legacy.cluster_scope = None
    assert compare(legacy, base)["scope_notice"] is None


def test_scope_record_round_trips_through_a_snapshot(tmp_path: Path) -> None:
    env = Environment(cluster_scope=ClusterScope(
        selected=[ScopeCluster(uuid="a", name="A")],
        excluded=[ScopeCluster(uuid="b", name="B")], registered_count=2))
    store = SnapshotStore(tmp_path)
    loaded = store.load(store.save(env, "s"))
    assert loaded.cluster_scope == env.cluster_scope
    assert describe(loaded.cluster_scope) == "1 of 2 clusters (A)"
