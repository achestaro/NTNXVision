"""Report sections: selection and custom order (exporters/sections.py).

Covers the order normalisation (unknown and missing ids), the ExportOptions round trip,
chapter order and renumbering in the shared content model, the Word and PDF chapter
order, the Excel sheet grouping, cross references after reordering, the availability
predicates against the builders, and the TypeScript catalogue mirror.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest

from ntnxvision_core.exporters import ExportOptions, export_environment
from ntnxvision_core.exporters import sections as S
from ntnxvision_core.exporters.content import (
    Figure,
    FigureRow,
    Paragraph,
    Report,
    SceneFigure,
    Table,
    build_report,
)
from ntnxvision_core.models import Environment

REPO = Path(__file__).resolve().parents[2]

#: A custom order: services first, design last; the fixed sections are left out on
#: purpose (they keep their place anyway).
CUSTOM = ["data_protection", "networking", "workloads", "health", "lcm", "security",
          "physical_design", "logical_design", "prism_central"]


@pytest.fixture(scope="module")
def estates(_demo_state: Any) -> dict[str, Environment]:
    from ntnxvision_core.api import PrismCredentials
    from ntnxvision_core.api.element_client import ElementCredentials
    from ntnxvision_core.engine import collect_element_environment, collect_environment
    from ntnxvision_core.synthetic import demo
    from ntnxvision_core.synthetic.scenarios import collect_scenario

    demo.reset()
    out = {
        "pc01": collect_environment(PrismCredentials(host="pc01.demo.local",
                                                     username="admin", password="demo")),
        "pe01": collect_element_environment(ElementCredentials(
            host="pe01.demo.local", username="admin", password="demo")),
        "medium": collect_scenario("medium"),
    }
    demo.reset()
    return out


# ─── order normalisation ──────────────────────────────────────────────────


def test_order_default_and_none() -> None:
    assert S.order(None) == S.DEFAULT_ORDER
    assert S.order([]) == S.DEFAULT_ORDER
    assert S.DEFAULT_ORDER[0] == "executive_summary" and S.DEFAULT_ORDER[-1] == "appendix"


def test_order_ignores_unknown_and_appends_missing() -> None:
    got = S.order(["lcm", "bogus", "health", "lcm", 42, "networking"])
    assert got[:4] == ("executive_summary", "lcm", "health", "networking")
    assert got[-1] == "appendix"
    assert sorted(got) == sorted(S.DEFAULT_ORDER)  # every id once
    rest = [s for s in S.DEFAULT_ORDER[1:-1] if s not in ("lcm", "health", "networking")]
    assert list(got[4:-1]) == rest  # missing ones in default order


def test_fixed_sections_cannot_move() -> None:
    got = S.order(["appendix", "executive_summary", "ndb"])
    assert got[0] == "executive_summary" and got[1] == "ndb" and got[-1] == "appendix"


def test_options_round_trip() -> None:
    opts = ExportOptions.coerce({"chapters": ["workloads", "networking"],
                                 "section_order": ["networking", "workloads"]})
    assert opts.chapters == frozenset({"workloads", "networking"})
    assert opts.section_order == ("networking", "workloads")
    assert opts.includes_chapter("workloads", "workloads_storage")
    assert not opts.includes_chapter("catalog", "workloads_storage")
    plain = ExportOptions.coerce({"sections": ["clusters_hosts"]})
    assert plain.chapters is None and plain.section_order is None
    assert plain.includes_chapter("health", "clusters_hosts")
    assert not plain.includes_chapter("workloads", "workloads_storage")
    # Both filters apply together.
    both = ExportOptions.coerce({"sections": ["clusters_hosts"], "chapters": ["workloads"]})
    assert not both.includes_chapter("workloads", "workloads_storage")


# ─── content model ────────────────────────────────────────────────────────


def _labels(report: Report) -> list[tuple[str, str]]:
    """(chapter number, caption label) for every table and figure, in document order."""
    out: list[tuple[str, str]] = []

    def visit(num: str, blocks: list[Any]) -> None:
        for b in blocks:
            if isinstance(b, (Table, Figure, SceneFigure)):
                out.append((num, b.label))
            elif isinstance(b, FigureRow):
                visit(num, list(b.figures))

    for ch in report.chapters:
        visit(ch.number, ch.blocks)
        for s in ch.sections:
            visit(ch.number, s.blocks)
    return out


def _check_numbering(report: Report) -> None:
    body = [c for c in report.chapters if not c.appendix]
    assert [c.number for c in body] == [str(i) for i in range(1, len(body) + 1)]
    appx = [c for c in report.chapters if c.appendix]
    assert [c.number for c in appx] == [chr(ord("A") + i) for i in range(len(appx))]
    for ch in report.chapters:
        assert [s.number for s in ch.sections] == \
            [f"{ch.number}.{i}" for i in range(1, len(ch.sections) + 1)]
    counters: dict[tuple[str, str], int] = {}
    for num, label in _labels(report):
        kind = label.split(" ", 1)[0]
        counters[(num, kind)] = counters.get((num, kind), 0) + 1
        assert label == f"{kind} {num}-{counters[(num, kind)]}."


@pytest.mark.parametrize("name", ["pc01", "pe01", "medium"])
def test_custom_order_reorders_and_renumbers(estates: dict[str, Environment],
                                              name: str) -> None:
    env = estates[name]
    default = build_report(env, ExportOptions(), "x")
    custom = build_report(env, ExportOptions(section_order=tuple(CUSTOM)), "x")
    keys = [c.key for c in custom.chapters]
    present = {c.key for c in default.chapters}
    assert keys == [k for k in S.order(CUSTOM) if k in present]
    assert sorted(keys) == sorted(c.key for c in default.chapters)
    _check_numbering(default)
    _check_numbering(custom)


def test_default_order_option_is_identical(estates: dict[str, Environment]) -> None:
    env = estates["medium"]
    a = build_report(env, ExportOptions(), "x")
    b = build_report(env, ExportOptions(section_order=S.DEFAULT_ORDER,
                                        chapters=frozenset(S.DEFAULT_ORDER)), "x")
    assert a.outline() == b.outline()
    assert _labels(a) == _labels(b)


def test_selection_drops_chapters(estates: dict[str, Environment]) -> None:
    env = estates["pc01"]
    keep = {"executive_summary", "networking", "workloads", "appendix"}
    rep = build_report(env, ExportOptions(chapters=frozenset(keep),
                                          section_order=("networking", "workloads")), "x")
    assert [c.key for c in rep.chapters] == ["executive_summary", "networking", "workloads",
                                             "appendix"]
    assert [c.number for c in rep.chapters] == ["1", "2", "3", "A"]
    _check_numbering(rep)


def _dp_xref(rep: Report) -> str | None:
    (dp,) = [c for c in rep.chapters if c.key == "data_protection"]
    texts = [b.text for b in dp.blocks if isinstance(b, Paragraph)]
    return next((t for t in texts if "Solution Conceptual Overview" in t), None)


def test_cross_reference_follows_reordering(estates: dict[str, Environment]) -> None:
    env = estates["medium"]
    for order in (None, ("data_protection", "logical_design")):
        rep = build_report(env, ExportOptions(section_order=order), "x")
        (ld,) = [c for c in rep.chapters if c.key == "logical_design"]
        first = next(b for s in ld.sections for b in s.blocks if isinstance(b, SceneFigure))
        xref = _dp_xref(rep)
        assert xref is not None and first.label.rstrip(".") in xref
        assert first.label.startswith(f"Figure {ld.number}-")


def test_cross_reference_dropped_without_logical_design(
        estates: dict[str, Environment]) -> None:
    env = estates["medium"]
    chapters = frozenset(S.DEFAULT_ORDER) - {"logical_design"}
    rep = build_report(env, ExportOptions(chapters=chapters), "x")
    assert _dp_xref(rep) is None


# ─── availability ─────────────────────────────────────────────────────────


@pytest.mark.parametrize("name", ["pc01", "pe01", "medium"])
def test_availability_matches_builders(estates: dict[str, Environment], name: str) -> None:
    env = estates[name]
    built = {c.key for c in build_report(env, ExportOptions(), "x").chapters}
    avail = {a["id"] for a in S.availability(env) if a["available"]}
    assert avail == built


def test_availability_hints() -> None:
    pe = Environment(source_type="prism_element")
    by_id = {a["id"]: a for a in S.availability(pe)}
    assert by_id["prism_central"]["hint"] == "Not part of a Prism Element collection"
    assert by_id["ndb"] == {"id": "ndb", "available": False,
                            "hint": "No data in this snapshot"}
    assert by_id["executive_summary"]["available"]


# ─── renderers ────────────────────────────────────────────────────────────


def test_word_chapter_order(estates: dict[str, Environment], tmp_path: Path) -> None:
    from docx import Document

    env = estates["medium"]
    opts = {"section_order": CUSTOM}
    path = export_environment(env, "word", tmp_path, "r", opts)
    heads = [p.text for p in Document(str(path)).paragraphs
             if p.style is not None and p.style.name == "Heading 1"]
    expected = [c.heading for c in build_report(env, ExportOptions.coerce(opts), "x").chapters]
    assert heads == expected
    assert heads[0].startswith("1. Executive Summary")
    assert heads[1] == "2. Data Protection & Disaster Recovery"


def test_pdf_chapter_order(estates: dict[str, Environment], tmp_path: Path,
                           monkeypatch: pytest.MonkeyPatch) -> None:
    from ntnxvision_core.exporters import pdf as P

    seen: list[str] = []
    orig = P._Pdf.start_section

    def spy(self: Any, name: str, level: int = 0, strict: bool = True) -> Any:
        if level == 0:
            seen.append(name)
        return orig(self, name, level, strict)

    monkeypatch.setattr(P._Pdf, "start_section", spy)
    env = estates["medium"]
    opts = {"section_order": CUSTOM, "chapters": [*S.DEFAULT_ORDER[:-1]]}
    export_environment(env, "pdf", tmp_path, "r", opts)
    expected = [c.heading for c in build_report(env, ExportOptions.coerce(opts), "x").chapters]
    assert seen[1:] == expected  # after "Contents"
    assert not any(h.startswith("Appendix") for h in seen)
    assert seen[2] == "2. Data Protection & Disaster Recovery"


def test_excel_sheets_follow_order_and_selection(estates: dict[str, Environment],
                                                 tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from ntnxvision_core.exporters.excel import _SPECS, sheet_specs

    assert sheet_specs(ExportOptions()) == list(_SPECS)
    env = estates["pc01"]
    default = load_workbook(export_environment(env, "excel", tmp_path, "a")).sheetnames
    opts = {"section_order": ["networking", "workloads"],
            "chapters": [s for s in S.DEFAULT_ORDER if s != "identity"]}
    wb = load_workbook(export_environment(env, "excel", tmp_path, "b", opts))
    names = wb.sheetnames
    assert names[0] == "Overview"
    assert "Users" in default and "Users" not in names
    assert names.index("Subnets") < names.index("VMs") < names.index("Clusters")
    # The Overview index says why an excluded type has no sheet.
    values = [c.value for row in wb["Overview"].iter_rows() for c in row]
    assert "Not included in this export" in values


# ─── TypeScript mirror ────────────────────────────────────────────────────


def test_typescript_catalogue_matches() -> None:
    src = (REPO / "app" / "src" / "shared" / "reportSections.ts").read_text("utf-8")
    ids = re.findall(r"\{ id: '([a-z_]+)', label: '([^']+)'(?:, fixed: '(first|last)')? \}",
                     src)
    assert [(i, lbl, f or None) for i, lbl, f in ids] == \
        [(s.id, s.label, s.fixed) for s in S.CATALOG]
