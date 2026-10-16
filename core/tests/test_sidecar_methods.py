"""Every sidecar method the Runway UI calls, exercised in-process on synthetic data."""

from __future__ import annotations

import csv
import io
import shutil
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document

import ntnxvision_core.sidecar as sc
from ntnxvision_core.api.errors import NutanixApiError

AREA_STATUSES = {"collected", "notes", "skipped", "failed"}


def _run(host: str, name: str, data_dir: Path, **extra: Any) -> dict[str, Any]:
    return sc._m_inventory(None, {
        "host": host, "username": "admin", "password": "demo", "port": 9440,
        "insecure": False, "name": name, "formats": extra.pop("formats", []),
        "export_dir": str(data_dir / "exports"), **extra,
    })


@pytest.fixture
def two_snapshots(demo: Any, data_dir: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    first = _run("pc01.demo.local", "pc01-week1", data_dir, formats=["excel"])
    second = _run("pc01.demo.local", "pc01-week2", data_dir)  # drifted one epoch
    return first, second


def test_list_areas_catalogue() -> None:
    result = sc._m_list_areas(None, {})
    ids = [a["id"] for a in result["areas"]]
    assert len(ids) == len(set(ids)) == 21
    groups = [a["group"] for a in result["areas"]]
    assert groups == sorted(groups, key=["Infrastructure", "Platform services",
                                         "Operations and governance"].index)
    assert {a["id"] for a in result["areas"] if a["pe_supported"]} == {
        "clusters", "hosts", "storage_containers", "subnets", "vms"}
    assert result["typical_seconds"] == {"prism_central": 240, "prism_element": 60}


def test_inventory_result_carries_run_metadata(two_snapshots: Any) -> None:
    first, _ = two_snapshots
    assert Path(first["snapshot_path"]).is_file()
    assert first["snapshot_name"] == "pc01-week1"
    assert first["duration_ms"] >= 0 and first["started_at"] <= first["finished_at"]
    assert len(first["areas"]) == 21
    assert {a["status"] for a in first["areas"]} <= AREA_STATUSES
    assert all(a["duration_ms"] is not None for a in first["areas"])
    assert first["reports"][0]["format"] == "excel"
    assert first["reports"][0]["detail"].endswith("sheets")
    assert first["reports"][0]["size_bytes"] > 0
    vms = next(f for f in first["figures"] if f["label"] == "VMs")
    assert vms["qualifier"] == "3 skipped" and vms["qualifier_tone"] == "warn"
    apis = {row["api"] for row in first["api_routing"]}
    assert {"v4", "v3"} <= apis
    stops = [row for row in first["api_routing"] if row["after_sunset"]["tone"] == "warn"]
    assert any("X-Play" in row["areas"] for row in stops)
    assert first["compat"] == {"version": "pc.2024.3", "supported": True,
                               "range": "pc.2024.1 – 7.5.1.6"}


def test_overview(demo: Any, data_dir: Path) -> None:
    empty = sc._m_get_overview(None, {})
    assert empty == {"latest": None, "figures": [], "recent": [], "coverage": None,
                     "api_routing": None, "sunset_areas": [], "compare_baseline": None}
    first = _run("pc01.demo.local", "a", data_dir)
    _run("pc01.demo.local", "b", data_dir)
    ov = sc._m_get_overview(None, {})
    assert ov["latest"]["name"] == "b"
    assert ov["compare_baseline"] == first["snapshot_path"]
    figures = {f["label"]: f for f in ov["figures"]}
    assert figures["Hosts"]["qualifier"].startswith("+2 since")
    assert figures["Clusters"]["qualifier"] == "AOS 6.10 – 7.0"
    assert figures["Storage"]["unit"] == "TB"
    assert figures["Protected VMs"]["value"].endswith("%")
    assert ov["coverage"]["total"] == 21 and ov["coverage"]["skipped"] == 1
    assert ov["api_routing"]["v4"] > ov["api_routing"]["v3"]
    assert "X-Play playbooks" in ov["sunset_areas"]


def test_list_and_get_snapshot(two_snapshots: Any, data_dir: Path) -> None:
    listing = sc._m_list_snapshots(None, {})
    assert [s["name"] for s in listing["snapshots"]] == ["pc01-week2", "pc01-week1"]
    assert listing["dir"].startswith(str(data_dir))
    assert listing["total_bytes"] == sum(s["size_bytes"] for s in listing["snapshots"])
    first = listing["snapshots"][1]
    assert first["kind"] == "pc" and first["source"] == "pc01.demo.local"
    detail = sc._m_get_snapshot(None, {"snapshot": first["path"]})
    assert detail["is_combined"] is False and detail["source_type"] == "prism_central"
    assert detail["run"]["areas"] and detail["run"]["reports"][0]["format"] == "excel"
    assert set(detail["export_estimates"]) == {"excel", "pdf", "word"}
    assert detail["collection_notes"]["files"].startswith("skipped")


def test_query_inventory_paging_filter_sort(two_snapshots: Any) -> None:
    _, second = two_snapshots
    path = second["snapshot_path"]

    def q(**kw: Any) -> dict[str, Any]:
        return sc._m_query_inventory(None, {"snapshot": path, "object_type": "vms",
                                            "filter": kw.get("filter", {}),
                                            "sort": kw.get("sort"),
                                            "page": kw.get("page", {"offset": 0, "limit": 25})})

    page1 = q()
    assert page1["total"] == page1["matching"] == len(page1["rows"]) * 0 + page1["total"]
    assert len(page1["rows"]) == 25 and page1["label"] == "Virtual machines"
    assert [c["key"] for c in page1["columns"]][:3] == ["name", "power", "cluster"]
    page2 = q(page={"offset": 25, "limit": 25})
    assert {r["id"] for r in page1["rows"]}.isdisjoint({r["id"] for r in page2["rows"]})

    k8s = q(filter={"text": "K8S-WORKER"})
    assert k8s["matching"] > 0 and all("k8s-worker" in r["cells"]["name"]
                                       for r in k8s["rows"])
    off = q(filter={"facets": {"power": "Off"}})
    assert all(r["cells"]["power"]["label"] == "Off" for r in off["rows"])
    power = next(f for f in page1["facets"] if f["key"] == "power")
    assert sum(o["count"] for o in power["options"]) == page1["total"]

    desc = q(sort={"key": "memory", "dir": "desc"})["rows"]
    mems = [r["cells"]["memory"] for r in desc]
    assert mems == sorted(mems, reverse=True)
    asc = q(sort={"key": "name", "dir": "asc"}, page={"offset": 0, "limit": 50})["rows"]
    names = [r["cells"]["name"].lower() for r in asc]
    assert names == sorted(names)

    counts = {t["type"]: t["count"] for t in page1["type_counts"]}
    assert counts["vms"] == page1["total"] and counts["file_servers"] is None
    assert counts["object_stores"] == 1


def test_query_other_types_and_detail(two_snapshots: Any) -> None:
    path = two_snapshots[1]["snapshot_path"]
    for object_type in ("clusters", "hosts", "storage_containers", "images", "alerts",
                        "lcm_entities", "security_policies", "users"):
        res = sc._m_query_inventory(None, {"snapshot": path, "object_type": object_type,
                                           "filter": {}, "sort": None,
                                           "page": {"offset": 0, "limit": 5}})
        assert res["rows"], object_type
        detail = sc._m_get_object_detail(None, {"snapshot": path, "object_type": object_type,
                                                "id": res["rows"][0]["id"]})
        assert detail["fields"] and detail["name"]
    vms = sc._m_query_inventory(None, {"snapshot": path, "object_type": "vms",
                                       "filter": {"text": "db-01"}, "sort": None,
                                       "page": {"offset": 0, "limit": 1}})
    vm = sc._m_get_object_detail(None, {"snapshot": path, "object_type": "vms",
                                        "id": vms["rows"][0]["id"]})
    labels = [f["label"] for f in vm["fields"]]
    for expected in ("Host", "Cluster", "Size", "Storage", "Protection", "Disk 1", "NIC 1",
                     "UUID"):
        assert expected in labels
    assert vm["categories"] and vm["prism_url"].startswith("https://")
    with pytest.raises(NutanixApiError):
        sc._m_get_object_detail(None, {"snapshot": path, "object_type": "vms",
                                       "id": "no-such-vm"})


def test_compare_and_export_comparison(two_snapshots: Any, data_dir: Path) -> None:
    a, b = (s["snapshot_path"] for s in two_snapshots)
    result = sc._m_compare_snapshots(None, {"baseline": a, "current": b})
    assert result["baseline"]["name"] == "pc01-week1"
    assert result["summary"]["total"] == 53
    assert result["figures"][0]["label"] == "Changes"
    assert sum(r["count"] for r in result["changes"] if r["kind"] == "added") == 36
    xlsx = sc._m_export_comparison(None, {"baseline": a, "current": b, "format": "excel",
                                          "export_dir": str(data_dir / "exp")})
    assert Path(xlsx["path"]).name == "changes-pc01-week1-to-pc01-week2.xlsx"
    out = sc._m_export_comparison(None, {"baseline": a, "current": b, "format": "csv",
                                         "export_dir": str(data_dir / "exp")})
    rows = list(csv.reader(io.StringIO(Path(out["path"]).read_text(encoding="utf-8"))))
    assert rows[0][0] == "Change" and len(rows) - 1 >= 53


def _make_dotx(tmp: Path) -> Path:
    """A customer template: a .docx with the template content type and a header."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "ACME CONFIDENTIAL"
    doc.add_paragraph("Template boilerplate that must not survive")
    src = tmp / "base.docx"
    doc.save(str(src))
    dotx = tmp / "acme.dotx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dotx, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml",
                                    b"wordprocessingml.template.main+xml")
            zout.writestr(item, data)
    return dotx


def test_export_options_sections_template_and_cover(two_snapshots: Any, data_dir: Path,
                                                    tmp_path: Path) -> None:
    path = two_snapshots[0]["snapshot_path"]
    dotx = _make_dotx(tmp_path)
    res = sc._m_export(None, {
        "snapshot": path, "formats": ["word", "pdf"], "export_dir": str(data_dir / "exp"),
        "sections": ["clusters_hosts", "collection_notes"],
        "word": {"template": str(dotx), "prepared_for": "Acme Corp", "author": "A. Tester"},
    })
    assert res["errors"] == {}
    assert {f["format"] for f in res["files"]} == {"word", "pdf"}
    word = Document(next(f["path"] for f in res["files"] if f["format"] == "word"))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "Template boilerplate" not in text
    assert word.sections[0].header.paragraphs[0].text == "ACME CONFIDENTIAL"
    assert "Prism Central — Management & Sizing" in text
    assert "Executive Summary" not in text and "Data Protection" not in text
    assert word.core_properties.author == "A. Tester"
    assert "Acme Corp" in word.core_properties.subject
    # Export history is recorded on the snapshot.
    detail = sc._m_get_snapshot(None, {"snapshot": path})
    assert {e["format"] for e in detail["exports"]} >= {"excel", "word", "pdf"}


def test_export_with_missing_template_reports_an_error(two_snapshots: Any,
                                                       data_dir: Path) -> None:
    res = sc._m_export(None, {"snapshot": two_snapshots[0]["snapshot_path"],
                              "formats": ["word", "excel"], "export_dir": str(data_dir),
                              "word": {"template": "/nope/missing.dotx"}})
    assert "word" in res["errors"] and res["paths"].get("excel")


def test_rename_delete_clear(two_snapshots: Any, data_dir: Path) -> None:
    a, b = (s["snapshot_path"] for s in two_snapshots)
    renamed = sc._m_rename_snapshot(None, {"snapshot": b, "new_name": "baseline"})
    assert renamed["name"] == "baseline" and Path(renamed["path"]).is_file()
    assert not Path(b).exists()
    names = [s["name"] for s in sc._m_list_snapshots(None, {})["snapshots"]]
    assert "baseline" in names
    with pytest.raises(NutanixApiError):
        sc._m_rename_snapshot(None, {"snapshot": a, "new_name": "baseline"})
    assert sc._m_delete_snapshot(None, {"snapshot": a}) == {"deleted": True}
    handle = sc._runs.begin("busy-run")
    try:
        with pytest.raises(sc.Busy):
            sc._m_clear_snapshots(None, {})
    finally:
        sc._runs.end(handle.run_id)
    assert sc._m_clear_snapshots(None, {}) == {"deleted": 1}
    assert sc._m_list_snapshots(None, {})["snapshots"] == []
    assert (data_dir / "exports").exists()  # exports are never touched


def test_cancel_unknown_run() -> None:
    assert sc._m_cancel_run(None, {"run_id": "nope"}) == {"cancelled": False}


def test_cancelled_run_saves_nothing(demo: Any, data_dir: Path) -> None:
    handle = sc._runs.begin("cancel-me")
    handle.cancel()
    with pytest.raises(sc.RunCancelled):
        _run("tiny.demo.local", "never", data_dir, run_id="cancel-me", formats=["excel"])
    assert sc._m_list_snapshots(None, {})["snapshots"] == []
    assert not (data_dir / "exports").exists() or not any((data_dir / "exports").iterdir())


def test_committed_run_cannot_be_cancelled() -> None:
    handle = sc.RunHandle("r")
    handle.commit()
    assert handle.cancel() is False


def test_prism_element_inventory(demo: Any, data_dir: Path) -> None:
    res = sc._m_inventory_cluster(None, {"host": "pe01.demo.local", "username": "admin",
                                         "password": "x", "name": "phx-dr-01",
                                         "formats": []})
    assert res["source_type"] == "prism_element"
    assert [a["api"] for a in res["areas"]] == ["v2"] * 5
    snap = sc._m_list_snapshots(None, {})["snapshots"][0]
    assert snap["kind"] == "pe" and snap["version"] == "AOS 7.0"


def test_multi_pc_parallel_with_a_failing_target(demo: Any, data_dir: Path) -> None:
    res = sc._m_inventory_multi(None, {
        "targets": [
            {"host": "tiny.demo.local", "username": "admin", "password": "x"},
            {"host": "down.demo.local", "username": "admin", "password": "x"},
            {"host": "pc03.demo.local", "username": "admin", "password": "x"},
        ],
        "combine": True, "combined_name": "estate", "formats": [],
    })
    ok = {t["host"]: t["success"] for t in res["targets"]}
    assert ok == {"tiny.demo.local": True, "down.demo.local": False, "pc03.demo.local": True}
    combined = res["combined"]
    assert combined["summary"]["source_environments"] == 2
    assert combined["duration_ms"] is not None
    assert combined["areas"][0]["id"] == "tiny.demo.local: clusters"
    assert combined["areas"][0]["api"] == "v4"
    assert {a["id"].split(": ")[0] for a in combined["areas"]} == {
        "tiny.demo.local", "pc03.demo.local"}
    kinds = sorted(s["kind"] for s in sc._m_list_snapshots(None, {})["snapshots"])
    assert kinds == ["combined", "pc", "pc"]


def test_snapshot_paths_stay_confined(data_dir: Path, tmp_path: Path) -> None:
    outside = tmp_path / "x.ntnxsnap.json"
    shutil.copy(__file__, outside)
    with pytest.raises(NutanixApiError):
        sc._m_query_inventory(None, {"snapshot": str(outside), "object_type": "vms"})
