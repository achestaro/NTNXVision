"""The Compare page's "Export changes" output: Excel, CSV and PDF.

The workbook follows the inventory workbook's Runway styling (light header fill, iris
rule, frozen header, auto-filter, N/A in ink-3, change kinds in their status colour); the
PDF reuses the inventory PDF renderer and keeps the Cluster column first.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from ntnxvision_core.compare import export_comparison
from ntnxvision_core.exporters import theme as TH
from ntnxvision_core.exporters.naming import export_date
from ntnxvision_core.numfmt import NA


@pytest.fixture
def pair(demo: Any) -> tuple[Any, Any]:
    from ntnxvision_core.api import PrismCredentials
    from ntnxvision_core.engine import collect_environment
    from ntnxvision_core.synthetic import demo as demo_mode

    creds = PrismCredentials(host="pc01.demo.local", username="a", password="p")
    base = collect_environment(creds)
    demo_mode.advance("pc01.demo.local")
    return base, collect_environment(creds)


def test_excel_follows_the_inventory_workbook_style(pair: Any, tmp_path: Path) -> None:
    base, cur = pair
    wb = load_workbook(export_comparison(base, cur, "week1", "week2", "excel", tmp_path))
    assert wb.sheetnames == ["Summary", "Changes"]
    summary = wb["Summary"]
    assert summary["A3"].value == "COMPARISON WORKBOOK"
    assert summary["A1"].fill.fgColor.rgb.endswith(TH.IRIS)
    labels = {r[0]: r[1] for r in summary.iter_rows(values_only=True) if r[0]}
    assert labels["Total"] == 53 and labels["Added"] == 36
    assert labels["Baseline"] == "week1" and labels["Current"] == "week2"
    assert "Virtual machines" in labels  # the per-type count table

    ws = wb["Changes"]
    headers = [c.value for c in ws[1]]
    assert headers == ["Change", "Object type", "Object", "Location", "Field", "Before",
                       "After", "Detail", "ID"]
    head = ws["A1"]
    assert head.fill.fgColor.rgb.endswith(TH.HEADER_FILL)
    assert head.border.bottom.color.rgb.endswith(TH.IRIS)
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    rows = list(ws.iter_rows(min_row=2))
    kinds = [r[0].value for r in rows]
    assert kinds == sorted(kinds, key=["Changed", "Added", "Removed"].index)
    assert set(kinds) == {"Changed", "Added", "Removed"}
    added = next(r for r in rows if r[0].value == "Added")
    assert added[0].font.color.rgb.endswith(TH.OK)
    assert added[4].value == NA and added[4].font.color.rgb.endswith(TH.INK_3)


def test_csv_has_the_same_rows_with_na(pair: Any, tmp_path: Path) -> None:
    base, cur = pair
    out = export_comparison(base, cur, "week1", "week2", "csv", tmp_path)
    rows = list(csv.reader(out.read_text(encoding="utf-8").splitlines()))
    assert rows[0][0] == "Change" and len(rows) - 1 >= 53
    assert all(cell != "" for r in rows[1:] for cell in r[:7])


def test_pdf_report(pair: Any, tmp_path: Path) -> None:
    from ntnxvision_core.exporters.compare_pdf import build_comparison_report
    from ntnxvision_core.exporters.content import Status, Table

    base, cur = pair
    out = export_comparison(base, cur, "week1", "week2", "pdf", tmp_path)
    assert out.name == f"changes-week1-to-week2-{export_date()}.pdf"
    assert out.read_bytes().startswith(b"%PDF") and out.stat().st_size > 10_000

    report = build_comparison_report(base, cur, "week1", "week2")
    assert [c.title for c in report.chapters] == ["Summary", "Changes"]
    tables = [b for ch in report.chapters for s in ch.sections for b in s.blocks
              if isinstance(b, Table)]
    by_caption = {t.caption.split(":")[0]: t for t in tables}
    vms = by_caption["Virtual machines"]
    assert vms.headers[:3] == ["Cluster", "Change", "VM"]
    kinds = {r[1].text for r in vms.rows if isinstance(r[1], Status)}
    assert kinds == {"Added", "Removed", "Changed"}
    # Grouped like the screen: 32 added workers collapse into one row.
    assert any(r[3] == 32 for r in vms.rows)
    clusters = by_caption["Clusters"]
    assert clusters.headers[0] == "Cluster"
    aos = next(r for r in clusters.rows if r[2] == "AOS")
    assert (aos[3], aos[4]) == ("6.10.1", "7.0")
    for t in tables:
        if "Cluster" in t.headers:
            assert t.headers[0] == "Cluster", t.caption


def test_pdf_with_no_changes(demo: Any, tmp_path: Path) -> None:
    from ntnxvision_core.api import PrismCredentials
    from ntnxvision_core.engine import collect_environment

    env = collect_environment(PrismCredentials(host="pc01.demo.local", username="a",
                                               password="p"))
    out = export_comparison(env, env, "a", "b", "pdf", tmp_path)
    assert out.stat().st_size > 5_000
    wb = load_workbook(export_comparison(env, env, "a", "b", "excel", tmp_path))
    assert wb["Changes"]["A2"].value == "No changes between the two snapshots."


def test_csv_neutralises_formulas_and_quotes_per_rfc4180(
    pair: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R2-01: formula-like strings get a leading apostrophe, numbers stay numbers, and
    fields with CR / LF / quotes / commas round-trip through a CSV reader."""
    from ntnxvision_core import compare as C

    hostile = ["=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)", "\tTAB", "\rCR", "\nLF",
               "＝fullwidth", "＋fw", "－fw", "＠fw", "a,b", "say \"hi\"", "multi\r\nline", "safe"]
    rows = [["Changed", "VM", h, None, "f", 5, -3.5, "d", "id"] for h in hostile]
    monkeypatch.setattr(C, "_export_rows", lambda raw: rows)
    base, cur = pair
    out = export_comparison(base, cur, "w1", "w2", "csv", tmp_path)
    with out.open(newline="", encoding="utf-8") as fh:
        parsed = list(csv.reader(fh))[1:]
    got = [r[2] for r in parsed]
    for h, g in zip(hostile, got, strict=True):
        if h[0] in "=+-@\t\r\n＝＋－＠":
            assert g == "'" + h
        else:
            assert g == h
    assert parsed[0][3] == NA
    assert parsed[0][5] == "5" and parsed[0][6] == "-3.5"  # numbers are not prefixed
